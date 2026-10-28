"""Publish a focused, portable complaint draft from accepted human review decisions."""

from __future__ import annotations

import copy
import os
import shutil
import tempfile
from collections.abc import Callable
from math import ceil
from pathlib import Path, PureWindowsPath
from urllib.parse import quote
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from edf_bill_fetcher.application.record_review import review_rows
from edf_bill_fetcher.domain.json_values import object_value
from edf_bill_fetcher.helpers.excel_utils import literal_cell
from edf_bill_fetcher.helpers.identity import account_key
from edf_bill_fetcher.io.persistence.run_manifest import artifact_record
from edf_bill_fetcher.io.writers.evidence_bundle import (
    build_bundle_manifest,
    resolve_bundle_paths,
    save_evidence_files,
)
from edf_bill_fetcher.models.claim_tables import (
    complaint_paragraphs,
    group_table,
    overlap_table,
    reading_table,
    total_table,
)
from edf_bill_fetcher.models.rebilling_tables import history_table, schedule_table
from edf_bill_fetcher.models.report_view import ReportView
from edf_bill_fetcher.processors.preparation import identity

REFERENCE_MARKER = "EDF_COMPLAINT_SCHEDULE_REFERENCE"
NOTICE = (
    "DRAFT FOR HUMAN REVIEW — this pack presents a complaint and a proposed adjustment, "
    "not a determination of entitlement or current liability. Check eligibility, dates, "
    "amounts and evidence before sending. Whole original source files are retained: "
    "this is not a redacted pack, and shared PDFs may contain other pages. "
    "The accompanying schedule is a reference output; use the main review workbook or app to edit."
)
TableData = tuple[str, str, list[list[str]]]


def requested(view: ReportView) -> bool:
    """Only opted-in reviews with active challenges produce a complaint draft."""
    model = view.claims
    return bool(
        model
        and model["input"]["context"]["include_pack"]
        and (
            any(g["input"]["status"] != "withdrawn" for g in model["groups"])
            or model["unresolved_positions"]
        )
    )


def _compatible(a: dict, b: dict) -> bool:
    """Curate a history without guessing transfers between unrelated accounts/meters."""
    ar, br = a["record"], b["record"]
    if not account_key(identity(ar.get("Account"))) or account_key(
        identity(ar.get("Account"))
    ) != account_key(identity(br.get("Account"))):
        return False
    return not any(
        identity(ar.get(field))
        and identity(br.get(field))
        and identity(ar[field]) != identity(br[field])
        for field in ("Fuel", "Meter", "Register")
    )


def _selection(view: ReportView) -> tuple[dict, dict, list[str]]:
    """Retain complete linked/overlapping history, and honour record/position exclusions."""
    assert view.claims is not None
    assert view.rebilling_ledger is not None
    model, ledger = copy.deepcopy(view.claims), copy.deepcopy(view.rebilling_ledger)
    docs = {d["row_id"]: d for d in ledger["documents"] if d["included"]}
    seeds = {
        key
        for key, doc in docs.items()
        if doc["position"] and doc["position"]["status"] == "challenged"
    }
    groups, notes = [], []
    for group in model["groups"]:
        data = group["input"]
        if data["status"] == "withdrawn":
            continue
        keys = set(data["members"]) | {c["row_id"] for c in data["credits"]}
        if not keys <= docs.keys() or any(key not in seeds for key in data["members"]):
            notes.append(
                f"Claim {data['id']} has excluded or inactive members; its details were omitted pending review."
            )
            continue
        groups.append(group)
        seeds.update(keys)
    model["groups"] = groups
    kept_ids = {g["input"]["id"] for g in groups}
    model["overlaps"] = [o for o in model["overlaps"] if set(o["groups"]) <= kept_ids]
    selected = set(seeds)
    linked_pairs = {
        frozenset((link["older"], link["newer"]))
        for link in ledger["comparisons"]
        if link["status"] != "rejected"
    }
    while True:
        previous = set(selected)
        for key, doc in docs.items():
            if key in selected:
                continue
            for seed in previous:
                other = docs[seed]
                linked = frozenset((key, seed)) in linked_pairs
                periods = [
                    doc["period_from"],
                    doc["period_to"],
                    other["period_from"],
                    other["period_to"],
                ]
                overlap = not all(periods) or max(periods[0], periods[2]) <= min(
                    periods[1], periods[3]
                )
                if linked or (_compatible(doc, other) and overlap):
                    selected.add(key)
                    break
        if selected == previous:
            break
    ledger["documents"] = [doc for doc in ledger["documents"] if doc["row_id"] in selected]
    ledger["comparisons"] = [
        link for link in ledger["comparisons"] if {link["older"], link["newer"]} <= selected
    ]
    # Excluded/withdrawn group inputs are never included in an outward machine payload.
    model["input"]["groups"] = [g["input"] for g in groups]
    model["input"]["separations"] = [
        s for s in model["input"]["separations"] if set(s["groups"]) <= kept_ids
    ]
    for doc in ledger["documents"]:
        for field in ("Attachment Name", "Original Filename"):
            if doc["record"].get(field):
                doc["record"][field] = PureWindowsPath(str(doc["record"][field])).name
    return model, ledger, notes


def _evidence(
    view: ReportView, ledger: dict, root: Path, roots: list[Path]
) -> tuple[dict, list[str]]:
    """Copy identities only when their recorded hash verifies; never substitute by filename."""
    state = object_value(view.snapshot.prepared.extraction.legacy_state)
    all_rows = {item["id"]: item["record"] for item in review_rows(view.snapshot)}
    rows = [dict(all_rows[doc["row_id"]]) for doc in ledger["documents"]]
    manifest = state.get("source_manifest", {})
    source_ids = {identity(row.get("Source ID")) for row in rows}
    filtered = {key: data for key, data in manifest.items() if key in source_ids}
    paths = {}
    for directory in roots:
        found, _ = resolve_bundle_paths({"source_manifest": filtered}, str(directory))
        paths.update({key: path for key, path in found.items() if key not in paths})
    for key, metadata in filtered.items():
        # An original is allowed only when no saved copy was promised. A changed
        # or missing saved copy must remain a visible gap.
        if not metadata.get("Saved Path") and metadata.get("Original Path"):
            original = Path(metadata["Original Path"])
            try:
                if (
                    metadata.get("Content Hash")
                    and artifact_record(str(original))["sha256"] == metadata["Content Hash"]
                ):
                    paths[key] = str(original)
            except OSError:
                pass
    notes = []
    for row in rows:
        key = identity(row.get("Source ID"))
        metadata = filtered.get(key, {})
        expected = row.get("Content Hash") or metadata.get("Content Hash")
        if (
            not key
            or not expected
            or (metadata.get("Content Hash") and expected != metadata["Content Hash"])
        ):
            paths.pop(key, None)
            notes.append(
                f"Recorded source identity/hash unavailable or inconsistent: {identity(row.get('Invoice #')) or 'unnumbered document'}."
            )
        row["Content Hash"] = expected
        for field in ("Attachment Name", "Original Filename"):
            row[field] = PureWindowsPath(
                str(row.get(field) or metadata.get(field) or "source.pdf")
            ).name
    frame = pd.DataFrame(rows)
    saved = save_evidence_files(frame, paths, str(root / "evidence_files"))
    bundle = build_bundle_manifest(frame, saved, str(root))
    for source in bundle["files"]:
        if source["status"] != "saved":
            notes.append(
                f"Original evidence unavailable or changed: {source['attachment_name']} (source {source['source_id'] or 'unassigned'})."
            )
    for doc in ledger["documents"]:
        row = doc["record"]
        if not row.get("Page From") or not row.get("Page To"):
            notes.append(
                f"Original page range unverified: {identity(row.get('Invoice #')) or doc['row_id']}."
            )
        if (
            doc["position"]
            and doc["position"]["status"] == "challenged"
            and not doc["first_demand_date"]
        ):
            notes.append(
                f"Earliest demand not verified: {identity(row.get('Invoice #')) or doc['row_id']}. Obtain EDF's dated bill or statement."
            )
    return bundle, notes


def _source_rows(bundle: dict) -> list[list[str]]:
    rows = [
        [
            "Original source / identity",
            "Invoices",
            "Original PDF pages",
            "Included file / verification",
        ]
    ]
    for source in bundle["files"]:
        rows.append(
            [
                source["attachment_name"] + "\n" + (source["source_id"] or "Unassigned"),
                ", ".join(source["invoices"]),
                ", ".join("–".join(map(str, r)) for r in source["page_ranges"]) or "Unverified",
                (source["saved_path"] + "\nSHA256: " + source["sha256"])
                if source["status"] == "saved"
                else "UNAVAILABLE — evidence gap",
            ]
        )
    return rows


def _write_documents(
    root: Path, paragraphs: list[tuple[str, str]], tables: list[TableData], bundle: dict
) -> None:
    """Use one set of accepted strings in each outward format; escape all user text."""
    styles = getSampleStyleSheet()
    doc = Document()
    doc.add_heading("EDF billing complaint — review draft", 0)
    doc.add_paragraph(NOTICE)
    story = [
        Paragraph("EDF billing complaint — review draft", styles["Title"]),
        Paragraph(escape(NOTICE), styles["BodyText"]),
    ]
    html = [
        '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>EDF complaint draft</title><style>body{font:16px system-ui;color:#253849;max-width:1150px;margin:2rem auto;padding:1rem}table{width:100%;border-collapse:collapse;font-size:14px}td,th{padding:10px;border:1px solid #cad5de;white-space:pre-wrap;vertical-align:top;overflow-wrap:anywhere}th{background:#18354a;color:white;text-align:left}h1,h2{color:#18354a}@media print{thead{display:table-header-group}body{margin:0}}</style><h1>EDF billing complaint — review draft</h1><p>'
        + escape(NOTICE)
        + "</p>"
    ]
    for title, text in paragraphs:
        doc.add_heading(title, 1)
        doc.add_paragraph(text)
        story.extend(
            [
                Paragraph(escape(title), styles["Heading2"]),
                Paragraph(escape(text).replace("\n", "<br/>"), styles["BodyText"]),
                Spacer(1, 0.2 * cm),
            ]
        )
        html.append(f"<h2>{escape(title)}</h2><p>{escape(text).replace(chr(10), '<br>')}</p>")
    # Keep the complaint narrative readable as a letter; supporting schedules
    # start on a fresh page rather than splitting a table across its last lines.
    doc.add_page_break()
    story.append(PageBreak())
    for title, note, rows in tables:
        doc.add_heading(title, 1)
        doc.add_paragraph(note)
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        for cells, values in [(table.rows[0].cells, rows[0])] + [
            (table.add_row().cells, row) for row in rows[1:]
        ]:
            for cell, value in zip(cells, values, strict=True):
                cell.text = value
        story.extend(
            [
                Paragraph(escape(title), styles["Heading2"]),
                Paragraph(escape(note), styles["BodyText"]),
            ]
        )
        pdf_rows = [
            [
                Paragraph(escape(str(value)).replace("\n", "<br/>"), styles["BodyText"])
                for value in row
            ]
            for row in rows
        ]
        pdf_table = Table(
            pdf_rows,
            colWidths=[16.5 * cm / len(rows[0])] * len(rows[0]),
            repeatRows=1,
            splitInRow=1,
        )
        pdf_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf1f5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cad5de")),
                ]
            )
        )
        story.extend([pdf_table, Spacer(1, 0.3 * cm)])
        html.append(
            f"<h2>{escape(title)}</h2><p>{escape(note)}</p><table><thead><tr>"
            + "".join(f"<th>{escape(c)}</th>" for c in rows[0])
            + "</tr></thead><tbody>"
            + "".join(
                "<tr>" + "".join(f"<td>{escape(c)}</td>" for c in row) + "</tr>" for row in rows[1:]
            )
            + "</tbody></table>"
        )
    html.append("<h2>Open included originals</h2><ul>")
    for source in bundle["files"]:
        if source["status"] == "saved":
            href = quote(source["saved_path"], safe="/")
            first_page = source["page_ranges"][0][0] if source["page_ranges"] else None
            if first_page:
                href += f"#page={first_page}"
            html.append(
                f'<li><a href="{href}">{escape(source["attachment_name"])}</a> — whole original; see page ranges above</li>'
            )
    html.append("</ul></html>")
    (root / "Complaint_Draft.html").write_text("\n".join(html), encoding="utf-8")
    doc.save(str(root / "Complaint_Draft.docx"))
    SimpleDocTemplate(
        str(root / "Complaint_Draft.pdf"), pagesize=A4, leftMargin=2.25 * cm, rightMargin=2.25 * cm
    ).build(story)
    workbook = Workbook()
    workbook.remove(workbook.active)
    overview = workbook.create_sheet("Read first")
    literal_cell(overview, 1, 1, REFERENCE_MARKER)
    literal_cell(overview, 2, 1, NOTICE)
    overview.column_dimensions["A"].width = 110
    overview["A2"].alignment = Alignment(wrap_text=True)
    overview.row_dimensions[2].height = 110
    for title, text in paragraphs:
        literal_cell(overview, overview.max_row + 2, 1, title)
        cell = literal_cell(overview, overview.max_row + 1, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        overview.row_dimensions[cell.row].height = 100
    names = (
        "Proposal",
        "Consolidated claims",
        "Overlap decisions",
        "Disputed documents",
        "Reading and demand dates",
        "Billing history",
        "Evidence index",
    )
    for name, (_, _, rows) in zip(names, tables, strict=True):
        ws = workbook.create_sheet(name)
        for row_number, values in enumerate(rows, 1):
            for column, value in enumerate(values, 1):
                cell = literal_cell(ws, row_number, column, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if row_number == 1:
                    cell.fill = PatternFill("solid", fgColor="18354A")
                    cell.font = Font(color="FFFFFF", bold=True)
            lines = max(
                sum(max(1, ceil(len(line) / 44)) for line in str(value).split("\n"))
                for value in values
            )
            ws.row_dimensions[row_number].height = (
                38 if row_number == 1 else min(409, max(40, lines * 15 + 8))
            )
        for column in range(1, len(rows[0]) + 1):
            ws.column_dimensions[get_column_letter(column)].width = 48
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.protection.sheet = True
        ws.protection.selectLockedCells = False
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_title_rows = "1:1"
    workbook.save(root / "Claim_Schedule.xlsx")
    workbook.close()


def publish(
    view: ReportView,
    destination: str,
    *,
    source_roots: tuple[str, ...] = (),
    cancelled: Callable[[], bool] = lambda: False,
) -> tuple[dict[str, str], tuple[str, ...]]:
    """Stage the complete pack and archive, publishing neither after failure/cancellation."""
    from edf_bill_fetcher.application.run_service import check_cancelled

    if not requested(view):
        return {}, ()
    target = Path(destination)
    archive = target.with_name("Complaint_pack.zip")
    if target.exists() or archive.exists():
        raise FileExistsError("Complaint pack destinations already exist; choose a new run folder")
    target.parent.mkdir(parents=True, exist_ok=True)
    state = object_value(view.snapshot.prepared.extraction.legacy_state)
    roots = [target.parent, *(Path(r) for r in source_roots)]
    if state.get("source_bundle_path"):
        roots.append(Path(state["source_bundle_path"]))
    with tempfile.TemporaryDirectory(prefix=".complaint-", dir=target.parent) as temporary:
        stage = Path(temporary) / "pack"
        stage.mkdir()
        check_cancelled(cancelled)
        model, ledger, notes = _selection(view)
        bundle, source_notes = _evidence(view, ledger, stage, roots)
        notes.extend(source_notes)
        if not model["ready"]:
            notes.append(
                "Claim reconciliation is incomplete: resolve overlaps, credit allocations and stale or ungrouped challenges before asserting one total."
            )
        notes = list(dict.fromkeys(notes))
        paragraphs = complaint_paragraphs(model)
        if notes:
            paragraphs.insert(0, ("Review gaps — resolve before sending", "\n".join(notes)))
        tables = [
            ("Consolidated claim proposal", model["notice"], total_table(model)),
            (
                "Claim groups and credit allocations",
                "Each economic challenge is entered once; cancellation offsets are not relief received.",
                group_table(model, ledger),
            ),
            (
                "Claim overlap review",
                "Unresolved overlap prevents assertion of one total.",
                overlap_table(model),
            ),
        ]
        tables.extend(
            [
                (
                    "Disputed documents",
                    "Per-document positions are not additive. Consolidated amounts appear above.",
                    schedule_table(ledger),
                ),
                (
                    "Reading and first-demand context",
                    "A printed reading date does not establish when EDF received it. No day-based allocation or automatic entitlement is inferred.",
                    reading_table(ledger),
                ),
                (
                    "Relevant earlier billing history",
                    "Includes linked and overlapping earlier bills/credits, including higher estimates. Whole source PDFs retain all original pages. Excluded records are not separate claims.",
                    history_table(ledger),
                ),
                (
                    "Evidence index",
                    "Paths are relative to this folder. Hashes verify the included bytes, not their accuracy or completeness.",
                    _source_rows(bundle),
                ),
            ]
        )
        check_cancelled(cancelled)
        _write_documents(stage, paragraphs, tables, bundle)
        check_cancelled(cancelled)
        staged_zip = Path(temporary) / "Complaint_pack.zip"
        with ZipFile(staged_zip, "w", compression=ZIP_DEFLATED) as zipped:
            for path in sorted(stage.rglob("*")):
                check_cancelled(cancelled)
                if path.is_file():
                    zipped.write(path, path.relative_to(stage).as_posix())
        check_cancelled(cancelled)
        if target.exists() or archive.exists():
            raise FileExistsError("Complaint pack destinations appeared during generation")
        os.rename(stage, target)
        try:
            os.rename(staged_zip, archive)
        except BaseException:
            shutil.rmtree(target)
            raise
    artifacts = {"complaint_zip": str(archive)}
    artifacts.update(
        {
            "complaint_" + path.suffix.lstrip("."): str(path)
            for path in target.glob("Complaint_Draft.*")
        }
    )
    artifacts["complaint_schedule"] = str(target / "Claim_Schedule.xlsx")
    artifacts["complaint_evidence_manifest"] = str(target / "evidence_manifest.json")
    return artifacts, tuple(notes)
