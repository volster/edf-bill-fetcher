"""Truthful documents when human review excludes every canonical evidence row."""

from html import escape
from pathlib import Path

from edf_bill_fetcher.models.report_view import ReportView


def render_empty_review(view: ReportView, output_path: str, fmt: str) -> str:
    """Publish an explicit empty review without running nonempty-only layouts."""
    title = "EDF evidence review — no included records"
    paragraphs = [
        "No canonical evidence records are included in this reviewed analysis. "
        "No conclusions about the account balance can be drawn from this selection.",
        view.review_notice or "",
        "Analysis date: " + view.snapshot.prepared.extraction.request.as_of.isoformat(),
        "Use Review & Corrections in the updated workbook to reinstate records. "
        "The workbook and result.json retain the review history, original observations "
        "and any SAP dump data. Earlier evidence packs remain unchanged.",
    ]
    if fmt == "html":
        Path(output_path).write_text(
            '<!doctype html><html lang="en"><head><meta charset="utf-8">'
            f"<title>{escape(title)}</title></head><body><h1>{escape(title)}</h1>"
            + "".join(f"<p>{escape(text)}</p>" for text in paragraphs)
            + "</body></html>",
            encoding="utf-8",
        )
    elif fmt == "docx":
        from docx import Document

        document = Document()
        document.add_heading(title, 0)
        for text in paragraphs:
            document.add_paragraph(text)
        document.save(output_path)
    elif fmt == "pdf":
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        styles = getSampleStyleSheet()
        elements = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
        for text in paragraphs:
            elements.extend([Paragraph(escape(text), styles["BodyText"]), Spacer(1, 12)])
        SimpleDocTemplate(output_path).build(elements)
    else:
        raise ValueError("Unsupported empty-review document format")
    return output_path
