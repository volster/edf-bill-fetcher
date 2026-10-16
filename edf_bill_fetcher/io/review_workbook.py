"""Self-contained Excel review exchange with strict, bounded JSON provenance."""

from __future__ import annotations

import base64
import hashlib
import math
import zlib
from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from edf_bill_fetcher.application.record_review import (
    ENTRY_TYPES,
    REVIEW_FIELDS,
    apply_review,
    input_value,
    review_inputs,
    review_rows,
)
from edf_bill_fetcher.domain.json_values import canonical_text, object_value, strict_loads
from edf_bill_fetcher.helpers.excel_utils import literal_cell
from edf_bill_fetcher.io.persistence.result_codec import decode_result, encode_result
from edf_bill_fetcher.io.writers.workbook_presentation import (
    polish_reference_sheets,
    write_workbook_dashboard,
)
from edf_bill_fetcher.models.report_view import ReportView

REVIEW_SHEET = "Review & Corrections"
PAYLOAD_SHEET = "_EDF_Review_Data"
REVIEW_MARKER = "EDF_REVIEW_WORKBOOK_1"
HEADER_ROW = 4
MAX_BYTES = 128 * 1024 * 1024
MAX_ROWS = 100_000
READ_ONLY = ("Source file", "Source", "OCR review", "Review Row ID")
HEADERS = ("Include", "Reason for change", *REVIEW_FIELDS, *READ_ONLY)
NAVY, TEAL, PALE, GREY = "18354A", "007F82", "EAF7F5", "637585"


def _display(field: str, record: dict) -> Any:
    value = record.get(field)
    kind = REVIEW_FIELDS[field]
    if kind == "bool":
        return "Yes" if value is True else "No"
    if kind == "date" and value:
        try:
            return datetime.strptime(input_value(field, value), "%d/%m/%Y")
        except ValueError:
            return str(value)
    if kind == "text" and value is not None:
        return ILLEGAL_CHARACTERS_RE.sub("", str(value))[:32000]
    return value


def _reference(item: dict) -> tuple:
    record = item["record"]
    return (
        record.get("Attachment Name"),
        record.get("Source"),
        "Check original PDF" if record.get("OCR Review Required") else "",
        item["id"],
    )


def _heading(ws: Worksheet, row: int, text: str, end: int = 10) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end)
    cell = literal_cell(ws, row, 1, text)
    cell.font = Font(name="Calibri", size=18 if row == 1 else 11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY if row == 1 else TEAL)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 38


def write_review_sheets(
    wb: Workbook, view: ReportView, saved_links: dict[str, str], workbook_path: str
) -> None:
    """Add an editable table, navigation, audit trail and a verified baseline snapshot."""
    snapshot = view.snapshot
    wb.properties.identifier = REVIEW_MARKER
    items = review_rows(snapshot)
    if len(items) > MAX_ROWS:
        raise ValueError("Too many records for the Excel review exchange")
    hidden = wb.create_sheet(PAYLOAD_SHEET)
    hidden.sheet_state = "veryHidden"

    start = wb.create_sheet("Workbook guide")
    start.sheet_view.showGridLines = False
    start.sheet_properties.tabColor = TEAL
    for col in range(1, 11):
        start.column_dimensions[get_column_letter(col)].width = 14
    _heading(start, 1, "EDF evidence • review workspace")
    _heading(start, 3, "1   Check sources     2   Correct or exclude     3   Regenerate your pack")
    state = object_value(snapshot.prepared.extraction.legacy_state)
    notices = [
        (
            5,
            f"Analysis as of {snapshot.prepared.extraction.request.as_of.isoformat()}  •  {sum(item['included'] for item in items)} included records  •  {sum(not item['included'] for item in items)} excluded records",
        ),
        (
            7,
            "Open Review & Corrections. Pale green cells are inputs. Compare OCR and questionable amounts with the original file before changing them.",
        ),
        (
            9,
            "Choose Include = No to remove a record from the next analysis. Give a reason for every correction, exclusion or reinstatement. Keep all rows and row IDs; sorting and filtering whole table rows is supported.",
        ),
        (
            11,
            "Balance is the amount owed at a point in time. Period charge is the bill's cost. Transaction is the signed movement: charges/refunds positive, payments/credits negative. Changing a charge also updates its matching transaction unless you edit that transaction separately.",
        ),
        (
            14,
            "Save a copy of this workbook beside the evidence_files folder. In Generate reports choose the amended .xlsx and run. The app recalculates all analysis and creates a fresh workbook, reports and bundle in a new folder, using the original analysis date and settings.",
        ),
        (
            17,
            "Report, summary, SAP and filtered tabs are reference outputs, not correction inputs. Excel does not recalculate the app's analysis. They show the previous result until you regenerate. Correct canonical evidence in Review & Corrections; new manual records and SAP dump editing are not supported here.",
        ),
        (
            20,
            "Exclusions remove records and files used only by those records from the new bundle. Shared PDFs remain whole: this does not redact pages. Original extraction data and review history remain in result.json. Send a separately redacted pack if confidentiality requires it.",
        ),
        (
            23,
            "Source files are not embedded in Excel. Keep evidence_files with the workbook for portable links. Missing or changed sources are reported; the app never substitutes a different same-named file. Worksheet protection guides editing; it is not encryption or a digital signature.",
        ),
    ]
    for row, text in notices:
        start.merge_cells(
            start_row=row, start_column=1, end_row=row + (1 if row >= 11 else 0), end_column=10
        )
        cell = literal_cell(start, row, 1, text)
        cell.font = Font(name="Calibri", size=11, color=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        start.row_dimensions[row].height = 32 if row < 11 else 42
    for row, title in enumerate(
        (
            REVIEW_SHEET,
            "Review history",
            "EDF Evidence Report",
            "Annual Summary",
            "Data Quality Report",
            "Parse Errors",
        ),
        27,
    ):
        if title in wb.sheetnames or title in {REVIEW_SHEET, "Review history"}:
            cell = literal_cell(start, row, 1, "Open " + title)
            cell.hyperlink = f"#'{title}'!A1"
            cell.font = Font(color=TEAL, underline="single", bold=True)
    start.freeze_panes = "A5"
    start.print_options.horizontalCentered = True
    start.print_area = "A1:J33"
    start.sheet_properties.pageSetUpPr.fitToPage = True
    start.page_setup.orientation = "portrait"
    start.page_setup.paperSize = start.PAPERSIZE_A4
    start.page_setup.fitToWidth = 1
    start.page_setup.fitToHeight = 1

    ws = wb.create_sheet(REVIEW_SHEET, 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.tabColor = TEAL
    _heading(ws, 1, "Review & Corrections • edit the pale green cells", len(HEADERS))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    literal_cell(
        ws,
        2,
        1,
        "Give a reason for changes. Include = No excludes a row; do not delete it. Save and regenerate in the app to update all reports. Dates: DD/MM/YYYY. Blanks mean unknown, not zero.",
    )
    ws.cell(2, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32
    for column, title in enumerate(HEADERS, 1):
        cell = literal_cell(ws, HEADER_ROW, column, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL if title not in READ_ONLY else GREY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        width = 32 if title in {"Reason for change", "Details", "Source file"} else 18
        if title == "Include":
            width = 11
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[HEADER_ROW].height = 34
    for row_number, item in enumerate(items, HEADER_ROW + 1):
        record = item["record"]
        values = (
            "Yes" if item["included"] else "No",
            "",
            *(_display(field, record) for field in REVIEW_FIELDS),
            *_reference(item),
        )
        for column, value in enumerate(values, 1):
            cell = literal_cell(ws, row_number, column, value)
            editable = column <= len(HEADERS) - len(READ_ONLY)
            cell.protection = Protection(locked=not editable)
            cell.fill = PatternFill("solid", fgColor=PALE if editable else "F2F4F6")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if 3 <= column <= len(REVIEW_FIELDS) + 2:
                kind = REVIEW_FIELDS[HEADERS[column - 1]]
                cell.number_format = (
                    "dd/mm/yyyy"
                    if kind == "date"
                    else '#,##0.00;[Red]-#,##0.00;"—"'
                    if kind == "money"
                    else "0.######"
                    if kind == "number"
                    else "@"
                    if kind == "text"
                    else "General"
                )
        source_column = HEADERS.index("Source file") + 1
        source = record.get("Source ID") or record.get("Attachment Name")
        link = saved_links.get(source) or saved_links.get(record.get("Attachment Name"))
        if link:
            if ws.cell(row_number, source_column).value:
                ws.cell(row_number, source_column).hyperlink = link
                ws.cell(row_number, source_column).font = Font(color=TEAL, underline="single")
            invoice_cell = ws.cell(row_number, HEADERS.index("Invoice #") + 1)
            if invoice_cell.value:
                invoice_cell.hyperlink = link
                invoice_cell.font = Font(color=TEAL, underline="single")
                invoice_cell.comment = Comment(
                    "Click to open the original source. Press F2 to correct the invoice reference. The source file itself is never edited.",
                    "EDF Collector",
                )
        ws.row_dimensions[row_number].height = 32
    last = HEADER_ROW + len(items)
    if items:
        table = Table(
            displayName="EDFReviewInputs",
            ref=f"A{HEADER_ROW}:{get_column_letter(len(HEADERS))}{last}",
        )
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(table)
        for name, formula in (
            ("Include", '"Yes,No"'),
            ("Entry Type", '"' + ",".join(ENTRY_TYPES) + '"'),
            ("Cancel/Rebill Admitted", '"Yes,No"'),
            ("Amount Kind", '"Balance,Transaction,Unknown"'),
            ("Fuel", '"Electricity,Gas,Dual,Unknown"'),
        ):
            col = get_column_letter(HEADERS.index(name) + 1)
            validation = DataValidation(type="list", formula1=formula, allow_blank=name == "Fuel")
            validation.errorTitle = "Choose a listed value"
            validation.error = "Use the dropdown; regeneration validates inputs again."
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            ws.add_data_validation(validation)
            validation.add(f"{col}{HEADER_ROW + 1}:{col}{last}")
        for field, kind in REVIEW_FIELDS.items():
            if kind not in {"money", "number"}:
                continue
            col = get_column_letter(HEADERS.index(field) + 1)
            validation = DataValidation(
                type="decimal",
                operator="between",
                formula1="-1000000000000",
                formula2="1000000000000",
                allow_blank=True,
            )
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            validation.error = (
                "Enter a number without a currency symbol, or leave blank for unknown."
            )
            ws.add_data_validation(validation)
            validation.add(f"{col}{HEADER_ROW + 1}:{col}{last}")
        ws.conditional_formatting.add(
            f"A5:{get_column_letter(len(HEADERS))}{last}",
            FormulaRule(
                formula=['$A5="No"'],
                fill=PatternFill("solid", fgColor="FBE7E5"),
                font=Font(color="874740"),
            ),
        )
    ws.freeze_panes = "F5"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    # Excel cannot sort a protected range containing locked IDs. The guide
    # allows unprotecting (no password) before sorting the complete table.
    ws.cell(2, 1).comment = Comment(
        "To sort in Excel, Review > Unprotect Sheet (no password), then sort the whole table. Do not edit grey source/identity columns.",
        "EDF Collector",
    )
    ws.column_dimensions[get_column_letter(len(HEADERS))].hidden = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"

    # Keep less frequent fields available behind column outlines. The core
    # financial inputs and source link remain visible without a 27-column scan.
    for title in (
        "Energy Unit Rate (p/kWh)",
        "Standing Chg (p/day)",
        "Reading",
        "Register Reading (kWh)",
        "Meter",
        "Register",
        "Tariff",
        "Amount Kind",
        "Cancel/Rebill Admitted",
        "Details",
        "Source",
    ):
        dimension = ws.column_dimensions[get_column_letter(HEADERS.index(title) + 1)]
        dimension.hidden = True
        dimension.outlineLevel = 1
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.cell(
        2, 1
    ).comment.text += " Expand the column outlines for rates, meter/register detail and notes. If Excel disables the outline buttons, unprotect the sheet (no password) or use Unhide Columns."

    history = wb.create_sheet("Review history", 1)
    history.sheet_view.showGridLines = False
    history.sheet_properties.tabColor = GREY
    columns = (
        "Reviewed at (UTC)",
        "Base fingerprint",
        "Row ID",
        "Invoice",
        "Field",
        "Before",
        "After",
        "Reason",
    )
    for column, title in enumerate(columns, 1):
        literal_cell(history, 1, column, title).font = Font(bold=True, color="FFFFFF")
        history.cell(1, column).fill = PatternFill("solid", fgColor=NAVY)
        history.column_dimensions[get_column_letter(column)].width = 25 if column < 5 else 35
    for row, change in enumerate(state.get("review_history", []), 2):
        for column, key in enumerate(
            ("at", "base_fingerprint", "row_id", "invoice", "field", "before", "after", "reason"), 1
        ):
            literal_cell(history, row, column, change.get(key))
    history.freeze_panes = "E2"
    history.auto_filter.ref = history.dimensions
    write_workbook_dashboard(wb, view)
    polish_reference_sheets(wb)
    # Calculated sheets are reference outputs. Protection prevents accidental
    # edits, while import also checks their values against a stored digest.
    for sheet in wb:
        if sheet.title not in {REVIEW_SHEET, PAYLOAD_SHEET}:
            sheet.protection.sheet = True
            sheet.protection.autoFilter = False
            sheet.sheet_view.showGridLines = False
    # Store the reference values with the snapshot. Comparing numeric cells
    # allows Excel's serialization rounding, rather than hashing binary floats.
    source_root_relative = None
    source_root = state.get("source_bundle_path")
    if isinstance(source_root, str) and source_root:
        parent = Path(workbook_path).resolve().parent
        root = Path(source_root).resolve()
        if root.is_relative_to(parent):
            source_root_relative = root.relative_to(parent).as_posix()
    payload = {
        "result": encode_result(snapshot, snapshot.prepared.extraction.request.run_id + "-excel"),
        "references": _reference_values(wb),
        "source_root_relative": source_root_relative,
    }
    data = canonical_text(payload).encode("utf-8")
    if len(data) > MAX_BYTES:
        raise ValueError("Analysis exceeds the 128 MiB Excel review payload limit")
    encoded = base64.b64encode(zlib.compress(data)).decode("ascii")
    literal_cell(hidden, 1, 1, REVIEW_MARKER)
    literal_cell(hidden, 2, 1, hashlib.sha256(data).hexdigest())
    literal_cell(hidden, 3, 1, len(data))
    for index in range(0, len(encoded), 30000):
        literal_cell(hidden, index // 30000 + 4, 1, encoded[index : index + 30000])
    wb.active = 0


def _reference_values(wb: Workbook) -> list[list]:
    result = []
    for ws in sorted(wb, key=lambda sheet: sheet.title):
        if ws.title in {REVIEW_SHEET, PAYLOAD_SHEET}:
            continue
        # Sparse iteration avoids materializing large formatted empty ranges.
        for (row, col), cell in sorted(ws._cells.items()):
            value = cell.value
            if value is None or value == "":
                continue
            if isinstance(value, datetime):
                value = value.isoformat()
            elif cell.data_type == "n":
                value = float(value)
            result.append([ws.title, row, col, cell.data_type, value])
    return result


def _references_match(wb: Workbook, expected: object) -> bool:
    if not isinstance(expected, list):
        return False
    actual = _reference_values(wb)
    if len(expected) != len(actual):
        return False
    for old, current in zip(expected, actual, strict=True):
        if not isinstance(old, list) or len(old) != 5 or old[:4] != current[:4]:
            return False
        if old[3] == "n":
            if type(old[4]) not in (int, float) or not math.isclose(
                old[4], current[4], rel_tol=2e-15, abs_tol=1e-9
            ):
                return False
        elif old[4] != current[4]:
            return False
    return True


def load_review_workbook(path: str) -> tuple[list[dict], dict, str] | None:
    """Return preserved originals and validated review state; None means a legacy file."""
    # Bound ZIP expansion before openpyxl loads the workbook into memory.
    with ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_BYTES * 4:
            raise ValueError("Workbook is too large to import safely")
    wb = load_workbook(path, data_only=False, keep_links=False)
    try:
        present = {REVIEW_SHEET, PAYLOAD_SHEET}.intersection(wb.sheetnames)
        if not present:
            if wb.properties.identifier == REVIEW_MARKER:
                raise ValueError(
                    "Review workbook sheets were removed; restore the original review workbook"
                )
            return None
        if len(present) != 2:
            raise ValueError(
                "Review workbook is incomplete; restore its corrections and hidden provenance sheets"
            )
        hidden = wb[PAYLOAD_SHEET]
        if hidden.cell(1, 1).value != REVIEW_MARKER or hidden.max_row > (
            MAX_BYTES * 2 // 30000 + 4
        ):
            raise ValueError("Invalid review workbook metadata")
        length = hidden.cell(3, 1).value
        if type(length) is not int or not 1 <= length <= MAX_BYTES:
            raise ValueError("Invalid review payload size")
        try:
            encoded = "".join(hidden.cell(row, 1).value for row in range(4, hidden.max_row + 1))
            compressed = base64.b64decode(encoded, validate=True)
            inflater = zlib.decompressobj()
            raw = inflater.decompress(compressed, length + 1)
            if (
                len(raw) != length
                or not inflater.eof
                or inflater.unused_data
                or inflater.unconsumed_tail
            ):
                raise ValueError("Invalid review payload length")
            if hashlib.sha256(raw).hexdigest() != hidden.cell(2, 1).value:
                raise ValueError("Review payload checksum does not match")
            payload = strict_loads(raw.decode("utf-8"))
            if not isinstance(payload, dict) or set(payload) != {
                "result",
                "references",
                "source_root_relative",
            }:
                raise ValueError("Invalid review payload fields")
            relative_root = payload["source_root_relative"]
            if relative_root is not None and (
                not isinstance(relative_root, str)
                or Path(relative_root).is_absolute()
                or ".." in Path(relative_root).parts
                or "\\" in relative_root
                or ":" in relative_root
            ):
                raise ValueError("Invalid relative source bundle location")
            snapshot = decode_result(payload["result"])
        except (TypeError, UnicodeError, zlib.error, ValueError) as error:
            raise ValueError(f"Cannot read preserved review analysis: {error}") from error
        if not _references_match(wb, payload["references"]):
            raise ValueError(
                "A reference report tab was changed. Restore it and make corrections in Review & Corrections; only that table drives regeneration."
            )
        ws = wb[REVIEW_SHEET]
        if (
            ws.max_row > MAX_ROWS + HEADER_ROW
            or tuple(ws.cell(HEADER_ROW, col).value for col in range(1, len(HEADERS) + 1))
            != HEADERS
        ):
            raise ValueError("Review table headers were changed or row limit exceeded")
        baseline = review_rows(snapshot)
        by_id = {item["id"]: item for item in baseline}
        decisions = []
        for row_number in range(HEADER_ROW + 1, ws.max_row + 1):
            cells = [ws.cell(row_number, col) for col in range(1, len(HEADERS) + 1)]
            if all(cell.value is None for cell in cells):
                continue
            try:
                if any(cell.data_type in {"f", "e"} for cell in cells):
                    raise ValueError(
                        "Formulas and Excel errors are not valid review inputs; paste values"
                    )
                identity = cells[-1].value
                if not isinstance(identity, str) or identity not in by_id:
                    raise ValueError(
                        "Unknown row ID; do not add rows or edit grey identity columns"
                    )
                original = by_id[identity]
                if tuple(cell.value or None for cell in cells[-len(READ_ONLY) :]) != tuple(
                    value or None for value in _reference(original)
                ):
                    raise ValueError(
                        "Source/identity cells were changed; only pale green inputs are editable"
                    )
                inclusion = cells[0].value
                if inclusion not in {"Yes", "No"}:
                    raise ValueError("Include must be Yes or No")
                reason = cells[1].value or ""
                changes = {}
                for index, field in enumerate(REVIEW_FIELDS, 2):
                    value = cells[index].value
                    old = _display(field, original["record"])
                    if REVIEW_FIELDS[field] in {"money", "number"}:
                        if type(value) is bool:
                            raise ValueError(
                                f"{field}: enter a finite number, without currency symbols"
                            )
                        if (
                            type(value) in (int, float)
                            and type(old) in (int, float)
                            and math.isclose(value, old, rel_tol=2e-15, abs_tol=1e-9)
                        ):
                            continue
                    if value == old or (value is None and old == ""):
                        continue
                    normalized = input_value(field, value)
                    try:
                        if normalized == input_value(field, old):
                            continue
                    except ValueError:
                        pass
                    changes[field] = normalized
                decision = {
                    "row_id": identity,
                    "include": inclusion == "Yes",
                    "changes": changes,
                    "reason": reason,
                }
                decisions.append(decision)
                # Validate each decision here so error messages identify its Excel row.
                apply_review(
                    {
                        "version": 1,
                        "base_fingerprint": snapshot.input_fingerprint,
                        "rows": [original],
                        "decisions": [decision],
                    }
                )
            except (TypeError, ValueError) as error:
                raise ValueError(f"{REVIEW_SHEET}, row {row_number}: {error}") from error
        return review_inputs(
            snapshot,
            decisions,
            input_path=path,
            relative_root=relative_root,
        )
    finally:
        wb.close()


def read_workbook(path: str) -> tuple[list[dict], dict, str | None]:
    """Prefer the lossless review exchange; explicitly label older imports as lossy."""
    reviewed = load_review_workbook(path)
    if reviewed is not None:
        return reviewed
    import pandas as pd

    from edf_bill_fetcher.models.snapshot_adapter import to_snapshot_json

    with pd.ExcelFile(path) as workbook:
        records = pd.read_excel(workbook, sheet_name="EDF Evidence Report").to_dict("records")
        metadata = {}
        if "Provenance" in workbook.sheet_names:
            frame = pd.read_excel(workbook, sheet_name="Provenance", header=None)
            metadata = {
                str(row[0]): to_snapshot_json(row[1])
                for row in frame.itertuples(index=False, name=None)
                if len(row) > 1 and pd.notna(row[0])
            }
        filtered = (
            pd.read_excel(workbook, sheet_name="Filtered (Below Min)").to_dict("records")
            if "Filtered (Below Min)" in workbook.sheet_names
            else []
        )
    return (
        records,
        {
            "records": records,
            "filtered_records": filtered,
            "pdf_count": 0,
            "email_count": 0,
            "source_manifest": {},
            "imported_workbook": path,
            "imported_workbook_metadata": metadata,
            "provenance": "lossy_workbook_import",
            "error_log": [
                "Workbook import: original observations, source files and analytical decisions are unavailable."
            ],
        },
        str(metadata["as_of"]) if metadata.get("as_of") else None,
    )
