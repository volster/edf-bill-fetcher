"""Editable dispute schedules bound to the workbook's preserved canonical evidence."""

from __future__ import annotations

from datetime import date, datetime

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, PatternFill, Protection
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from edf_bill_fetcher.application.record_review import corrected_record, review_rows
from edf_bill_fetcher.domain.analysis import AnalysisSnapshot
from edf_bill_fetcher.helpers.excel_utils import literal_cell
from edf_bill_fetcher.io.writers.rebilling_ledger import reference_sheet, source_link
from edf_bill_fetcher.models.output_tables import ledger_label
from edf_bill_fetcher.models.rebilling_tables import position_state
from edf_bill_fetcher.models.report_view import RenderSelection, ReportView, build_report_view
from edf_bill_fetcher.processors.rebilling_ledger import (
    SCHEDULE_FIELDS,
    STATUSES,
    schedule_digest,
    validate_schedule,
)
from edf_bill_fetcher.processors.revisions import BILLS

SHEET = "Dispute schedule"
EDITABLE = (
    "Status",
    "Challenged amount (£)",
    "Proposed reduction (£)",
    "First demand date",
    "Basis / scope",
    "EDF must substantiate",
)
HEADERS = (
    *EDITABLE,
    "Invoice",
    "Invoice date",
    "Account / ledger",
    "Included",
    "Period from",
    "Period to",
    "Period charge (£)",
    "Gross electricity, printed (£)",
    "Source file",
    "Pages",
    "Evidence status",
    "Row ID",
)


def _values(doc: dict) -> tuple:
    p, row = doc["position"] or {}, doc["record"]
    return (
        p.get("status", "unreviewed"),
        float(p["challenged_amount"]) if p.get("challenged_amount") is not None else None,
        float(p["proposed_reduction"]) if p.get("proposed_reduction") is not None else None,
        p.get("first_demand_date"),
        p.get("basis", ""),
        p.get("edf_request", ""),
        row.get("Invoice #"),
        doc["invoice_date"],
        ledger_label(row),
        "Yes" if doc["included"] else "No",
        doc["period_from"],
        doc["period_to"],
        row.get("Period Charge (£)"),
        row.get("Gross Electricity Charges (£)"),
        row.get("Attachment Name"),
        "–".join(str(row[k]) for k in ("Page From", "Page To") if row.get(k) is not None),
        position_state(doc),
        doc["row_id"],
    )


def write_schedule(wb: Workbook, view: ReportView, links: dict[str, str]) -> None:
    """Put editable positions beside protected original facts and stable row identities."""
    model = view.rebilling_ledger
    if model is None:
        return
    ws = reference_sheet(
        wb,
        SHEET,
        model["notice"]
        + " Edit pale green inputs, give the scope and evidence in Basis, then regenerate. First demand is a date you verify for the additional amount; leaving it blank does not use the invoice date. Re-save changed positions in the app to acknowledge changed evidence.",
        HEADERS,
    )
    docs = [doc for doc in model["documents"] if doc["record"].get("Entry Type") in BILLS]
    for n, doc in enumerate(docs, 5):
        for c, value in enumerate(_values(doc), 1):
            cell = literal_cell(ws, n, c, value)
            cell.protection = Protection(locked=c > len(EDITABLE))
            cell.fill = PatternFill("solid", fgColor="EAF7F5" if c <= len(EDITABLE) else "F2F4F6")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c in (2, 3, 13, 14):
                cell.number_format = "£#,##0.00;[Red]-£#,##0.00"
        source_link(ws.cell(n, 15), doc["record"], links)
        ws.row_dimensions[n].height = 65
    if docs:
        table = Table(displayName="EDFDisputePositions", ref=f"A4:R{4 + len(docs)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(table)
        status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"')
        status.errorStyle, status.showErrorMessage = "stop", True
        ws.add_data_validation(status)
        status.add(f"A5:A{4 + len(docs)}")
        for column in ("B", "C"):
            validation = DataValidation(
                type="decimal",
                operator="between",
                formula1="0",
                formula2="1000000000000",
                allow_blank=True,
            )
            validation.errorStyle, validation.showErrorMessage = "stop", True
            ws.add_data_validation(validation)
            validation.add(f"{column}5:{column}{4 + len(docs)}")
    ws.freeze_panes = "G5"
    ws.column_dimensions["R"].hidden = True
    for column in ("E", "F"):
        ws.column_dimensions[column].width = 65
    for n in range(1, 7):
        ws.cell(4, n).comment = Comment(
            "Editable dispute position. Keep all rows and IDs. Unprotect the sheet (no password) to sort the whole table. Regenerate to update reports; reference figures do not update in Excel.",
            "EDF Collector",
        )


def read_schedule(
    wb: Workbook, snapshot: AnalysisSnapshot, record_decisions: list[dict]
) -> list[dict] | None:
    """Reject invented/removed rows, altered source facts and formulas before importing."""
    model = build_report_view(snapshot, RenderSelection()).rebilling_ledger
    if model is None:
        if SHEET in wb.sheetnames:
            raise ValueError("This analysis has no preserved dispute schedule baseline")
        return None
    if SHEET not in wb.sheetnames:
        raise ValueError("Dispute schedule was removed; restore it from the original workbook")
    ws = wb[SHEET]
    if (
        ws.max_row > 100004
        or tuple(ws.cell(4, col).value for col in range(1, len(HEADERS) + 1)) != HEADERS
    ):
        raise ValueError("Dispute schedule headers were changed or row limit exceeded")
    docs = {
        doc["row_id"]: doc for doc in model["documents"] if doc["record"].get("Entry Type") in BILLS
    }
    indexed = {decision["row_id"]: decision for decision in record_decisions}
    baseline = review_rows(snapshot)
    if len(indexed) != len(record_decisions) or set(indexed) != {item["id"] for item in baseline}:
        raise ValueError("Correction rows are missing or duplicated; restore every original row")
    items = [
        {
            "id": item["id"],
            "included": indexed[item["id"]]["include"],
            "record": corrected_record(item["record"], indexed[item["id"]]["changes"]),
        }
        for item in baseline
    ]
    from edf_bill_fetcher.domain.json_values import object_value

    revisions = object_value(snapshot.prepared.extraction.legacy_state).get(
        "bill_revision_decisions", []
    )
    result, seen = [], set()
    for n in range(5, ws.max_row + 1):
        cells = [ws.cell(n, c) for c in range(1, len(HEADERS) + 1)]
        if all(cell.value is None for cell in cells):
            continue
        try:
            if any(cell.data_type in {"f", "e"} for cell in cells):
                raise ValueError(
                    "Formulas and Excel errors are not valid schedule inputs; paste values"
                )
            key = cells[-1].value
            if not isinstance(key, str) or key not in docs or key in seen:
                raise ValueError("Unknown or duplicate schedule row ID; keep all original rows")
            seen.add(key)
            doc = docs[key]
            if tuple(
                None if cell.value in (None, "") else cell.value for cell in cells[6:]
            ) != tuple(None if value in (None, "") else value for value in _values(doc)[6:]):
                raise ValueError("Grey dispute source/identity cells were changed")
            values = [cell.value for cell in cells[:6]]
            if isinstance(values[3], (date, datetime)):
                if isinstance(values[3], datetime) and values[3].time().isoformat() != "00:00:00":
                    raise ValueError("First demand date must not include a time")
                values[3] = values[3].strftime("%Y-%m-%d")
            for index in (4, 5):
                values[index] = "" if values[index] is None else values[index]
            previous = doc["position"]
            if values[0] == "unreviewed" and not any(
                value not in (None, "") for value in values[1:]
            ):
                continue
            decision = {
                "row_id": key,
                "source_digest": previous["source_digest"]
                if previous
                else schedule_digest(items, revisions, key),
                **dict(zip(SCHEDULE_FIELDS, values, strict=True)),
            }
            normalized = validate_schedule(items, [decision])[0]
            if previous is not None and normalized != previous:
                normalized["source_digest"] = schedule_digest(items, revisions, key)
            if previous is not None or any(
                normalized[field] not in (None, "", "unreviewed") for field in SCHEDULE_FIELDS
            ):
                result.append(normalized)
        except (ValueError, TypeError) as error:
            raise ValueError(f"{SHEET}, row {n}: {error}") from error
    if seen != docs.keys():
        raise ValueError("Dispute schedule rows were removed; restore every original row")
    return result
