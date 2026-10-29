"""Navigable, filterable rebilling history and component worksheets."""

from __future__ import annotations

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from edf_bill_fetcher.helpers.excel_utils import literal_cell, text
from edf_bill_fetcher.io.writers.sheet_layout import (
    write_banner,
    write_header_row,
    write_merged_text,
)
from edf_bill_fetcher.models.output_tables import ledger_label
from edf_bill_fetcher.models.rebilling_tables import comparison_table, position_state
from edf_bill_fetcher.models.report_view import ReportView


def source_link(cell: Cell, row: dict, links: dict[str, str]) -> None:
    """Link to this source occurrence and the first known original PDF page."""
    identity = row.get("Source ID")
    if identity is not None and not isinstance(identity, str):
        return
    path = links.get(identity) if identity else links.get(str(row.get("Attachment Name", "")))
    if path:
        page = row.get("Page From")
        cell.hyperlink = path + (
            f"#page={int(page)}"
            if isinstance(page, (int, float))
            and not isinstance(page, bool)
            and 0 < page < 1_000_000
            and page == int(page)
            else ""
        )
        cell.font = Font(color="007F82", underline="single")


def reference_sheet(wb: Workbook, title: str, note: str, headers: tuple[str, ...]) -> Worksheet:
    """Keep multi-line reference tables consistent and printable."""
    ws = wb.create_sheet(title)
    write_banner(ws, title.upper(), len(headers))
    write_merged_text(ws, 2, note, len(headers), height=65)
    write_header_row(ws, 4, list(headers), height=36)
    ws.freeze_panes = "D5"
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(index)].width = (
            42
            if any(word in header for word in ("detail", "Source", "Relations", "ledger"))
            else 23
        )
    return ws


def write_ledger_sheets(wb: Workbook, view: ReportView, links: dict[str, str]) -> None:
    """Render stored document and component facts, preserving credits exactly once."""
    model = view.rebilling_ledger
    if model is None:
        return
    headers = (
        "Invoice",
        "Invoice date",
        "Account / ledger",
        "Included",
        "Entry type",
        "Consumption from",
        "Consumption to",
        "Reading date",
        "Closing read type",
        "Register reading (kWh)",
        "Usage (kWh)",
        "Period charge (£)",
        "Transaction (£)",
        "Gross electricity, printed (£)",
        "Observed balance (£)",
        "Relations / decisions",
        "Dispute status",
        "Source file",
        "Pages",
        "Row ID",
    )
    ws = reference_sheet(wb, "Rebilling ledger", model["notice"], headers)
    indexed = {doc["row_id"]: doc for doc in model["documents"]}
    for n, doc in enumerate(model["documents"], 5):
        row = doc["record"]
        relations = []
        for direction, edges, endpoint in (
            ("From", doc["incoming"], "older"),
            ("To", doc["outgoing"], "newer"),
        ):
            for link in edges:
                other = indexed[link[endpoint]]["record"]
                relations.append(
                    f"{direction} {other.get('Invoice #')}: {link['kind']} ({link['status']})"
                )
        values = (
            row.get("Invoice #"),
            doc["invoice_date"],
            ledger_label(row),
            "Yes" if doc["included"] else "No",
            row.get("Entry Type"),
            doc["period_from"],
            doc["period_to"],
            doc["reading_date"],
            row.get("Closing Reading Type") or row.get("Reading"),
            row.get("Register Reading (kWh)"),
            row.get("Units (kWh)"),
            row.get("Period Charge (£)"),
            row.get("Transaction Amount (£)"),
            row.get("Gross Electricity Charges (£)"),
            row.get("Balance (£)"),
            "\n".join(relations),
            position_state(doc),
            row.get("Attachment Name"),
            "–".join(str(row[k]) for k in ("Page From", "Page To") if row.get(k) is not None),
            doc["row_id"],
        )
        for c, value in enumerate(values, 1):
            cell = literal_cell(ws, n, c, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if 12 <= c <= 15:
                cell.number_format = "£#,##0.00;[Red]-£#,##0.00"
        source_link(ws.cell(n, 18), row, links)
        ws.row_dimensions[n].height = 65
    ws.auto_filter.ref = f"A4:T{max(4, ws.max_row)}"
    ws.column_dimensions["T"].hidden = True
    part_headers = (
        "Invoice",
        "Invoice date",
        "Account / ledger",
        "Included",
        "Component",
        "Period from",
        "Period to",
        "Units / days",
        "Rate (p or %)",
        "Component amount (£)",
        "Source detail",
        "Source file",
        "Pages",
        "Row ID",
    )
    parts = reference_sheet(
        wb,
        "Rebilling components",
        "Observed components only. Energy and standing charges exclude VAT. VAT, discounts and cancellations remain separate; these are not additional account transactions. Components may be incomplete and must not be added to the period charge again. Page ranges identify the original document, not a verified exact component page.",
        part_headers,
    )
    n = 5
    for doc in model["documents"]:
        row = doc["record"]
        for component in doc["components"]:
            component_values = (
                row.get("Invoice #"),
                doc["invoice_date"],
                ledger_label(row),
                "Yes" if doc["included"] else "No",
                component["kind"],
                component["period_from"],
                component["period_to"],
                component["units"],
                component["rate"],
                float(component["amount"]) if component["amount"] is not None else None,
                component["detail"],
                row.get("Attachment Name"),
                "–".join(str(row[k]) for k in ("Page From", "Page To") if row.get(k) is not None),
                doc["row_id"],
            )
            for c, value in enumerate(component_values, 1):
                text(parts, n, c, value, wrap=True) if c not in (8, 9, 10) else literal_cell(
                    parts, n, c, value
                )
            parts.cell(n, 10).number_format = "£#,##0.00;[Red]-£#,##0.00"
            source_link(parts.cell(n, 12), row, links)
            n += 1
    parts.auto_filter.ref = f"A4:N{max(4, parts.max_row)}"
    parts.column_dimensions["N"].hidden = True

    comparison_headers = (
        "Earlier → later / decision",
        "Comparison scope",
        "Component changes",
        "Method / limits",
        "Earlier source",
        "Later source",
        "Earlier row ID",
        "Later row ID",
    )
    comparisons = reference_sheet(
        wb,
        "Rebilling comparisons",
        "Every stored comparison, including partial overlaps. Differences explain observations; they do not create refunds. Original extracted components remain distinct from corrected scalar fields. Verify discrepancies against the original source.",
        comparison_headers,
    )
    for n, (link, comparison_values) in enumerate(
        zip(model["comparisons"], comparison_table(model)[1:], strict=True), 5
    ):
        a, b = indexed[link["older"]]["record"], indexed[link["newer"]]["record"]
        for c, value in enumerate(
            [
                *comparison_values,
                a.get("Attachment Name"),
                b.get("Attachment Name"),
                link["older"],
                link["newer"],
            ],
            1,
        ):
            text(comparisons, n, c, value, wrap=True)
        source_link(comparisons.cell(n, 5), a, links)
        source_link(comparisons.cell(n, 6), b, links)
        comparisons.row_dimensions[n].height = 100
    for column in ("A", "B", "C", "D"):
        comparisons.column_dimensions[column].width = 65
    comparisons.column_dimensions.group("G", "H", hidden=True)
    comparisons.auto_filter.ref = f"A4:H{max(4, comparisons.max_row)}"
