"""Workbook orientation and navigation without changing evidence cell addresses."""

from __future__ import annotations

from math import ceil

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook

from edf_bill_fetcher.helpers.excel_utils import literal_cell
from edf_bill_fetcher.models.output_overview import build_output_overview, report_account_label
from edf_bill_fetcher.models.report_view import ReportView

NAVY, TEAL, INK, GREY = "18354A", "007F82", "253849", "526576"


def _band(ws, row, start, end, value, *, fill=None, size=11, bold=False, color=INK):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    cell = literal_cell(ws, row, start, value)
    cell.font = Font(name="Calibri", size=size, bold=bold, color=color)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    if fill:
        for column in range(start, end + 1):
            ws.cell(row, column).fill = PatternFill("solid", fgColor=fill)
    return cell


def _jump(cell, sheet):
    cell.hyperlink = f"#'{sheet.replace(chr(39), chr(39) * 2)}'!A1"
    cell.font = Font(name="Calibri", size=11, bold=True, color=TEAL, underline="single")


def write_workbook_dashboard(wb: Workbook, view: ReportView) -> None:
    """Put financial meaning, concrete review work and a complete sheet directory first."""
    overview = build_output_overview(view)
    ws = wb.create_sheet("Start here", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = TEAL
    for column in range(1, 13):
        ws.column_dimensions[get_column_letter(column)].width = 11
    _band(
        ws, 1, 1, 12, "EDF evidence • case overview", fill=NAVY, color="FFFFFF", size=21, bold=True
    )
    ws.row_dimensions[1].height = 44
    _band(
        ws,
        2,
        1,
        12,
        f"{report_account_label(view)}   |   Analysis as of {overview.analysis_date}   |   Record dates: {overview.coverage}",
        color=GREY,
    )
    ws.row_dimensions[2].height = 34
    _band(
        ws,
        3,
        1,
        12,
        "Edit Review & Corrections or Dispute schedule, then regenerate in the app. These figures show the last generated result until you do.",
        fill="EAF7F5",
        color=NAVY,
        bold=True,
    )
    ws.row_dimensions[3].height = 32

    summary = view.snapshot.summary
    cards = (
        ("Recorded charges", summary.charges, "Costs recorded in the supplied bills"),
        ("Net payments / credits", summary.net_credits, "After refunds; supplied evidence only"),
        (
            "Recorded net movement",
            summary.charges - summary.net_credits,
            "Charges less net payments / credits",
        ),
        (
            "First observed balance",
            summary.opening_balance,
            "First dated balance in the selected ledger",
        ),
        (
            "Last observed balance",
            summary.closing_balance,
            "Last dated observation; not a current balance",
        ),
        (
            "Included evidence",
            overview.record_count,
            f"{overview.source_count:,} inventoried source{'s' if overview.source_count != 1 else ''} • {overview.ledger_count:,} ledger{'s' if overview.ledger_count != 1 else ''}",
        ),
    )
    for index, (title, value, hint) in enumerate(cards):
        top = 5 + index // 3 * 4
        left = 1 + index % 3 * 4
        _band(ws, top, left, left + 3, title, fill="EFF3F7", bold=True, color=GREY)
        cell = _band(
            ws,
            top + 1,
            left,
            left + 3,
            float(value) if value is not None else "Unavailable",
            fill="EFF3F7",
            size=22,
            bold=True,
            color=NAVY,
        )
        cell.number_format = "#,##0" if index == 5 else "£#,##0.00;[Red]-£#,##0.00;£0.00"
        _band(ws, top + 2, left, left + 3, hint, fill="EFF3F7", size=10, color=GREY)
        ws.row_dimensions[top].height = 25
        ws.row_dimensions[top + 1].height = 36
        ws.row_dimensions[top + 2].height = 32
    _band(ws, 13, 1, 12, overview.balance_note, color=GREY, size=10)
    ws.row_dimensions[13].height = 44
    _band(
        ws,
        14,
        1,
        12,
        f"{overview.exclusions:,} excluded record{'s' if overview.exclusions != 1 else ''}   •   {overview.duplicates:,} duplicate decision{'s' if overview.duplicates != 1 else ''}   •   {overview.notes:,} extraction / review note{'s' if overview.notes != 1 else ''}",
        color=GREY,
    )
    ws.row_dimensions[14].height = 25
    _band(ws, 16, 1, 12, "Where to focus next", fill=NAVY, color="FFFFFF", size=14, bold=True)
    ws.row_dimensions[16].height = 30
    row = 17
    for priority in overview.priorities:
        target = priority.sheet if priority.sheet in wb.sheetnames else "Workbook guide"
        cell = _band(ws, row, 1, 12, priority.title, bold=True)
        _jump(cell, target)
        _band(ws, row + 1, 1, 12, priority.action, color=GREY)
        ws.row_dimensions[row].height = 25
        ws.row_dimensions[row + 1].height = 36
        row += 3
    _band(ws, row, 1, 12, "Find the detail", fill=NAVY, color="FFFFFF", size=14, bold=True)
    ws.row_dimensions[row].height = 30
    row += 1
    review = [
        name
        for name in (
            "Review & Corrections",
            "EDF Evidence Report",
            "Dispute Flags",
            "Dispute schedule",
            "Rebilling ledger",
            "Rebilling components",
            "Rebilling comparisons",
            "Parse Errors",
            "Review history",
            "Workbook guide",
        )
        if name in wb.sheetnames
    ]
    money = [
        name
        for name in (
            "Annual Summary",
            "Period Charges",
            "Payment Analysis",
            "Balance Trend",
            "Reconciliation",
            "Reconciliation Drill-down",
            "SAP Financial Transactions",
        )
        if name in wb.sheetnames
    ]
    supporting = [
        sheet.title
        for sheet in wb
        if sheet.title not in {"Start here", *review, *money} and sheet.sheet_state == "visible"
    ]
    for column, title, names in (
        (1, "Review & corrections", review),
        (5, "Money & reconciliation", money),
        (9, "Supporting analysis & audit", supporting),
    ):
        _band(ws, row, column, column + 3, title, fill="EAF7F5", bold=True)
        for offset, name in enumerate(names, 1):
            _jump(_band(ws, row + offset, column, column + 3, name), name)
            ws.row_dimensions[row + offset].height = 30
    ws.freeze_panes = "A4"
    ws.print_area = f"A1:L{row + max(map(len, (review, money, supporting)), default=0)}"
    ws.print_title_rows = "1:3"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def polish_reference_sheets(wb: Workbook) -> None:
    """Make long tables easier to scan/print while keeping their formulas and positions."""
    for ws in wb:
        if ws.sheet_state != "visible" or ws.title in {
            "Start here",
            "Workbook guide",
            "Review & Corrections",
        }:
            continue
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        ws.sheet_properties.tabColor = (
            "B76A13"
            if ws.title == "Parse Errors"
            else "7B8D9C"
            if ws.title in {"Provenance", "Duplicate Entries", "Filtered (Below Min)"}
            else NAVY
        )
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        if ws.title in {
            "Bill revisions",
            "Rebilling ledger",
            "Rebilling components",
            "Rebilling comparisons",
            "Dispute schedule",
        }:
            ws.print_title_rows = "1:4"
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        if ws.title == "Parse Errors":
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.column_dimensions["A"].width = 12
            ws.column_dimensions["B"].width = 38
            ws.column_dimensions["C"].width = 90
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                lines = max(
                    ceil(len(str(cell.value or "")) / width)
                    for cell, width in zip(row, (12, 38, 90), strict=True)
                )
                ws.row_dimensions[row[0].row].height = min(120, max(30, lines * 14))
    for name in ("EDF Evidence Report", "Duplicate Entries", "Filtered (Below Min)"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        ws.freeze_panes = "G2"
        ws.sheet_view.zoomScale = 85
        widths = {
            "Date": 14,
            "Period From": 14,
            "Period To": 14,
            "Invoice #": 23,
            "Entry Type": 18,
            "Attachment Name": 28,
            "Details": 40,
            "Account": 20,
        }
        hidden = {
            "Source",
            "Sender",
            "Amount (£)",
            "% Change",
            "Logic Used",
            "Anomaly Flag",
            "Sub Periods",
            "Amount Kind",
            "Register Reading (kWh)",
            "Cancel/Rebill Admitted",
            "Source ID",
            "Source IDs",
            "Content Hash",
            "Original Filename",
            "Page From",
            "Page To",
            "Text Extraction",
            "OCR Pages",
            "OCR Minimum Confidence (%)",
            "Meter",
            "Register",
            "Unit Rate (p/kWh)",
            "Standing Chg (p/day)",
            "Tariff",
        }
        for column, cell in enumerate(ws[1], 1):
            header = cell.value
            letter = get_column_letter(column)
            ws.column_dimensions[letter].width = widths.get(header, 18)
            if header in hidden:
                ws.column_dimensions[letter].hidden = True
                ws.column_dimensions[letter].outlineLevel = 1
            cell.fill = PatternFill("solid", fgColor=NAVY)
        ws.sheet_properties.outlinePr.summaryRight = False
        ws.row_dimensions[1].height = 36
