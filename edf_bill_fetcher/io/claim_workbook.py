"""Editable, bounded Excel claim groups with stable evidence bindings and credit sources."""

from __future__ import annotations

import copy
import re
from math import ceil

from openpyxl.styles import Alignment, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from edf_bill_fetcher.application.record_review import corrected_record, review_rows
from edf_bill_fetcher.domain.analysis import AnalysisSnapshot
from edf_bill_fetcher.domain.json_values import object_value
from edf_bill_fetcher.domain.money import parse_money
from edf_bill_fetcher.helpers.excel_utils import literal_cell
from edf_bill_fetcher.io.writers.rebilling_ledger import reference_sheet, source_link
from edf_bill_fetcher.models.claim_tables import report_tables
from edf_bill_fetcher.models.report_view import RenderSelection, ReportView, build_report_view
from edf_bill_fetcher.processors.claim_reconciliation import (
    GROUP_FIELDS,
    build_reconciliation,
    empty_state,
    group_digest,
    separation_digest,
    validate_state,
)
from edf_bill_fetcher.processors.rebilling_ledger import build_ledger
from edf_bill_fetcher.processors.revisions import build_revisions

GROUP_SHEET, CREDIT_SHEET, OVERLAP_SHEET = (
    "Claim groups",
    "Credit allocations",
    "Claim overlap review",
)
INPUT_SHEETS = {GROUP_SHEET, CREDIT_SHEET, OVERLAP_SHEET}
GROUP_HEADERS = (
    "Claim ID",
    "Title",
    "Review state",
    "Amount challenged (£)",
    "Reduction before received relief (£)",
    "Scope / evidence",
    "Account-link evidence",
)
CREDIT_HEADERS = (
    "Claim ID",
    "Credit source reference",
    "Allocated (£)",
    "Treatment",
    "Allocation reason",
)
OVERLAP_HEADERS = (
    "First claim ID",
    "Second claim ID",
    "Review outcome",
    "Evidence for separate charges",
)
LIMIT = 1000


def member_groups(model: dict | None) -> dict[str, str]:
    """Represent all memberships so unchanged files preserve withdrawn and overlapping groups."""
    result: dict[str, list[str]] = {}
    if model:
        for group in model["input"]["groups"]:
            for key in group["members"]:
                result.setdefault(key, []).append(group["id"])
    return {key: "; ".join(sorted(groups)) for key, groups in result.items()}


def _credit_catalogue(view: ReportView) -> list[tuple[str, dict]]:
    model, ledger = view.claims, view.rebilling_ledger
    if not model or not ledger:
        return []
    allocated = {
        credit["row_id"] for group in model["input"]["groups"] for credit in group["credits"]
    }
    docs = [
        doc
        for doc in ledger["documents"]
        if doc["row_id"] in allocated
        or (
            (amount := parse_money(doc["record"].get("Transaction Amount (£)"))) is not None
            and amount < 0
        )
    ]
    return [(f"CRED{n:03d}", doc) for n, doc in enumerate(docs, 1)]


def _input_sheet(
    wb: Workbook,
    name: str,
    note: str,
    headers: tuple[str, ...],
    rows: list[tuple],
    money_columns: tuple[int, ...] = (),
) -> Worksheet:
    if len(rows) > LIMIT:
        raise ValueError(
            f"{name} exceeds the {LIMIT}-row Excel review limit; no entries were dropped"
        )
    ws = reference_sheet(wb, name, note, headers)
    # Preformat a useful working area, rather than thousands of empty cells
    # in every ordinary evidence workbook. Column defaults keep appended
    # inputs editable; parsing and dropdowns still enforce the hard limit.
    last_row = min(LIMIT + 4, max(24, len(rows) + 24))
    for column in range(1, len(headers) + 1):
        dimension = ws.column_dimensions[get_column_letter(column)]
        dimension.protection = Protection(locked=False)
        dimension.number_format = "£#,##0.00" if column in money_columns else "@"
    for row in range(5, last_row + 1):
        values = rows[row - 5] if row - 5 < len(rows) else ()
        for column in range(1, len(headers) + 1):
            cell = literal_cell(ws, row, column, values[column - 1] if values else None)
            cell.protection = Protection(locked=False)
            cell.fill = PatternFill("solid", fgColor="EAF7F5")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.number_format = "£#,##0.00" if column in money_columns else "@"
        if values:
            ws.row_dimensions[row].height = 65
    table = Table(
        displayName=name.replace(" ", "") + "Inputs",
        ref=f"A4:{get_column_letter(len(headers))}{last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    ws.add_table(table)
    ws.freeze_panes = "C5"
    ws.print_area = f"A1:{get_column_letter(len(headers))}{max(5, len(rows) + 4)}"
    return ws


def _list(ws: Worksheet, column: str, values: str) -> None:
    validation = DataValidation(type="list", formula1=values, allow_blank=True)
    validation.errorStyle, validation.showErrorMessage = "stop", True
    ws.add_data_validation(validation)
    validation.add(f"{column}5:{column}{LIMIT + 4}")


def write_claim_sheets(wb: Workbook, view: ReportView, links: dict[str, str]) -> None:
    """Export editable proposals and a protected catalogue of their exact credit sources."""
    model = view.claims
    if model is None:
        return
    groups = model["input"]["groups"]
    rows = [
        (
            g["id"],
            g["title"],
            g["status"],
            float(g["challenged"]) if g["challenged"] is not None else None,
            float(g["reduction"]) if g["reduction"] is not None else None,
            g["reason"],
            g["account_link_reason"],
        )
        for g in groups
    ]
    ws = _input_sheet(
        wb,
        GROUP_SHEET,
        "Create a group ID, then assign bills in Dispute schedule → Claim groups. Enter one consolidated challenge; do not sum repeated invoices. Reduction is before relief already received, including applicable VAT. Reviewed means you checked the evidence and allocation, not that EDF accepted it. Source changes make unchanged decisions stale; see Reconciled claims.",
        GROUP_HEADERS,
        rows,
        (4, 5),
    )
    _list(ws, "C", '"draft,reviewed,withdrawn"')
    for letter in ("B", "F", "G"):
        ws.column_dimensions[letter].width = 60
    catalogue = _credit_catalogue(view)
    refs = {doc["row_id"]: ref for ref, doc in catalogue}
    credits = [
        (g["id"], refs[c["row_id"]], float(c["amount"]), c["treatment"], c["reason"])
        for g in groups
        for c in g["credits"]
    ]
    ws = _input_sheet(
        wb,
        CREDIT_SHEET,
        "Choose the source reference from the protected catalogue on the right. already_applied_relief reduces the further reduction requested; cancellation_offset explains the earlier bill/credit history without deducting it again. Allocations cannot exceed the original recorded credit or spend it twice. Source references belong to this workbook's baseline, even if records are corrected in the same review.",
        CREDIT_HEADERS,
        credits,
        (3,),
    )
    _list(ws, "D", '"cancellation_offset,already_applied_relief"')
    if catalogue:
        _list(ws, "B", f"$H$5:$H${len(catalogue) + 4}")
    catalogue_headers = (
        "Source reference",
        "Invoice",
        "Date",
        "Account",
        "Recorded credit (£)",
        "Included",
        "Source file",
        "Pages",
        "Row ID",
    )
    for column, title in enumerate(catalogue_headers, 8):
        literal_cell(ws, 4, column, title)
        ws.column_dimensions[get_column_letter(column)].width = 27
    for row, (ref, doc) in enumerate(catalogue, 5):
        record = doc["record"]
        amount = parse_money(record.get("Transaction Amount (£)"))
        catalogue_values = (
            ref,
            record.get("Invoice #"),
            doc["invoice_date"],
            record.get("Account"),
            float(abs(amount)) if amount is not None and amount < 0 else None,
            "Yes" if doc["included"] else "No",
            record.get("Attachment Name"),
            "–".join(
                str(record[key]) for key in ("Page From", "Page To") if record.get(key) is not None
            ),
            doc["row_id"],
        )
        for column, value in enumerate(catalogue_values, 8):
            cell = literal_cell(ws, row, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        source_link(ws.cell(row, 14), record, links)
    ws.column_dimensions["P"].hidden = True
    ws.print_area = f"A1:O{max(5, len(credits) + 4, len(catalogue) + 4)}"
    ws.column_dimensions["E"].width = 65
    decisions = {tuple(item["groups"]): item for item in model["input"]["separations"]}
    suggested_pairs = sorted(
        {tuple(item["groups"]) for item in model["overlaps"]} - decisions.keys()
    )
    pairs = sorted(decisions) + suggested_pairs[: max(0, LIMIT - len(decisions))]
    rows = [
        (
            *pair,
            "separate_charges" if pair in decisions else "unreviewed",
            decisions.get(pair, {}).get("reason", ""),
        )
        for pair in pairs
    ]
    ws = _input_sheet(
        wb,
        OVERLAP_SHEET,
        "Potentially overlapping periods or account identities withhold a combined total. Put repeated demands in one group. Choose separate_charges only with source-backed evidence; shared documents cannot be waived. An unchanged decision becomes stale when either group's evidence or figures change. See Reconciled claims for current validity.",
        OVERLAP_HEADERS,
        rows,
    )
    _list(ws, "C", '"unreviewed,separate_charges"')
    ws.column_dimensions["D"].width = 90
    ws = reference_sheet(
        wb,
        "Reconciled claims",
        model["notice"],
        ("Item", "Amount / review", "Meaning / credits", "Scope / evidence"),
    )
    row = 5
    tables = report_tables(model, view.rebilling_ledger)
    if not tables:
        literal_cell(
            ws,
            row,
            1,
            "No claim groups yet. Enter challenged positions, create groups, and regenerate to obtain a reconciled proposal.",
        )
    for title, note, table in tables:
        literal_cell(ws, row, 1, title)
        literal_cell(ws, row + 1, 1, note)
        row += 2
        for table_values in table:
            for column, value in enumerate(table_values, 1):
                literal_cell(ws, row, column, value).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
            lines = max(
                sum(max(1, ceil(len(line) / 65)) for line in str(value).split("\n"))
                for value in table_values
            )
            ws.row_dimensions[row].height = min(409, max(32, lines * 15 + 8))
            row += 1
        row += 2
    for letter in "ABCD":
        ws.column_dimensions[letter].width = 70


def _rows(wb: Workbook, name: str, headers: tuple[str, ...]) -> list[tuple[int, list]]:
    if name not in wb.sheetnames:
        raise ValueError(f"{name} was removed; restore the original claim review sheets")
    ws = wb[name]
    if (
        any(
            row > LIMIT + 4 and column <= len(headers) and cell.value not in (None, "")
            for (row, column), cell in ws._cells.items()
        )
        or tuple(ws.cell(4, column).value for column in range(1, len(headers) + 1)) != headers
    ):
        raise ValueError(f"{name}: headers were changed or the {LIMIT}-row limit was exceeded")
    result = []
    for row in range(5, min(ws.max_row, LIMIT + 4) + 1):
        cells = [ws.cell(row, column) for column in range(1, len(headers) + 1)]
        if all(cell.value in (None, "") for cell in cells):
            continue
        if any(cell.data_type in {"f", "e"} for cell in cells):
            raise ValueError(
                f"{name}, row {row}: formulas and Excel errors are not valid review inputs; paste values"
            )
        result.append((row, [cell.value for cell in cells]))
    return result


def read_claim_state(
    wb: Workbook,
    snapshot: AnalysisSnapshot,
    record_decisions: list[dict],
    positions: list[dict] | None,
) -> dict | None:
    """Apply human changes while retaining evidence bindings for untouched decisions."""
    view = build_report_view(snapshot, RenderSelection())
    model = view.claims
    if model is None:
        if INPUT_SHEETS & set(wb.sheetnames):
            raise ValueError("This analysis has no preserved claim review baseline")
        return None
    baseline = review_rows(snapshot)
    edits = {item["row_id"]: item for item in record_decisions}
    if len(edits) != len(record_decisions) or set(edits) != {item["id"] for item in baseline}:
        raise ValueError("Record corrections are missing or duplicated")
    items = [
        {
            "id": item["id"],
            "included": edits[item["id"]]["include"],
            "record": corrected_record(item["record"], edits[item["id"]]["changes"]),
        }
        for item in baseline
    ]
    state = copy.deepcopy(model["input"])
    previous = {group["id"]: group for group in state["groups"]}
    groups = {}
    for row, values in _rows(wb, GROUP_SHEET, GROUP_HEADERS):
        key, title, status, challenged, reduction, reason, account_link = values
        if not isinstance(key, str) or key in groups:
            raise ValueError(f"{GROUP_SHEET}, row {row}: enter a unique text claim ID")
        groups[key] = {
            "id": key,
            "title": title or "",
            "status": status or "draft",
            "challenged": challenged,
            "reduction": reduction,
            "reason": reason or "",
            "account_link_reason": account_link or "",
            "members": [],
            "credits": [],
            "source_digest": previous.get(key, {}).get("source_digest", "0" * 64),
        }
    from edf_bill_fetcher.io.dispute_workbook import HEADERS as DISPUTE_HEADERS
    from edf_bill_fetcher.io.dispute_workbook import SHEET as DISPUTE_SHEET

    ws = wb[DISPUTE_SHEET]
    column = DISPUTE_HEADERS.index("Claim groups") + 1
    for row in range(5, ws.max_row + 1):
        key = ws.cell(row, len(DISPUTE_HEADERS)).value
        if key is None:
            continue
        cell = ws.cell(row, column)
        if cell.data_type in {"f", "e"} or (
            cell.value is not None and not isinstance(cell.value, str)
        ):
            raise ValueError(
                f"{DISPUTE_SHEET}, row {row}: claim memberships must be text IDs, not formulas"
            )
        owners = [value.strip() for value in re.split(r"[;,]", cell.value or "") if value.strip()]
        if len(set(owners)) != len(owners) or any(owner not in groups for owner in owners):
            raise ValueError(
                f"{DISPUTE_SHEET}, row {row}: unknown or duplicate claim group; create the group first"
            )
        for owner in owners:
            groups[owner]["members"].append(key)
    catalogue = {ref: doc["row_id"] for ref, doc in _credit_catalogue(view)}
    for row, values in _rows(wb, CREDIT_SHEET, CREDIT_HEADERS):
        owner, ref, amount, treatment, reason = values
        if (
            not isinstance(owner, str)
            or owner not in groups
            or not isinstance(ref, str)
            or ref not in catalogue
        ):
            raise ValueError(
                f"{CREDIT_SHEET}, row {row}: choose an existing claim ID and source reference from this workbook's catalogue"
            )
        groups[owner]["credits"].append(
            {
                "row_id": catalogue[ref],
                "amount": amount,
                "treatment": treatment,
                "reason": reason or "",
            }
        )
    state["groups"] = list(groups.values())
    old_decisions = {tuple(item["groups"]): item for item in state["separations"]}
    decisions = []
    for row, values in _rows(wb, OVERLAP_SHEET, OVERLAP_HEADERS):
        left, right, outcome, reason = values
        if outcome in (None, "", "unreviewed"):
            if reason not in (None, ""):
                raise ValueError(
                    f"{OVERLAP_SHEET}, row {row}: choose separate_charges to record this reason, or clear the unused reason"
                )
            continue
        if (
            outcome != "separate_charges"
            or not isinstance(left, str)
            or not isinstance(right, str)
            or left not in groups
            or right not in groups
            or left == right
        ):
            raise ValueError(
                f"{OVERLAP_SHEET}, row {row}: select two different existing groups and a supported review outcome"
            )
        pair = tuple(sorted((left, right)))
        decisions.append(
            {
                "groups": list(pair),
                "reason": reason or "",
                "source_digest": old_decisions.get(pair, {}).get("source_digest", "0" * 64),
            }
        )
    state["separations"] = decisions
    state = validate_state(items, state)
    legacy = object_value(snapshot.prepared.extraction.legacy_state)
    revisions = build_revisions(items, legacy.get("bill_revision_decisions", []))
    ledger = build_ledger(items, revisions, positions or [])
    for group in state["groups"]:
        old = previous.get(group["id"])
        if old is None or any(group[key] != old[key] for key in GROUP_FIELDS - {"source_digest"}):
            group["source_digest"] = group_digest(items, ledger, group)
            if group["status"] == "reviewed":
                single = {**empty_state(), "groups": [group]}
                checked = build_reconciliation(items, ledger, single)["groups"][0]
                if not checked["ready"]:
                    raise ValueError(
                        f"Claim {group['id']} cannot be marked reviewed: "
                        + " ".join(checked["issues"])
                    )
    indexed = {group["id"]: group for group in state["groups"]}
    for decision in state["separations"]:
        old = old_decisions.get(tuple(decision["groups"]))
        if old is None or decision["reason"] != old["reason"]:
            decision["source_digest"] = separation_digest(
                items, ledger, *(indexed[key] for key in decision["groups"])
            )
    return state
