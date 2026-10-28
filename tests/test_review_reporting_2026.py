"""End-to-end regressions for reporting and workbook evidence references."""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import cast

import openpyxl
import pandas as pd
import pdfplumber
import pytest
from bs4 import BeautifulSoup
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Table

from edf_bill_fetcher.domain.json_values import freeze_json
from edf_bill_fetcher.domain.requests import RunRequest
from edf_bill_fetcher.helpers.excel_utils import (
    _TEXT_SUPPRESSION_QUEUE,
    suppress_text_warning,
    suppress_text_warnings_post_save,
    text,
)
from edf_bill_fetcher.io.persistence.result_codec import migrate_legacy
from edf_bill_fetcher.io.persistence.result_store import load_result, save_result
from edf_bill_fetcher.io.reporters.docx_report import (
    _get_or_create_styles,
)
from edf_bill_fetcher.io.reporters.docx_report import (
    create_compensation_section as docx_compensation,
)
from edf_bill_fetcher.io.reporters.html_report import (
    create_compensation_section as html_compensation,
)
from edf_bill_fetcher.io.reporters.pdf_report import (
    CONTENT_WIDTH,
    create_anomaly_detail_section,
    create_appendix_full_evidence,
    create_appendix_glossary,
    create_evidence_index,
    create_timeline_section,
)
from edf_bill_fetcher.io.reporters.pdf_report import (
    create_compensation_section as pdf_compensation,
)
from edf_bill_fetcher.io.writers.evidence import write_evidence_sheet, write_summary_sheet
from edf_bill_fetcher.io.writers.export import export_to_excel, render_xlsx
from edf_bill_fetcher.io.writers.sap import write_sap_back_billing_sheets
from edf_bill_fetcher.models.config import ConfigDict
from edf_bill_fetcher.models.events import SapBackBillingEvent, SapEdfMatch
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view
from edf_bill_fetcher.models.snapshot_adapter import from_snapshot_json, to_snapshot_json
from edf_bill_fetcher.writers._helpers import detect_sap_back_billing_events


def _bill(invoice: str, *, charge: float = 500, account: str = "A-12345678") -> dict:
    return {
        "Source": "Local PDF Folder",
        "Source ID": invoice + account,
        "Account": account,
        "Fuel": "Electricity",
        "Date": "01/02/2025",
        "Period From": "01/01/2025",
        "Period To": "01/02/2025",
        "Invoice #": invoice,
        "Entry Type": "New Bill",
        "Amount (£)": 9000,
        "Period Charge (£)": charge,
        "Attachment Name": invoice + ".pdf",
        "Details": "Original evidence",
    }


def _sap_rows() -> list[dict]:
    return [
        {
            "Clearing Document": "CL1",
            "Clearing Date": "01/02/2025",
            "Posting Date": "01/02/2025",
            "Account": "A-12345678",
            "Amount": 500 if idx == 0 else 0,
            "Document No.": "D" + str(idx),
        }
        for idx in range(4)
    ]


@pytest.mark.parametrize("saved", [False, True], ids=["fresh", "saved-result"])
def test_sap_match_links_to_full_evidence_row_and_displays_match_amount(
    tmp_path: Path, saved: bool
) -> None:
    records = [
        _bill("SMALL", charge=10),
        _bill("MATCH"),
        _bill("LATER", charge=600),
    ]
    records[0]["Date"] = "01/01/2025"
    records[2].update(
        {"Date": "01/03/2025", "Period From": "01/02/2025", "Period To": "01/03/2025"}
    )
    config = cast(ConfigDict, {"analysis_min": 100, "use_dedup": False, "as_of": "2025-04-01"})
    output = tmp_path / "evidence.xlsx"
    if saved:
        request = RunRequest(
            "report-review",
            datetime(2025, 4, 1, tzinfo=timezone.utc),
            date(2025, 4, 1),
            (),
            freeze_json(config),
        )
        snapshot = migrate_legacy({"records": records, "sap_financial_rows": _sap_rows()}, request)
        source = tmp_path / "result.json"
        save_result(snapshot, source, result_id="report-review")
        original = source.read_bytes()
        view = build_report_view(load_result(source), RenderSelection(("xlsx",)))
        render_xlsx(view, str(output))
        assert source.read_bytes() == original
    else:
        export_to_excel(records, str(output), [], config, sap_rows={"financial": _sap_rows()})
    workbook = openpyxl.load_workbook(output)
    try:
        matched = workbook["SAP ↔ EDF Matched Events"]
        evidence = workbook["EDF Evidence Report"]
        headers = {cell.value: cell.column for cell in evidence[1]}
        rows = [row for row in matched.iter_rows(min_row=6) if row[3].value == "MATCH"]
        assert rows
        for row in rows:
            assert row[3].hyperlink is not None
            target = int(row[3].hyperlink.location.rsplit("!A", 1)[1])
            assert evidence.cell(target, headers["Invoice #"]).value == "MATCH"
            assert row[6].value == 500
            assert row[7].value == row[2].value - row[6].value
    finally:
        workbook.close()


@pytest.mark.parametrize("ambiguous", [False, True])
def test_sap_match_does_not_use_another_ledger_or_ambiguous_invoice(ambiguous: bool) -> None:
    target = _bill("REUSED")
    other = _bill("REUSED", account="A-99999999")
    candidates = [other, dict(target)]
    if ambiguous:
        candidates.append(dict(target))
    event = SapBackBillingEvent("CL1", pd.Timestamp("2025-02-01"), "", net_amount=500)
    match = SapEdfMatch(event, target, "High", 90, 0, 0, "")
    workbook = openpyxl.Workbook()
    write_sap_back_billing_sheets(
        workbook, [event], [match], 4, edf_rows=candidates, edf_first_row=2
    )
    link = workbook["SAP ↔ EDF Matched Events"].cell(6, 4).hyperlink
    if ambiguous:
        assert link is None
    else:
        assert link is not None
        assert link.location == "'EDF Evidence Report'!A3"


def test_excel_text_survives_xml_controls_without_becoming_a_formula(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    text(workbook.active, 1, 1, '\x0b=HYPERLINK("https://example.org")\x00')
    destination = tmp_path / "controls.xlsx"
    workbook.save(destination)
    restored = openpyxl.load_workbook(destination)
    try:
        assert restored.active["A1"].data_type == "s"
        assert restored.active["A1"].value == '\'=HYPERLINK("https://example.org")'
    finally:
        restored.close()


def test_annual_summary_does_not_invent_zero_balances_for_unobserved_years() -> None:
    workbook = openpyxl.Workbook()
    write_summary_sheet(workbook.active, [2024, 2025], "EDF Evidence Report", 10)
    sheet = workbook.active
    # Excel MAXIFS/MINIFS return zero for no matches, so IFERROR alone is insufficient.
    for row in (2, 3):
        for column in (2, 5, 6):
            assert f'IF(C{row}=0,"",' in sheet.cell(row, column).value
    for column in (2, 5, 6):
        assert 'IF(SUM(C2:C3)=0,"",' in sheet.cell(4, column).value


@pytest.mark.parametrize("builder", ["appendix", "index"])
def test_pdf_evidence_tables_have_positive_widths_within_the_page(builder: str) -> None:
    data = pd.DataFrame([_bill("LONG-INVOICE-REFERENCE")])
    elements = (
        create_appendix_full_evidence(data)
        if builder == "appendix"
        else create_evidence_index(data, None)
    )
    tables = [element for element in elements if isinstance(element, Table)]
    assert tables
    for table in tables:
        assert all(width > 0 for width in table._argW)
        assert sum(table._argW) <= CONTENT_WIDTH + 0.01


def test_multiple_excel_warning_ranges_share_one_schema_container(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    _TEXT_SUPPRESSION_QUEUE.clear()
    suppress_text_warning(worksheet, "A", 1, 5)
    suppress_text_warning(worksheet, "B", 1, 5)
    destination = tmp_path / "warnings.xlsx"
    workbook.save(destination)
    suppress_text_warnings_post_save(str(destination))
    with zipfile.ZipFile(destination) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    containers = root.findall("{*}ignoredErrors")
    assert len(containers) == 1
    assert {child.attrib["sqref"] for child in containers[0]} == {"A1:A5", "B1:B5"}
    assert _TEXT_SUPPRESSION_QUEUE == {}


def test_failed_excel_warning_rewrite_keeps_original_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = openpyxl.Workbook()
    _TEXT_SUPPRESSION_QUEUE.clear()
    suppress_text_warning(workbook.active, "A", 1, 5)
    destination = tmp_path / "preserved.xlsx"
    workbook.save(destination)
    original = destination.read_bytes()

    def fail_write(*args, **kwargs):
        raise OSError("Disk full")

    monkeypatch.setattr(zipfile.ZipFile, "writestr", fail_write)
    try:
        with pytest.raises(OSError, match="Disk full"):
            suppress_text_warnings_post_save(str(destination))
        assert destination.read_bytes() == original
        assert list(tmp_path.iterdir()) == [destination]
    finally:
        _TEXT_SUPPRESSION_QUEUE.clear()


def test_missing_authoritative_source_id_cannot_link_a_same_name_original() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    write_evidence_sheet(
        worksheet,
        pd.DataFrame([_bill("SAME-NAME")]),
        saved_paths={"SAME-NAME.pdf": "evidence_files/different-source.pdf"},
    )
    headers = {cell.value: cell.column for cell in worksheet[1]}
    assert worksheet.cell(2, headers["Attachment Name"]).hyperlink is None


def test_sap_same_clearing_document_keeps_accounts_and_saved_links_separate() -> None:
    first = _sap_rows()
    second = [
        dict(row, Account="A-99999999", Amount=200 if idx == 0 else 0)
        for idx, row in enumerate(first)
    ]
    events = detect_sap_back_billing_events(first + second)
    assert len(events) == 2
    assert [event.clearing_doc for event in events] == ["CL1", "CL1"]
    assert sorted(event.net_amount for event in events) == [200, 500]
    records = [_bill("FIRST"), _bill("SECOND", charge=200, account="A-99999999")]
    matches = [
        SapEdfMatch(event, records[idx], "High", 90, 0, 0, "") for idx, event in enumerate(events)
    ]
    # The renderer must handle independently decoded event and match models.
    events = from_snapshot_json(to_snapshot_json(events))
    matches = from_snapshot_json(to_snapshot_json(matches))
    workbook = openpyxl.Workbook()
    write_sap_back_billing_sheets(workbook, events, matches, 4, records, edf_first_row=2)
    matched = workbook["SAP ↔ EDF Matched Events"]
    summaries = workbook["SAP Back-billing Events"]
    for row in (6, 7):
        link = matched.cell(row, 1).hyperlink
        assert link is not None
        target = int(link.location.rsplit("!A", 1)[1])
        assert summaries.cell(target, 5).value == matched.cell(row, 3).value
    assert matched.cell(6, 1).hyperlink.location != matched.cell(7, 1).hyperlink.location


def test_sap_unknown_and_conflicting_ledgers_cannot_form_an_artificial_cluster() -> None:
    known = _sap_rows()[:2]
    unknown = [
        {key: value for key, value in row.items() if key != "Account"} for row in _sap_rows()[:2]
    ]
    assert detect_sap_back_billing_events(known + unknown) == []


def test_noncontiguous_compensation_categories_keep_every_row_under_its_own_heading() -> None:
    payload = [
        {
            "category": category,
            "invoice_ref": str(idx),
            "indicative_amount": idx,
            "legal_basis": "Indicative scenario requiring review",
        }
        for idx, category in enumerate(
            ("back_billing_excess", "credit_hold_interest", "back_billing_excess"), 1
        )
    ]
    pdf_tables = [
        element
        for element in pdf_compensation(pd.DataFrame(), payload=payload)
        if isinstance(element, Table)
    ]
    assert [len(table._cellvalues) - 1 for table in pdf_tables] == [2, 1]
    assert all(
        row[0].getPlainText() == "Potential back-billing adjustment"
        for row in pdf_tables[0]._cellvalues[1:]
    )
    document = Document()
    docx_compensation(document, _get_or_create_styles(document), pd.DataFrame(), payload=payload)
    assert [len(table.rows) - 1 for table in document.tables] == [2, 1]
    assert all(
        row.cells[0].text == "Potential back-billing adjustment"
        for row in document.tables[0].rows[1:]
    )
    soup = BeautifulSoup(html_compensation(pd.DataFrame(), payload=payload), "html.parser")
    tables = soup.select("table")
    assert [len(table.select("tr")) - 1 for table in tables] == [2, 1]
    assert [row.select("td")[0].text for row in tables[0].select("tr")[1:]] == [
        "Potential back-billing adjustment"
    ] * 2


@pytest.mark.parametrize("section", ["timeline", "anomaly", "glossary", "compensation"])
def test_pdf_long_explanations_and_legal_context_stay_inside_the_page(section: str) -> None:
    record = _bill("INVOICE")
    record["Details"] = "A very long explanation about the original evidence and its limits. " * 5
    frame = pd.DataFrame([record])
    flags = [("LARGE JUMP", "01/02/2025", 0, record["Details"], "HIGH")]
    if section == "timeline":
        elements = create_timeline_section(frame, flags)
    elif section == "anomaly":
        elements = create_anomaly_detail_section(flags, frame)
    elif section == "glossary":
        elements = create_appendix_glossary()
    else:
        elements = pdf_compensation(
            frame,
            payload=[{"category": "credit_hold_interest", "legal_basis": record["Details"]}],
        )
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer).build(elements)
    buffer.seek(0)
    with pdfplumber.open(buffer) as document:
        assert any(page.chars for page in document.pages)
        for page in document.pages:
            assert all(0 <= char["x0"] <= char["x1"] <= page.width for char in page.chars)
