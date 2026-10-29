"""Accepted results render honestly; failed sections preserve the previous report."""

from dataclasses import replace

import pandas as pd
import pdfplumber
import pytest
from bs4 import BeautifulSoup
from docx import Document
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate

from edf_bill_fetcher.domain.json_values import freeze_json, object_value
from edf_bill_fetcher.io.persistence.result_codec import decode_result, encode_result
from edf_bill_fetcher.io.reporters import docx_report as word
from edf_bill_fetcher.io.reporters import html_report as html
from edf_bill_fetcher.io.reporters import pdf_report as pdf
from edf_bill_fetcher.models.report_models import (
    OfgemComparison,
    OfgemRow,
    compute_forecast,
    compute_statistical_analysis,
)
from edf_bill_fetcher.models.report_sections import RenderContext
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view
from edf_bill_fetcher.models.snapshot_adapter import to_snapshot_json
from tests.synthetic_analysis import snapshot
from tests.test_financial_identity_boundaries import bill

MODULES = {"html": html, "docx": word, "pdf": pdf}


@pytest.fixture(scope="module")
def accepted():
    return snapshot([bill()])


def text(path):
    if path.suffix == ".html":
        return BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser").get_text(" ")
    if path.suffix == ".docx":
        doc = Document(path)
        return " ".join(
            [
                *(p.text for p in doc.paragraphs),
                *(cell.text for table in doc.tables for row in table.rows for cell in row.cells),
            ]
        )
    with pdfplumber.open(path) as document:
        contents = []
        for page in document.pages:
            contents.append(page.extract_text() or "")
            contents.extend(
                (cell or "").replace("\n", "")
                for table in page.extract_tables()
                for row in table
                for cell in row
            )
        return " ".join(contents)


def render(view, path, fmt):
    return getattr(MODULES[fmt], "render_" + fmt)(view, str(path))


@pytest.mark.parametrize("fmt", MODULES)
@pytest.mark.parametrize("section", ["toc", "data_quality"])
def test_failed_section_preserves_previous_report_and_returns_failure(
    tmp_path, monkeypatch, accepted, fmt, section
):
    module = MODULES[fmt]
    name = "create_table_of_contents" if section == "toc" else "create_data_quality_section"

    def fail(*args, **kwargs):
        raise RuntimeError("Simulated section failure")

    monkeypatch.setattr(module, name, fail)
    destination = tmp_path / ("report." + fmt)
    destination.write_bytes(b"Previous complete report")
    view = build_report_view(accepted, RenderSelection((fmt,), (section,)))
    with pytest.raises(RuntimeError, match="Simulated section failure"):
        render(view, destination, fmt)
    assert destination.read_bytes() == b"Previous complete report"


@pytest.mark.parametrize("fmt", MODULES)
def test_empty_unreviewed_analysis_cannot_replace_previous_report(tmp_path, fmt):
    destination = tmp_path / ("report." + fmt)
    destination.write_bytes(b"Previous report")
    view = build_report_view(snapshot([]), RenderSelection((fmt,)))
    with pytest.raises(ValueError, match="No records to report"):
        render(view, destination, fmt)
    assert destination.read_bytes() == b"Previous report"


@pytest.mark.parametrize("fmt", MODULES)
def test_sparse_accepted_evidence_frame_keeps_missing_values_unavailable(tmp_path, accepted, fmt):
    models = object_value(accepted.models)
    models["sap"]["payload"]["frames"]["evidence"] = to_snapshot_json(
        pd.DataFrame([{"Date": "01/02/2024"}])
    )
    historical = decode_result(
        encode_result(replace(accepted, models=freeze_json(models)), "sparse")
    )
    view = build_report_view(historical, RenderSelection((fmt,), ("appendix_full_evidence",)))
    destination = tmp_path / ("report." + fmt)
    render(view, destination, fmt)
    contents = text(destination)
    assert "01/02/2024" in contents
    assert "£0.00" not in contents
    assert historical.summary == accepted.summary


def appendix(tmp_path, fmt, frame, filtered=None, config=None):
    path = tmp_path / ("appendix." + fmt)
    if fmt == "html":
        path.write_text(
            html.create_appendix_full_evidence(frame, filtered, config), encoding="utf-8"
        )
    elif fmt == "docx":
        doc = Document()
        styles = word._get_or_create_styles(doc)
        word.create_appendix_full_evidence(doc, styles, frame, filtered)
        doc.save(path)
    else:
        SimpleDocTemplate(str(path), pagesize=landscape(A4)).build(
            pdf.create_appendix_full_evidence(frame, filtered, config)
        )
    return path


@pytest.mark.parametrize("fmt", MODULES)
def test_empty_appendix_explains_absence_of_evidence(tmp_path, fmt):
    path = appendix(tmp_path, fmt, pd.DataFrame())
    assert "No records available" in text(path)


@pytest.mark.parametrize("fmt", MODULES)
@pytest.mark.parametrize("config", [None, {"min_amount": 12.5}])
def test_filtered_appendix_preserves_all_rows_and_explicit_threshold(tmp_path, fmt, config):
    filtered = [bill(**{"Invoice #": name}) for name in ("FILTER-1", "FILTER-2")]
    path = appendix(tmp_path, fmt, pd.DataFrame([bill()]), filtered, config)
    contents = text(path)
    assert "FILTER-1" in contents
    assert "FILTER-2" in contents
    if fmt != "docx":
        assert ("£12.50" if config else "£500") in contents


def test_html_evidence_cap_is_explicit_and_chronological(tmp_path):
    rows = [
        bill(**{"Date": when.strftime("%d/%m/%Y"), "Invoice #": f"ROW-{n:03}"})
        for n, when in enumerate(pd.date_range("2024-01-01", periods=151))
    ]
    path = appendix(tmp_path, "html", pd.DataFrame(rows).iloc[::-1])
    contents = text(path)
    assert "151 rows > 150 row limit" in contents
    assert "ROW-000" in contents
    assert "ROW-149" in contents
    assert "ROW-150" not in contents


def test_empty_contents_and_absent_revision_models_do_not_invent_sections(tmp_path):
    ctx = RenderContext([])
    assert "No report sections selected" in html.create_table_of_contents(ctx)
    doc = Document()
    styles = word._get_or_create_styles(doc)
    word.create_table_of_contents(doc, styles, ctx)
    assert any("No report sections selected" in p.text for p in doc.paragraphs)
    path = tmp_path / "contents.pdf"
    SimpleDocTemplate(str(path)).build(pdf.create_table_of_contents(ctx))
    assert "No report sections selected" in text(path)
    assert html.create_revision_section(None) == ""
    assert pdf.create_revision_section(None) == []
    before = len(doc.paragraphs)
    word.create_revision_section(doc, styles, None)
    assert len(doc.paragraphs) == before


@pytest.mark.parametrize("section", ["statistical", "payment", "forecast"])
def test_html_compatibility_sections_explain_insufficient_data(section):
    function = {
        "statistical": html.create_statistical_analysis,
        "payment": html.create_payment_analysis,
        "forecast": html.create_forecast_section,
    }[section]
    contents = BeautifulSoup(function(pd.DataFrame([bill()])), "html.parser").get_text(" ").lower()
    expected = {
        "statistical": "no statistical conclusion",
        "payment": "no payment",
        "forecast": "no forecast",
    }
    assert expected[section] in contents


def test_default_section_contexts_and_missing_caption_style_remain_usable():
    doc = Document()
    doc.styles["Caption"].delete()
    styles = word._get_or_create_styles(doc)
    count = len(doc.styles)
    word._get_or_create_styles(doc)
    assert len(doc.styles) == count
    assert doc.styles["Caption"].font.size.pt == 9
    frame = pd.DataFrame([bill()])
    args = (
        frame,
        {"as_of": "2025-03-01"},
        "SYNTHETIC",
        {"MEDIUM": 1},
        1,
        100,
        0,
        "01/01/2024",
        "01/02/2024",
    )
    assert "1 MEDIUM-severity" in html.create_executive_summary(*args)
    word.create_executive_summary(doc, styles, *args)
    assert "No" in html.create_key_findings_table([])
    word.create_key_findings_table(doc, styles, [])
    assert "No specific anomalies" in html.create_anomaly_detail_section([], frame)
    word.create_anomaly_detail_section(doc, styles, [], frame)
    timeline = html.create_timeline_section(frame, [])
    assert "01/02/2024" in timeline
    assert "A-EXAMPLE" in timeline
    many = pd.DataFrame([bill(**{"Invoice #": f"EVENT-{n}"}) for n in range(52)])
    word.create_timeline_section(doc, styles, many, [])
    assert any("Showing first 50 of 52 events" in p.text for p in doc.paragraphs)
    word.create_statistical_analysis(doc, styles, frame)
    assert "Valuation date: 2025-03-01" in html.create_appendix_methodology({"as_of": "2025-03-01"})
    word.create_appendix_methodology(doc, styles, {"as_of": "2025-03-01"})
    assert "Period Charge" in html.create_appendix_glossary()


def test_pdf_retains_legacy_exceeds_status_and_optional_forecast_column(tmp_path):
    comparison = OfgemComparison([OfgemRow("2024-Q1", 50, 20, 30, "EXCEEDS")], 1, 0, 0, 50, 50)
    elements = pdf.create_ofgem_comparison(pd.DataFrame(), payload=comparison)
    forecast = replace(
        compute_forecast(pd.DataFrame()),
        n=12,
        linear_forecast=[70.0] * 6,
        ema_forecast=[80.0] * 6,
        hw_forecast=[90.0] * 6,
    )
    elements.extend(pdf.create_forecast_section(pd.DataFrame(), payload=forecast))
    path = tmp_path / "accepted-models.pdf"
    SimpleDocTemplate(str(path)).build(elements)
    contents = text(path)
    assert "EXCEEDS" in contents
    assert "Holt-Winters" in contents
    assert "£90.00" in contents


def test_pdf_explains_unavailable_normality_test(tmp_path):
    frame = pd.DataFrame(
        [bill(**{"Date": f"0{n}/02/2024", "Balance (£)": n * 100}) for n in (1, 2, 3)]
    )
    payload = replace(compute_statistical_analysis(frame), shapiro_stat=None, shapiro_p=None)
    path = tmp_path / "statistics.pdf"
    SimpleDocTemplate(str(path)).build(pdf.create_statistical_analysis(frame, payload=payload))
    assert "Scipy not available" in text(path)


@pytest.mark.parametrize("blank", [False, True])
def test_invalid_percentage_is_not_displayed_as_zero(blank):
    assert pdf.fmt_pct("unavailable", blank_if_na=blank) == ("" if blank else "unavailable")


def test_empty_html_table_is_absent_instead_of_a_fake_empty_record():
    assert html._table([]) == ""
