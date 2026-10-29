"""Corrupt saved histories and claims must fail without recalculating accepted figures."""

import copy

import pytest
from openpyxl import Workbook

from edf_bill_fetcher.domain.json_values import object_value
from edf_bill_fetcher.io.persistence.result_codec import (
    decode_result,
    encode_result,
    migrate_legacy,
)
from edf_bill_fetcher.io.writers.rebilling_ledger import write_ledger_sheets
from edf_bill_fetcher.models.events import SapBackBillingEvent
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view
from edf_bill_fetcher.models.snapshot_adapter import to_snapshot_json
from edf_bill_fetcher.processors.claim_reconciliation import (
    build_reconciliation,
    empty_state,
    separation_digest,
)
from edf_bill_fetcher.processors.revisions import build_revisions
from tests.synthetic_analysis import snapshot
from tests.test_claim_reconciliation import evidence, group, make_ledger
from tests.test_financial_identity_boundaries import bill
from tests.test_saved_result_validation import request


@pytest.fixture(scope="module")
def accepted_models():
    items, _ = evidence()
    for item in items[:2]:
        amount = item["record"]["Period Charge (£)"]
        item["record"]["Sub Periods"] = f"01/01/2024|31/01/2024|{amount}|100|{amount}"
        item["record"]["Rebilling Components"] = [
            {"kind": "vat", "amount": 5, "detail": "Observed VAT"}
        ]
    revisions = build_revisions(items)
    ledger = make_ledger(items)
    state = empty_state()
    state["groups"] = [group(items, ledger, "C1", ("a",)), group(items, ledger, "C2", ("b",))]
    state["separations"] = [
        {
            "groups": ["C1", "C2"],
            "reason": "Verified distinct charge allocations on the original bills",
            "source_digest": separation_digest(items, ledger, *state["groups"]),
        }
    ]
    claims = build_reconciliation(items, ledger, state)
    assert claims["ready"]
    assert claims["proposed_reduction"] == "240.00"
    original = snapshot([item["record"] for item in items])
    payload = encode_result(original, "accepted-integrity")
    sap = payload["analysis"]["models"]["sap"]["payload"]
    sap.update(
        rebilling_ledger=to_snapshot_json(ledger),
        bill_revisions=to_snapshot_json(revisions),
        claim_reconciliation=to_snapshot_json(claims),
    )
    # Verify the complete fixture through the public decoder before corrupting it.
    assert encode_result(decode_result(payload), "accepted-integrity") == payload
    return payload


def corrupt(payload, path, value):
    for key in path[:-1]:
        payload = payload[key]
    payload[path[-1]] = value


@pytest.mark.parametrize(
    ("path", "value", "message"),
    [
        (("documents",), {}, "ledger entries"),
        (("comparisons",), {}, "ledger entries"),
        (("documents", 0, "invoice_date"), 20240201, "rebilling date"),
        (("documents", 0, "invoice_date"), "20240201", "rebilling date"),
        (("documents", 0, "first_demand_date"), "2024-01-01", "Noncanonical stored dispute"),
        (("documents", 0, "position", "challenged_amount"), "200", "Noncanonical stored dispute"),
        (("documents", 2, "cutoff_date"), "2024-01-01", "lack a dispute position"),
        (("documents", 2, "needs_review"), True, "lack a dispute position"),
        (("documents", 0, "components"), {}, "rebilling components"),
        (("documents", 0, "components", 0, "units"), True, "component quantity"),
        (("documents", 0, "components", 0, "rate"), "100", "component quantity"),
        (("documents", 0, "components", 0, "kind"), "invented", "component kind"),
        (("documents", 0, "components", 0, "kind"), [], "component kind"),
        (("documents", 0, "incoming"), {}, "rebilling relationships"),
        (("documents", 0, "outgoing"), {}, "rebilling relationships"),
        (("documents", 0, "outgoing", 0, "newer"), [], "unknown document"),
        (("documents", 0, "outgoing", 0, "status"), [], "unknown document"),
        (("comparisons", 0, "status"), [], "rebilling comparison"),
        (("comparisons", 0, "other_components"), {}, "other component comparisons"),
        (("comparisons", 0, "other_components", 0, "kind"), "invented", "comparison kind"),
        (("comparisons", 0, "other_components", 0, "kind"), [], "comparison kind"),
        (("comparisons", 0, "components", 0, "status"), "invented", "comparison status"),
        (("comparisons", 0, "components", 0, "status"), [], "comparison status"),
        (("comparisons", 0, "components", 0, "residual"), "-0.00", "Signed zero"),
    ],
)
def test_corrupt_accepted_ledger_is_rejected_without_mutation(
    accepted_models, path, value, message
):
    payload = copy.deepcopy(accepted_models)
    corrupt(payload["analysis"]["models"]["sap"]["payload"]["rebilling_ledger"], path, value)
    before = copy.deepcopy(payload)
    with pytest.raises(ValueError, match=message):
        decode_result(payload)
    assert payload == before


@pytest.mark.parametrize(
    ("path", "value", "message"),
    [
        (("version",), True, "Unsupported claim reconciliation"),
        (("ready",), 1, "Unsupported claim reconciliation"),
        (("unresolved_positions",), "a", "Unknown unresolved claim"),
        (("unresolved_positions",), ["unknown"], "Unknown unresolved claim"),
        (("unresolved_positions",), ["a", "a"], "Duplicate unresolved claim"),
        (("groups",), [], "group membership disagrees"),
        (("groups", 0, "input", "title"), "Forged title", "disagrees with preserved inputs"),
        (("groups", 0, "needs_review"), 0, "group review status"),
        (("groups", 0, "needs_review"), True, "cannot be ready for totaling"),
        (("groups", 0, "ready"), False, "cannot publish a full proposed total"),
        (("groups", 0, "issues"), [1], "group review status"),
        (("groups", 0, "scope", "accounts"), ["one", "one"], "ledger scope"),
        (("overlaps",), {}, "overlap or ledger results"),
        (("ledgers",), {}, "overlap or ledger results"),
        (("overlaps", 0, "groups"), ["C1", "unknown"], "overlap review"),
        (("overlaps", 0, "groups"), ["C2", "C1"], "overlap review"),
        (("overlaps", 0, "kind"), "invented", "overlap review"),
        (("overlaps", 0, "kind"), [], "overlap review"),
        (("overlaps", 0, "kind"), "shared_documents", "shared or stale"),
        (("overlaps", 0, "needs_review"), True, "shared or stale"),
        (("overlaps", 0, "resolved"), False, "cannot publish a full proposed total"),
        (("unresolved_positions",), ["a"], "cannot publish a full proposed total"),
        (("ledgers", 0, "groups"), ["unknown"], "Unknown ledger claim"),
        (("ledgers", 0, "groups"), {}, "Unknown ledger claim"),
    ],
)
def test_corrupt_accepted_claims_never_publish_a_total(accepted_models, path, value, message):
    payload = copy.deepcopy(accepted_models)
    corrupt(payload["analysis"]["models"]["sap"]["payload"]["claim_reconciliation"], path, value)
    before = copy.deepcopy(payload)
    with pytest.raises(ValueError, match=message):
        decode_result(payload)
    assert payload == before


def test_noncanonical_claim_inputs_and_partial_totals_are_rejected(accepted_models):
    payload = copy.deepcopy(accepted_models)
    model = payload["analysis"]["models"]["sap"]["payload"]["claim_reconciliation"]
    model["input"]["groups"][0]["challenged"] = "120"
    with pytest.raises(ValueError, match="inputs are not canonical"):
        decode_result(payload)
    payload = copy.deepcopy(accepted_models)
    model = payload["analysis"]["models"]["sap"]["payload"]["claim_reconciliation"]
    model.update(ready=False, proposed_reduction="240.00")
    with pytest.raises(ValueError, match="cannot publish a full proposed total"):
        decode_result(payload)


@pytest.mark.parametrize("key", ["ledger_overview", "bill_revisions"])
def test_bounded_accepted_assessments_reject_non_arrays(accepted_models, key):
    payload = copy.deepcopy(accepted_models)
    sap = payload["analysis"]["models"]["sap"]["payload"]
    if key == "bill_revisions":
        sap[key]["relationships"] = {}
    else:
        sap[key] = {}
    with pytest.raises(ValueError, match="Invalid"):
        decode_result(payload)


@pytest.mark.parametrize(
    ("field", "value"),
    [("scope", ["one"]), ("references", [""]), ("records", True), ("closing_date", "20240301")],
)
def test_corrupt_ledger_observation_identity_or_date_is_rejected(accepted_models, field, value):
    payload = copy.deepcopy(accepted_models)
    payload["analysis"]["models"]["sap"]["payload"]["ledger_overview"][0][field] = value
    with pytest.raises(ValueError, match="ledger|Ledger"):
        decode_result(payload)


def test_claims_cannot_survive_removal_of_their_preserved_ledger(accepted_models):
    payload = copy.deepcopy(accepted_models)
    del payload["analysis"]["models"]["sap"]["payload"]["rebilling_ledger"]
    with pytest.raises(ValueError, match="needs a preserved rebilling ledger"):
        decode_result(payload)


def test_malformed_reconciliation_detail_is_rejected(accepted_models):
    payload = copy.deepcopy(accepted_models)
    payload["analysis"]["models"]["reconciliation"]["payload"]["details"] = [
        {"entity": "Financial"}
    ]
    with pytest.raises(ValueError, match="Malformed reconciliation detail"):
        decode_result(payload)


def test_unavailable_balances_and_legacy_revision_absence_are_preserved(accepted_models):
    payload = copy.deepcopy(accepted_models)
    sap = payload["analysis"]["models"]["sap"]["payload"]
    del sap["bill_revisions"]
    sap["union_total"] = None
    for row in sap["ledger_overview"]:
        for field in ("opening_date", "opening_balance", "closing_date", "closing_balance"):
            row[field] = None
    for row in sap["claim_reconciliation"]["ledgers"]:
        for field in ("observed_balance", "balance_date", "illustrative_balance"):
            row[field] = None
    restored = encode_result(decode_result(payload), "accepted-integrity")
    assert restored == payload


def test_multirow_sap_event_keeps_exact_source_rows_in_order():
    rows = [{"Amount": 100, "Document": "FIRST"}, {"Amount": -20, "Document": "SECOND"}]
    accepted = snapshot([bill()], {"sap_financial_rows": rows})
    payload = encode_result(accepted, "multirow")
    sap = payload["analysis"]["models"]["sap"]["payload"]
    sap["events"] = [
        to_snapshot_json(SapBackBillingEvent("CLEAR-1", None, "Known event", rows=rows))
    ]
    sap["event_source_indexes"] = [[0, 1]]
    assert encode_result(decode_result(payload), "multirow") == payload


def test_legacy_originals_and_manifest_are_retained_without_missing_source_notes():
    manifest = {"source": {"Attachment Name": "fixture.pdf"}}
    result = migrate_legacy(
        {"records": [], "original_records": [], "source_manifest": manifest}, request()
    )
    state = object_value(result.prepared.extraction.legacy_state)
    assert state["source_manifest"] == manifest
    assert len(state["migration_notes"]) == 1
    assert "Analysis recomputed" in state["migration_notes"][0]


@pytest.mark.parametrize("value", [[], {"OLD": []}, {"OLD": ["", True]}, {"OLD": ["NEW", 1]}])
def test_corrupt_superseded_invoice_mapping_is_rejected(accepted_models, value):
    payload = copy.deepcopy(accepted_models)
    payload["analysis"]["models"]["sap"]["payload"]["domination"] = value
    with pytest.raises(ValueError, match="superseded invoice mapping"):
        decode_result(payload)


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("reason", None),
        ("earlier_included", 0),
        ("changed_fields", [1]),
        ("charge_difference", True),
    ],
)
def test_revision_metadata_cannot_change_type_in_a_saved_result(accepted_models, field, value):
    payload = copy.deepcopy(accepted_models)
    payload["analysis"]["models"]["sap"]["payload"]["bill_revisions"]["relationships"][0][field] = (
        value
    )
    with pytest.raises(ValueError, match="stored bill revision"):
        decode_result(payload)


def test_old_results_without_ledger_extensions_do_not_gain_recomputed_sheets(accepted_models):
    payload = copy.deepcopy(accepted_models)
    sap = payload["analysis"]["models"]["sap"]["payload"]
    for field in ("bill_revisions", "rebilling_ledger", "claim_reconciliation", "ledger_overview"):
        del sap[field]
    restored = decode_result(payload)
    view = build_report_view(restored, RenderSelection(("xlsx",)))
    wb = Workbook()
    write_ledger_sheets(wb, view, {})
    assert wb.sheetnames == ["Sheet"]
    assert encode_result(restored, "accepted-integrity") == payload


@pytest.mark.parametrize(
    ("path", "value", "message"),
    [
        (("frames",), {}, "precomputed frames"),
        (("analyses",), {}, "precomputed analyses"),
        (("events",), [None], "event source membership"),
        (("matches",), [None], "event match membership"),
        (("position",), {}, "position payload"),
        (("timeline",), [[1]], "timeline payload"),
        (("union_total",), "100.00", "union total"),
        (("bill_revisions", "version"), True, "Unsupported bill revision"),
    ],
)
def test_corrupt_accepted_sap_components_fail_explicitly(accepted_models, path, value, message):
    payload = copy.deepcopy(accepted_models)
    corrupt(payload["analysis"]["models"]["sap"]["payload"], path, value)
    with pytest.raises(ValueError, match=message):
        decode_result(payload)


def test_duplicate_saved_revision_pair_is_rejected(accepted_models):
    payload = copy.deepcopy(accepted_models)
    links = payload["analysis"]["models"]["sap"]["payload"]["bill_revisions"]["relationships"]
    links.append(copy.deepcopy(links[0]))
    with pytest.raises(ValueError, match="Duplicate or identical revision"):
        decode_result(payload)


def test_accepted_component_workbook_preserves_amounts_and_source_pages(accepted_models):
    view = build_report_view(decode_result(accepted_models), RenderSelection(("xlsx",)))
    wb = Workbook()
    write_ledger_sheets(wb, view, {})
    rows = list(wb["Rebilling components"].iter_rows(min_row=5, values_only=True))
    assert any(row[0] == "OLD" and row[4] == "energy" and row[9] == 100 for row in rows)
    assert any(row[0] == "REPLACEMENT" and row[4] == "vat" and row[9] == 5 for row in rows)


def test_unreadable_component_amount_stays_blank_in_the_workbook():
    accepted = snapshot([bill(**{"Sub Periods": "01/01/2024|31/01/2024|100|100|unreadable"})])
    wb = Workbook()
    write_ledger_sheets(wb, build_report_view(accepted, RenderSelection(("xlsx",))), {})
    assert wb["Rebilling components"]["J5"].value is None
    assert str(accepted.summary.charges) == "100.00"
