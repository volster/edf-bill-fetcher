"""Synthetic regressions for real-output failures in identity, provenance and reports."""

import copy
import hashlib
import json
from datetime import date, datetime, timezone
from zipfile import ZipFile

import pandas as pd
import pdfplumber
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from reportlab.platypus import SimpleDocTemplate

from edf_bill_fetcher.application.analysis_service import analyse_prepared
from edf_bill_fetcher.application.legacy_adapter import capture_legacy
from edf_bill_fetcher.application.preparation_service import prepare_snapshot
from edf_bill_fetcher.collectors.engine import EvidenceEngine
from edf_bill_fetcher.domain.json_values import freeze_json, rows_value
from edf_bill_fetcher.domain.requests import RunRequest
from edf_bill_fetcher.io.persistence.result_codec import decode_result, encode_result
from edf_bill_fetcher.io.reporters.docx_report import _get_or_create_styles
from edf_bill_fetcher.io.reporters.docx_report import create_evidence_index as word_index
from edf_bill_fetcher.io.reporters.pdf_report import create_evidence_index as pdf_index
from edf_bill_fetcher.io.reporters.source_links import source_links
from edf_bill_fetcher.io.writers.meter import write_meter_readings_sheet
from edf_bill_fetcher.models.events import SapBackBillingEvent
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view
from edf_bill_fetcher.processors.detection import detect_back_billing
from edf_bill_fetcher.processors.ledger_context import resolve_ledger_context
from edf_bill_fetcher.processors.matching import match_sap_events_to_edf
from edf_bill_fetcher.processors.meter_reads import sap_reading_basis, sap_reading_observations
from edf_bill_fetcher.processors.preparation import prepare_evidence, scope


@pytest.mark.parametrize("amount", [1234.56, 875.40])
def test_reminder_demand_beats_footer_threshold_without_becoming_another_charge(amount):
    engine = EvidenceEngine({"filter_below": False, "min_amount": 500}, lambda text: None)
    text = (
        f"Account number: 811122233344 We recently sent you a bill for £{amount:,.2f} "
        "but we've not received payment yet. If your bill is estimated, send a reading. "
        "Consider a smart meter. Late charges apply for outstanding balances over £50."
    )
    engine.process_text(text, "Email Body", "Payment reminder", "01/02/2025")
    row = engine.records[0]
    assert row["Balance (£)"] == amount
    assert row["Transaction Amount (£)"] is None
    assert row["Reading"] == "Unknown"
    assert row["Document Role"] == "Payment reminder"
    assert row["Demand Amount (£)"] == amount
    engine.cleanup()


@pytest.mark.parametrize(
    ("sentence", "basis"),
    [
        ("This bill is based on estimated readings.", "Estimated"),
        ("Your bill uses smart meter readings.", "Smart"),
        ("Your bill is based on your actual meter reading.", "Actual"),
    ],
)
def test_specific_email_reading_statement_is_retained(sentence, basis):
    engine = EvidenceEngine({"filter_below": False}, lambda text: None)
    engine.process_text(
        "Current balance £120 in debit. " + sentence, "Email Body", "Bill", "01/02/2025"
    )
    assert engine.records[0]["Reading"] == basis
    engine.cleanup()


def _alias_sources():
    return [
        {
            "SAP Account Number": "811122233344",
            "Kraken ID": "A-45678901",
            "Account": "A-45678901",
            "Source ID": "synthetic-header",
            "Fuel": "Elec",
            "Register": "001",
        }
    ]


def _identity(account="811122233344", **fields):
    return {"Account": account, "Fuel": "Electricity", "Meter": "TEST001", **fields}


def test_alias_and_unambiguous_context_group_without_rewriting_observations():
    records = [_identity(), {"Account": "A-45678901"}, {"Account": "", "Fuel": "Electricity"}]
    before = copy.deepcopy(records)
    resolved, sap = resolve_ledger_context(records, _alias_sources())
    assert records == before
    assert scope(resolved[0]) == scope(resolved[1]) == scope(sap[0])
    assert resolved[1]["Account"] == "A-45678901"
    assert "Meter" not in resolved[1]
    assert scope(resolved[2])[0] == ""
    assert "synthetic-header" in resolved[1]["Ledger Identity Basis"]


@pytest.mark.parametrize("conflict", ["alias", "fuel", "meter"])
def test_conflicting_identity_evidence_does_not_fill_unknown_fields(conflict):
    sources = _alias_sources()
    records = [_identity(), {"Account": "A-45678901"}]
    if conflict == "alias":
        sources.append({**sources[0], "Kraken ID": "A-98765432"})
    elif conflict == "fuel":
        records.append(_identity(Fuel="Gas", Meter="GAS001"))
    else:
        records.append(_identity(Meter="TEST002"))
    resolved, _ = resolve_ledger_context(records, sources)
    if conflict == "alias":
        assert scope(resolved[0])[0] != scope(resolved[1])[0]
    elif conflict == "fuel":
        assert scope(resolved[1])[1] == ""
    else:
        assert scope(resolved[1])[2] == ""


def test_edited_identity_discards_obsolete_derived_context():
    records, _ = resolve_ledger_context([_identity(), {"Account": "A-45678901"}], _alias_sources())
    records[1]["Account"] = "A-77777777"
    resolved, _ = resolve_ledger_context(records, _alias_sources())
    assert scope(resolved[1])[0] == "77777777"
    assert scope(resolved[1])[2] == ""


def _posting(day, source, source_id, **fields):
    return {
        **_identity(),
        "Date": day,
        "Entry Type": "Payment",
        "Amount Kind": "Transaction",
        "Amount (£)": -100,
        "Transaction Amount (£)": -100,
        "Source": source,
        "Source ID": source_id,
        **fields,
    }


def test_payment_repeated_in_single_and_merged_invoice_is_one_event():
    row = _posting(
        "01/01/2025",
        "Local PDF Folder",
        "single",
        **{
            "Source Invoice #": "T90000000001",
            "Document Date": "05/01/2025",
            "Source Row": "payment-history:1",
        },
    )
    prepared = prepare_evidence([row, {**row, "Source ID": "merged"}])
    assert len(prepared.records) == 1
    assert len(prepared.duplicates) == 1
    assert set(prepared.records[0]["Source IDs"]) == {"single", "merged"}


@pytest.mark.parametrize("different", ["Source Invoice #", "Source Row", "Document Date"])
def test_equal_payments_in_different_parent_observations_are_not_merged(different):
    row = _posting(
        "01/01/2025",
        "Local PDF Folder",
        "single",
        **{
            "Source Invoice #": "T90000000001",
            "Document Date": "05/01/2025",
            "Source Row": "payment-history:1",
        },
    )
    assert (
        len(
            prepare_evidence([row, {**row, different: "different", "Source ID": "another"}]).records
        )
        == 2
    )


def _histories(days):
    rows = []
    for i, day in enumerate(days):
        rows.append(
            _posting(
                day,
                "Statement Reconciliation",
                "statement",
                **{
                    "Document Role": "Historical account statement",
                    "Source Row": i + 1,
                    "Invoice #": "STATEMENT",
                },
            )
        )
        rows.append(
            _posting(
                day,
                "HTM Account History",
                "history",
                **{"Posting Date": day, "Balance (£)": 1000 - i * 100, "Source Row": i + 1},
            )
        )
    return rows


def test_unique_postings_need_multiple_dates_to_establish_common_history():
    assert len(prepare_evidence(_histories(["01/01/2025"])).records) == 2
    prepared = prepare_evidence(_histories(["01/01/2025", "02/01/2025"]))
    assert len(prepared.records) == 2
    assert len(prepared.duplicates) == 2
    assert all(len(r["Source IDs"]) == 2 for r in prepared.records)


def test_repeated_equal_event_within_statement_remains_ambiguous():
    rows = _histories(["01/01/2025", "02/01/2025", "03/01/2025"])
    rows.append({**rows[0], "Source Row": 99})
    prepared = prepare_evidence(rows)
    assert len(prepared.records) == 5
    assert sum(bool(r.get("Transaction Identity Review")) for r in prepared.records) == 3


def test_conflicting_running_balances_prevent_sequence_link():
    rows = _histories(["01/01/2025", "02/01/2025", "03/01/2025"])
    rows[0]["Balance (£)"] = 999
    prepared = prepare_evidence(rows)
    assert len(prepared.records) == 4
    assert sum(bool(r.get("Transaction Identity Review")) for r in prepared.records) == 2


def test_matching_amount_cannot_link_events_years_apart():
    event = SapBackBillingEvent("CLEAR-1", pd.Timestamp("2017-01-01"), "Reversal", net_amount=100)
    invoice = {
        "Invoice #": "KI-FICTIONAL",
        "Period From": "01/01/2025",
        "Period To": "31/01/2025",
        "Period Charge (£)": 100,
    }
    assert match_sap_events_to_edf([event], [invoice]) == []
    event.clearing_date = pd.Timestamp("2025-01-15")
    assert match_sap_events_to_edf([event], [invoice])[0].confidence_band == "High"


@pytest.mark.parametrize(
    ("source", "code", "basis"),
    [
        ("Deemed or estimated", "A", "Deemed / estimated"),
        ("Routine", "A", "Unknown"),
        ("Meter reading by customer", "E", "Customer"),
        ("Internet reading", "", "Customer"),
        ("Meter reading by utility", "", "Actual"),
        ("Automatic estimation", "A", "Estimated"),
    ],
)
def test_sap_read_source_takes_precedence_over_raw_code(source, code, basis):
    assert sap_reading_basis({"Read Source": source, "Read Type": code}) == basis


def test_repeated_deemed_opening_reading_is_one_display_observation_with_count():
    row = {
        **_identity(),
        "Meter Read Date": "2017-01-01",
        "Reading (kWh)": "12345",
        "Read Type": "A",
        "Read Source": "Deemed or estimated",
    }
    original = [copy.deepcopy(row) for _ in range(3)]
    combined = sap_reading_observations(original)
    assert len(combined) == 1
    assert combined[0]["Reading Occurrences"] == 3
    assert all(r["Read Type"] == "A" for r in original)
    ws = Workbook().active
    write_meter_readings_sheet(ws, pd.DataFrame(), pd.DataFrame(), sap_readings=original)
    assert ws.cell(8, 3).value == "D"
    assert "3 occurrence(s)" in ws.cell(8, 6).value
    assert ws.cell(9, 2).value is None


def test_positive_net_rebill_uses_explained_gross_replacement_base():
    row = {
        **_identity(),
        "Invoice #": "T90000000001",
        "Date": "01/01/2025",
        "Period From": "01/01/2023",
        "Period To": "01/01/2025",
        "Entry Type": "New Bill",
        "Period Charge (£)": 50,
        "Gross Electricity Charges (£)": 500,
        "Rebilling Components": [{"kind": "cancellation", "amount": -450}],
    }
    result = detect_back_billing(pd.DataFrame([row])).iloc[0]
    assert result["Period Charge (£)"] == 500
    assert result["Unlawful Charge (£)"] == 250.34
    assert "net bill adjustment £50.00" in result["Value Source"]
    assert row["Period Charge (£)"] == 50


def _snapshot():
    request = RunRequest(
        "fixture",
        datetime(2025, 2, 1, tzinfo=timezone.utc),
        date(2025, 2, 1),
        (),
        freeze_json({"as_of": "2025-02-01"}),
    )
    rows = [
        {
            **_identity(),
            "Date": "01/01/2025",
            "Entry Type": "New Bill",
            "Amount (£)": 100,
            "Period Charge (£)": 80,
        },
        {
            "Account": "A-45678901",
            "Date": "01/02/2025",
            "Entry Type": "Ongoing Balance",
            "Amount (£)": 120,
        },
    ]
    return analyse_prepared(
        prepare_snapshot(capture_legacy(rows, {"sap_meter_rows": _alias_sources()}, request))
    )


def test_ledger_observations_survive_strict_roundtrip_and_legacy_absence():
    result = _snapshot()
    assert len(result.prepared.extraction.observations) == 2
    assert rows_value(result.prepared.records)[1]["Account"] == "A-45678901"
    payload = encode_result(result, "fixture-result")
    view = build_report_view(decode_result(payload), RenderSelection())
    assert len(view.ledger_overview) == 1
    assert view.ledger_overview[0]["closing_balance"] == "120.00"
    assert view.ledger_overview[0]["references"] == ["811122233344", "A-45678901"]
    del payload["analysis"]["models"]["sap"]["payload"]["ledger_overview"]
    assert build_report_view(decode_result(payload), RenderSelection()).ledger_overview == []


def test_ledger_amount_and_date_validation_cannot_be_bypassed():
    payload = encode_result(_snapshot(), "fixture-result")
    payload["analysis"]["models"]["sap"]["payload"]["ledger_overview"][0]["closing_date"] = None
    with pytest.raises(ValueError, match="dated observation"):
        decode_result(payload)


def test_word_and_pdf_source_links_use_verified_relative_manifest_targets(tmp_path):
    source = tmp_path / "evidence_files" / "fictional original.pdf"
    source.parent.mkdir()
    source.write_bytes(b"synthetic-source-bytes")
    manifest = {
        "files": [
            {
                "status": "saved",
                "source_id": "fixture-source",
                "attachment_name": "original.pdf",
                "saved_path": "evidence_files/fictional original.pdf",
                "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            }
        ]
    }
    (tmp_path / "evidence_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    frame = pd.DataFrame(
        [
            {
                **_identity(),
                "Source ID": "fixture-source",
                "Source": "Local PDF Folder",
                "Date": "01/01/2025",
                "Invoice #": "INV-TEST",
                "Attachment Name": "original.pdf",
                "Page From": 2,
            }
        ]
    )
    document = Document()
    word_index(document, _get_or_create_styles(document), frame, None, bundle_dir=str(tmp_path))
    path = tmp_path / "report.docx"
    document.save(path)
    with ZipFile(path) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        assert 'Target="evidence_files/fictional%20original.pdf#page=2"' in rels
    path = tmp_path / "report.pdf"
    SimpleDocTemplate(str(path)).build(pdf_index(frame, None, bundle_dir=str(tmp_path)))
    with pdfplumber.open(path) as pdf:
        assert any(
            link["uri"] == "evidence_files/fictional%20original.pdf#page=2"
            for page in pdf.pages
            for link in page.hyperlinks
        )
    source.write_bytes(b"changed bytes")
    assert source_links(str(tmp_path)) == {}


def test_posting_inside_long_invoice_span_has_zero_period_distance():
    event = SapBackBillingEvent(
        "CLEAR-1",
        pd.Timestamp("2025-02-01"),
        "Reversal",
        rows=[{"Posting Date": "2022-01-15"}],
        net_amount=100,
    )
    invoice = {
        "Invoice #": "T-TEST",
        "Period From": "01/01/2022",
        "Period To": "31/12/2023",
        "Period Charge (£)": 100,
    }
    result = match_sap_events_to_edf([event], [invoice])
    assert result[0].confidence_band == "High"
    assert result[0].date_delta_days == 0
    assert "Posting date inside" in result[0].notes


def test_nonmonetary_nan_rows_do_not_create_financial_ledgers():
    from edf_bill_fetcher.processors.ledger_overview import ledger_overview

    frame = pd.DataFrame(
        [
            _posting("01/01/2025", "Local PDF Folder", "source"),
            {"Entry Type": "Supporting Document", "Amount Kind": "Not monetary"},
        ]
    )
    rows = ledger_overview(frame)
    assert len(rows) == 1
    assert rows[0]["records"] == 1


@pytest.mark.parametrize(
    "threshold", ["outstanding balances over £50", "a balance above £50", "your balance below £50"]
)
def test_fee_threshold_without_a_stated_demand_is_not_an_observed_balance(threshold):
    engine = EvidenceEngine({"filter_below": False, "use_large": False}, lambda text: None)
    engine.process_text("Late charges apply for " + threshold, "Email Body", "Advice", "01/02/2025")
    assert engine.records == []
    engine.cleanup()


def test_conditional_current_bill_reading_language_does_not_assert_a_basis():
    engine = EvidenceEngine({"filter_below": False}, lambda text: None)
    engine.process_text(
        "Current balance £120 in debit. If your bill is based on estimated readings, send a reading.",
        "Email Body",
        "Advice",
        "01/02/2025",
    )
    assert engine.records[0]["Reading"] == "Unknown"
    engine.cleanup()


def test_single_matching_posting_is_visible_as_unresolved_work():
    prepared = prepare_evidence(_histories(["01/01/2025"]))
    assert len(prepared.records) == 2
    assert all(row["Transaction Identity Review"] for row in prepared.records)


@pytest.mark.parametrize("second_account", [False, True])
def test_annual_summary_ignores_nonmonetary_sources_but_keeps_ledger_boundary(
    tmp_path, second_account
):
    from edf_bill_fetcher.io.writers.export import export_to_excel

    rows = [
        {
            **_identity(),
            "Date": "01/01/2025",
            "Entry Type": "Ongoing Balance",
            "Amount (£)": 120,
        },
        {
            "Date": "01/01/2024",
            "Entry Type": "Supporting Document",
            "Amount Kind": "Not monetary",
        },
    ]
    if second_account:
        rows.append({**rows[0], "Account": "999988887777", "Amount (£)": 200})
    output = tmp_path / "summary.xlsx"
    export_to_excel(rows, str(output), [], {"analysis_min": 0, "use_dedup": True})
    workbook = load_workbook(output)
    sheet = workbook["Annual Summary"]
    if second_account:
        assert "multiple ledgers" in sheet["A1"].value
    else:
        assert sheet["A1"].value == "Year"
        assert sheet["A2"].value == 2025
        assert sheet["A3"].value == "OVERALL"
        assert sheet.max_row == 3
    workbook.close()
