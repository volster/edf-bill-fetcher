"""Desktop/Excel parity and portable complaint artifacts with independent expected amounts."""

import copy
import hashlib
import json
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pdfplumber
import pytest
from docx import Document
from openpyxl import load_workbook
from reportlab.pdfgen.canvas import Canvas

from edf_bill_fetcher.application import complaint_pack
from edf_bill_fetcher.application.analysis_service import analyse_prepared
from edf_bill_fetcher.application.legacy_adapter import capture_legacy
from edf_bill_fetcher.application.preparation_service import prepare_snapshot
from edf_bill_fetcher.application.review_session import ReviewSession
from edf_bill_fetcher.application.run_service import RunCancelled, run
from edf_bill_fetcher.application.workbook_review import regenerate_workbook
from edf_bill_fetcher.domain.json_values import freeze_json, object_value
from edf_bill_fetcher.io.claim_workbook import CREDIT_SHEET, GROUP_SHEET, OVERLAP_SHEET
from edf_bill_fetcher.io.dispute_workbook import HEADERS, SHEET
from edf_bill_fetcher.io.persistence.result_codec import decode_result, encode_result
from edf_bill_fetcher.io.review_workbook import load_review_workbook, read_workbook
from edf_bill_fetcher.io.writers.export import render_xlsx
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view
from tests.test_claim_reconciliation import credit, evidence
from tests.test_native_review import analyse_session
from tests.test_qt_desktop import result_snapshot
from tests.test_rebilling_ledger import set_position


def rid(value):
    return hashlib.sha256(("review-" + value).encode()).hexdigest()


def set_claim(session, key="C1", members=("a", "b"), **changes):
    session.set_claim(
        key,
        title="Historic additional charges",
        status="reviewed",
        members=[rid(k) for k in members],
        challenged="120",
        reduction="120",
        credits=[
            credit("70", row_id=rid("c")),
            credit("30", "already_applied_relief", row_id=rid("c")),
        ],
        reason="One challenge to the same economic charges, verified against the entire earlier billing history.",
        account_link_reason="",
        **changes,
    )


@pytest.fixture
def reviewed(tmp_path):
    root = tmp_path / "original"
    root.mkdir()
    items, _ = evidence()
    # Keep an even higher estimate in the earlier history. It must survive curation.
    earlier = copy.deepcopy(items[0])
    earlier["id"] = "prior"
    earlier["record"].update(
        {
            "Invoice #": "HIGHER-ESTIMATE",
            "Date": "01/01/2024",
            "Period Charge (£)": 400,
            "Transaction Amount (£)": 400,
        }
    )
    items.append(earlier)
    unrelated = copy.deepcopy(items[0])
    unrelated["id"] = "unrelated"
    unrelated["record"].update(
        {"Invoice #": "UNRELATED-ACCOUNT", "Account": "OTHER-ACCOUNT", "Meter": "OTHER-METER"}
    )
    items.append(unrelated)
    manifest = {}
    for item in items:
        source = root / "evidence_files" / (item["id"] + ".pdf")
        source.parent.mkdir(exist_ok=True)
        pdf = Canvas(str(source))
        pdf.drawString(40, 700, item["record"]["Invoice #"])
        pdf.save()
        key = hashlib.sha256(item["id"].encode()).hexdigest()
        metadata = {
            "Source ID": key,
            "Content Hash": hashlib.sha256(source.read_bytes()).hexdigest(),
            "Attachment Name": source.name,
            "Original Filename": source.name,
            "Saved Path": "evidence_files/" + source.name,
            "Original Path": str(source),
        }
        manifest[key] = metadata
        item["record"].update(
            {**metadata, "Review Row ID": rid(item["id"]), "Page From": 1, "Page To": 1}
        )
    request = result_snapshot().prepared.extraction.request
    snapshot = analyse_prepared(
        prepare_snapshot(
            capture_legacy(
                [i["record"] for i in items],
                {"source_manifest": manifest, "source_bundle_path": str(root)},
                request,
            )
        )
    )
    session = ReviewSession(snapshot)
    for key in ("a", "b"):
        set_position(session, rid(key), first_demand_date="01/02/2024")
    set_claim(session)
    session.set_complaint_context(
        "Additional historic charges are disputed <script>unsafe()</script>.",
        "Cancel the protected charges and substantiate the balance.",
        True,
    )
    snapshot = analyse_session(session, root / "source.json")
    return root, snapshot


def view(snapshot):
    return build_report_view(snapshot, RenderSelection(("xlsx",), ("detailed_findings",)))


def import_snapshot(path, baseline):
    imported = load_review_workbook(str(path))
    assert imported is not None
    records, state, _ = imported
    request = replace(
        baseline.prepared.extraction.request, config=freeze_json(state["review_config"])
    )
    return analyse_prepared(prepare_snapshot(capture_legacy(records, state, request)))


def workbook(snapshot, root):
    path = root / "Review.xlsx"
    render_xlsx(view(snapshot), str(path))
    return path


def edit(path, action):
    wb = load_workbook(path)
    action(wb)
    wb.save(path)
    wb.close()


def test_unchanged_claim_workbook_and_saved_result_preserve_every_decision(reviewed):
    root, snapshot = reviewed
    path = workbook(snapshot, root)
    imported = import_snapshot(path, snapshot)
    assert view(imported).claims == view(snapshot).claims
    assert imported.input_fingerprint == snapshot.input_fingerprint
    decoded = decode_result(encode_result(snapshot, result_id="claim-test"))
    assert view(decoded).claims == view(snapshot).claims
    wb = load_workbook(path)
    assert wb[GROUP_SHEET]["D5"].protection.locked is False
    assert wb[CREDIT_SHEET]["H5"].protection.locked is True
    assert wb[CREDIT_SHEET].column_dimensions["P"].hidden
    assert "90.00" in " ".join(str(c.value) for row in wb["Reconciled claims"] for c in row)
    wb.close()


def test_excel_and_native_change_one_consolidated_proposal_identically(reviewed):
    root, snapshot = reviewed
    path = workbook(snapshot, root)

    def change(wb):
        wb[GROUP_SHEET]["D5"] = 140
        wb[GROUP_SHEET]["E5"] = 130
        wb[CREDIT_SHEET]["C5"] = 20  # sorted: already-applied relief before cancellation offset
        wb[CREDIT_SHEET]["C6"] = 80

    edit(path, change)
    excel = import_snapshot(path, snapshot)
    native = ReviewSession(snapshot)
    data = {k: v for k, v in native.claims["C1"].items() if k not in {"id", "source_digest"}}
    data.update(
        challenged="140",
        reduction="130",
        credits=[
            credit("80", row_id=rid("c")),
            credit("20", "already_applied_relief", row_id=rid("c")),
        ],
    )
    native.set_claim("C1", **data)
    accepted = analyse_session(native, path)
    assert view(excel).claims == view(accepted).claims
    assert view(excel).claims["proposed_reduction"] == "110.00"
    assert excel.summary == snapshot.summary


@pytest.mark.parametrize(
    "change",
    [
        "formula",
        "unknown-credit",
        "catalogue",
        "unknown-group",
        "orphan",
        "overflow",
        "missing-sheet",
        "duplicate-group",
    ],
)
def test_invalid_claim_workbooks_fail_before_a_run_folder_is_created(reviewed, tmp_path, change):
    root, snapshot = reviewed
    path = workbook(snapshot, root)

    def alter(wb):
        if change == "formula":
            wb[GROUP_SHEET]["E5"] = "=120"
        elif change == "unknown-credit":
            wb[CREDIT_SHEET]["B5"] = "CRED999"
        elif change == "catalogue":
            wb[CREDIT_SHEET]["I5"] = "DIFFERENT SOURCE"
        elif change == "unknown-group":
            wb[SHEET].cell(5, HEADERS.index("Claim groups") + 1, "C999")
        elif change == "orphan":
            wb[CREDIT_SHEET]["A5"] = "MISSING"
        elif change == "overflow":
            wb[GROUP_SHEET]["A1005"] = "LOST"
        elif change == "missing-sheet":
            del wb[OVERLAP_SHEET]
        elif change == "duplicate-group":
            for column in range(1, 8):
                wb[GROUP_SHEET].cell(6, column, wb[GROUP_SHEET].cell(5, column).value)

    edit(path, alter)
    destination = tmp_path / "new-runs"
    destination.mkdir()
    with pytest.raises(ValueError):
        regenerate_workbook(str(path), str(destination))
    assert not list(destination.iterdir())


def test_unchanged_claim_decision_becomes_stale_when_source_record_is_corrected(reviewed):
    from tests.test_excel_review import amend

    root, snapshot = reviewed
    path = workbook(snapshot, root)
    amend(
        path, "REPLACEMENT", **{"Period Charge (£)": 180, "Reason for change": "Verified original"}
    )
    imported = import_snapshot(path, snapshot)
    claims = view(imported).claims
    assert not claims["ready"]
    assert claims["proposed_reduction"] is None
    assert claims["groups"][0]["needs_review"]
    assert claims["input"] == view(snapshot).claims["input"]


def test_focused_pack_preserves_higher_history_exact_sources_and_one_adjustment(reviewed, tmp_path):
    root, snapshot = reviewed
    destination = tmp_path / "run" / "complaint_pack"
    artifacts, notes = complaint_pack.publish(view(snapshot), str(destination))
    assert not notes
    assert set(artifacts) == {
        "complaint_zip",
        "complaint_html",
        "complaint_pdf",
        "complaint_docx",
        "complaint_schedule",
        "complaint_evidence_manifest",
    }
    bundle = json.loads((destination / "evidence_manifest.json").read_text())
    invoices = {inv for source in bundle["files"] for inv in source["invoices"]}
    assert {"OLD", "REPLACEMENT", "CREDIT", "HIGHER-ESTIMATE"} <= invoices
    assert "UNRELATED-ACCOUNT" not in invoices
    assert "NEXT" not in invoices
    for source in bundle["files"]:
        path = destination / source["saved_path"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == source["sha256"]
        assert source["page_ranges"] == [[1, 1]]
        assert not Path(source["saved_path"]).is_absolute()
    html = (destination / "Complaint_Draft.html").read_text()
    assert "&lt;script&gt;" in html
    assert "<script>" not in html
    assert "#page=1" in html
    assert str(root) not in html
    assert "UNRELATED-ACCOUNT" not in html
    doc = Document(destination / "Complaint_Draft.docx")
    doc_text = " ".join(p.text for p in doc.paragraphs) + " ".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells
    )
    with pdfplumber.open(destination / "Complaint_Draft.pdf") as pdf:
        pdf_text = " ".join(p.extract_text() or "" for p in pdf.pages)
    for text in (html, doc_text, pdf_text):
        assert "90.00" in text
        assert "HIGHER-ESTIMATE" in text
        assert "Cancellation" in text or "cancellation" in text
        assert "current liability" in text
    with ZipFile(artifacts["complaint_zip"]) as archive:
        assert "result.json" not in archive.namelist()
        assert set(archive.namelist()) == {
            p.relative_to(destination).as_posix() for p in destination.rglob("*") if p.is_file()
        }
    with pytest.raises(ValueError, match="reference output"):
        read_workbook(str(destination / "Claim_Schedule.xlsx"))
    with pytest.raises(FileExistsError):
        complaint_pack.publish(view(snapshot), str(destination))


@pytest.mark.parametrize("failure", ["missing", "changed", "hashless"])
def test_missing_or_changed_saved_evidence_never_falls_back_to_an_original(
    reviewed, tmp_path, failure
):
    root, snapshot = reviewed
    state = object_value(snapshot.prepared.extraction.legacy_state)
    source_id = next(
        key
        for key, data in state["source_manifest"].items()
        if data["Original Filename"] == "a.pdf"
    )
    good_original = tmp_path / "original-a.pdf"
    good_original.write_bytes((root / "evidence_files/a.pdf").read_bytes())
    state["source_manifest"][source_id]["Original Path"] = str(good_original)
    if failure == "missing":
        (root / "evidence_files/a.pdf").unlink()
    elif failure == "changed":
        (root / "evidence_files/a.pdf").write_bytes(b"changed source")
    else:
        state["source_manifest"][source_id].pop("Content Hash")
    if failure == "hashless":
        # No authoritative source/row hash may be invented from today's bytes.
        raw = [object_value(o.raw) for o in snapshot.prepared.extraction.observations]
        for row in raw:
            if row.get("Source ID") == source_id:
                row.pop("Content Hash", None)
        state.pop("original_records", None)
        state.pop("human_review", None)
        base = analyse_prepared(
            prepare_snapshot(capture_legacy(raw, state, snapshot.prepared.extraction.request))
        )
    else:
        base = replace(
            snapshot,
            prepared=replace(
                snapshot.prepared,
                extraction=replace(snapshot.prepared.extraction, legacy_state=freeze_json(state)),
            ),
        )
    artifacts, notes = complaint_pack.publish(view(base), str(tmp_path / "complaint_pack"))
    assert any("unavailable" in n or "inconsistent" in n for n in notes)
    bundle = json.loads(Path(artifacts["complaint_evidence_manifest"]).read_text())
    assert (
        next(s for s in bundle["files"] if s["source_id"] == source_id)["status"] == "unavailable"
    )


def test_pack_cancellation_and_renderer_failure_leave_no_partial_pack(
    reviewed, tmp_path, monkeypatch
):
    _, snapshot = reviewed
    target = tmp_path / "cancelled" / "complaint_pack"
    with pytest.raises(RunCancelled):
        complaint_pack.publish(view(snapshot), str(target), cancelled=lambda: True)
    assert not target.exists()
    assert not target.with_name("Complaint_pack.zip").exists()
    assert not list(target.parent.iterdir())
    monkeypatch.setattr(
        complaint_pack,
        "_write_documents",
        lambda *args: (_ for _ in ()).throw(OSError("disk full")),
    )
    with pytest.raises(OSError, match="disk full"):
        complaint_pack.publish(view(snapshot), str(target))
    assert not list(target.parent.iterdir())


def test_run_manifest_records_pack_artifacts_gaps_and_keeps_analysis_after_failure(
    reviewed, tmp_path, monkeypatch
):
    root, snapshot = reviewed
    (root / "evidence_files/a.pdf").unlink()
    target = tmp_path / "run"
    target.mkdir()
    outcome = run(
        snapshot.prepared.extraction.request,
        {},
        result_path=str(target / "result.json"),
        manifest_path=str(target / "run_manifest.json"),
        accepted_result=snapshot,
        selection=RenderSelection(()),
    )
    assert outcome.status == "review_required", outcome.error
    manifest = object_value(outcome.manifest)
    assert manifest["artifacts"]["complaint_zip"]["sha256"]
    assert manifest["warnings"]
    target2 = tmp_path / "failed"
    target2.mkdir()
    monkeypatch.setattr(
        complaint_pack,
        "publish",
        lambda *args, **kwargs: (_ for _ in ()).throw(OSError("pack write failure")),
    )
    failed = run(
        snapshot.prepared.extraction.request,
        {},
        result_path=str(target2 / "result.json"),
        manifest_path=str(target2 / "run_manifest.json"),
        accepted_result=snapshot,
        selection=RenderSelection(()),
    )
    assert failed.status == "failed"
    assert failed.snapshot is snapshot
    assert (target2 / "result.json").is_file()


def test_excel_can_create_new_groups_memberships_and_credit_allocations_from_a_blank_review(
    reviewed,
):
    root, original = reviewed
    native = ReviewSession(original)
    native.clear_claim("C1")
    blank = analyse_session(native, root / "blank.json")
    path = workbook(blank, root)

    def fill(wb):
        group_values = (
            "C1",
            "Historic additional charges",
            "reviewed",
            120,
            120,
            "One challenge to the same economic charges, verified against the entire earlier billing history.",
            "",
        )
        for column, value in enumerate(group_values, 1):
            wb[GROUP_SHEET].cell(5, column, value)
        for row in range(5, wb[SHEET].max_row + 1):
            key = wb[SHEET].cell(row, len(HEADERS)).value
            if key in {rid("a"), rid("b")}:
                wb[SHEET].cell(row, HEADERS.index("Claim groups") + 1, "C1")
        for row, values in enumerate(
            (
                ("C1", "CRED001", 70, "cancellation_offset", credit("70")["reason"]),
                ("C1", "CRED001", 30, "already_applied_relief", credit("30")["reason"]),
            ),
            5,
        ):
            for column, value in enumerate(values, 1):
                wb[CREDIT_SHEET].cell(row, column, value)

    edit(path, fill)
    imported = import_snapshot(path, blank)
    assert view(imported).claims == view(original).claims
    assert view(imported).claims["proposed_reduction"] == "90.00"


def test_excel_overlap_decisions_bind_both_groups_and_become_stale_after_changes(reviewed):
    root, snapshot = reviewed
    native = ReviewSession(snapshot)
    fields = {k: v for k, v in native.claims["C1"].items() if k not in {"id", "source_digest"}}
    fields.update(members=[rid("a")], credits=[])
    native.set_claim("C1", **fields)
    fields.update(members=[rid("b")])
    native.set_claim("C2", **fields)
    baseline = analyse_session(native, root / "overlap.json")
    path = workbook(baseline, root)
    assert not view(baseline).claims["ready"]

    def separate(wb):
        wb[OVERLAP_SHEET]["C5"] = "separate_charges"
        wb[OVERLAP_SHEET]["D5"] = "Distinct charge lines separately identified on the originals."

    edit(path, separate)
    imported = import_snapshot(path, baseline)
    assert view(imported).claims["ready"]
    assert view(imported).claims["proposed_reduction"] == "240.00"
    path = workbook(imported, root)
    edit(path, lambda wb: setattr(wb[GROUP_SHEET]["E5"], "value", 110))
    stale = import_snapshot(path, imported)
    assert not view(stale).claims["ready"]
    assert view(stale).claims["overlaps"][0]["needs_review"]


def test_draft_incomplete_and_opted_out_cases_never_assert_a_combined_total(reviewed, tmp_path):
    root, snapshot = reviewed
    native = ReviewSession(snapshot)
    native.clear_claim("C1")
    incomplete = analyse_session(native, root / "incomplete.json")
    artifacts, notes = complaint_pack.publish(
        view(incomplete), str(tmp_path / "incomplete" / "complaint_pack")
    )
    assert any("reconciliation is incomplete" in note for note in notes)
    text = Path(artifacts["complaint_html"]).read_text()
    assert "combined reduction is not yet asserted" in text
    assert "WITHHELD" in text
    assert "90.00" not in text
    native.set_complaint_context("", "", False)
    opted_out = analyse_session(native, root / "opted-out.json")
    target = tmp_path / "disabled" / "complaint_pack"
    assert complaint_pack.publish(view(opted_out), str(target)) == ({}, ())
    assert not target.parent.exists()


@pytest.mark.parametrize(
    "corruption",
    [
        "scope",
        "ledger-scope",
        "group-ready",
        "unresolved-total",
        "unknown-member",
        "duplicate-ledger-group",
        "bad-money",
    ],
)
def test_corrupt_stored_claims_fail_closed_without_recomputing_accepted_numbers(
    reviewed, corruption, monkeypatch
):
    from edf_bill_fetcher.io.persistence.result_codec import _validate_claims
    from edf_bill_fetcher.processors import claim_reconciliation

    _, snapshot = reviewed
    accepted = view(snapshot)
    model = copy.deepcopy(accepted.claims)
    if corruption == "scope":
        model["groups"][0]["scope"]["accounts"] = "A"
    elif corruption == "ledger-scope":
        model["ledgers"][0]["scope"] = {}
    elif corruption == "group-ready":
        model["groups"][0]["needs_review"] = True
    elif corruption == "unresolved-total":
        model["unresolved_positions"] = [rid("a")]
    elif corruption == "unknown-member":
        model["input"]["groups"][0]["members"].append("missing")
    elif corruption == "duplicate-ledger-group":
        model["ledgers"][0]["groups"].append("C1")
    elif corruption == "bad-money":
        model["groups"][0]["further_reduction"] = float("nan")
    monkeypatch.setattr(
        claim_reconciliation,
        "build_reconciliation",
        lambda *args: pytest.fail("Accepted results must not recalculate"),
    )
    with pytest.raises(ValueError):
        _validate_claims(model, accepted.rebilling_ledger)
    _validate_claims(accepted.claims, accepted.rebilling_ledger)


def test_excluded_group_members_and_withdrawn_details_stay_out_of_focused_claim_tables(
    reviewed, tmp_path
):
    root, snapshot = reviewed
    session = ReviewSession(snapshot)
    session.update(
        rid("a"), {}, False, "Unrelated source excluded from this complaint", "needs_information"
    )
    excluded = analyse_session(session, root / "excluded.json")
    artifacts, notes = complaint_pack.publish(view(excluded), str(tmp_path / "complaint_pack"))
    assert any("excluded or inactive members" in note for note in notes)
    bundle = json.loads(Path(artifacts["complaint_evidence_manifest"]).read_text())
    assert "OLD" not in {inv for source in bundle["files"] for inv in source["invoices"]}
    text = Path(artifacts["complaint_html"]).read_text()
    assert "90.00" not in text
    assert "One challenge to the same economic charges" not in text


def test_claims_on_absorbed_duplicate_records_are_archived_without_transferring_amounts(reviewed):
    root, snapshot = reviewed
    session = ReviewSession(snapshot)
    before = session.claim_state()["groups"][0]
    session.update(
        rid("b"),
        {
            "Invoice #": "OLD",
            "Date": "01/02/2024",
            "Period Charge (£)": "100",
            "Balance (£)": "250",
        },
        True,
        "Same underlying bill; verified original identity",
        "verified",
    )
    merged = analyse_session(session, root / "merged.json")
    state = object_value(merged.prepared.extraction.legacy_state)
    assert state["archived_claim_groups"] == [before]
    assert not state["claim_reconciliation"]["groups"]
    assert view(merged).claims["proposed_reduction"] is None
    assert any(
        e["field"] == "Claim reconciliation" and e["after"] == "archived_after_deduplication"
        for e in state["review_history"]
    )


def test_main_outputs_and_focused_pack_share_the_accepted_reconciled_amount(reviewed, tmp_path):
    _, snapshot = reviewed
    destination = tmp_path / "complete-reviewed-run"
    destination.mkdir()
    outputs = {fmt: str(destination / ("Report." + fmt)) for fmt in ("xlsx", "html", "docx", "pdf")}
    outcome = run(
        snapshot.prepared.extraction.request,
        outputs,
        result_path=str(destination / "result.json"),
        manifest_path=str(destination / "run_manifest.json"),
        accepted_result=snapshot,
        selection=RenderSelection(tuple(outputs), ("detailed_findings",)),
    )
    assert outcome.status == "complete", outcome.error
    doc = Document(outputs["docx"])
    doc_text = " ".join(p.text for p in doc.paragraphs) + " ".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells
    )
    with pdfplumber.open(outputs["pdf"]) as pdf:
        pdf_text = " ".join(p.extract_text() or "" for p in pdf.pages)
    for content in (Path(outputs["html"]).read_text(), doc_text, pdf_text):
        assert "Consolidated claim proposal" in content
        assert "90.00" in content
        assert "Cancellation offsets" in content
        assert "Historical account" not in content
    model = view(decode_result(json.loads((destination / "result.json").read_text()))).claims
    assert model == view(snapshot).claims
    assert model["proposed_reduction"] == "90.00"


@pytest.mark.parametrize("kind", ["excess-credit", "unreviewed-member"])
def test_new_reviewed_excel_inputs_follow_the_native_evidence_rules(reviewed, tmp_path, kind):
    root, snapshot = reviewed
    path = workbook(snapshot, root)

    def invalid(wb):
        if kind == "excess-credit":
            wb[CREDIT_SHEET]["C5"] = 80
        else:
            for row in range(5, wb[SHEET].max_row + 1):
                if wb[SHEET].cell(row, len(HEADERS)).value == rid("d"):
                    wb[SHEET].cell(row, HEADERS.index("Claim groups") + 1, "C1")

    edit(path, invalid)
    destination = tmp_path / "invalid-financial-review"
    destination.mkdir()
    with pytest.raises(ValueError, match="cannot be marked reviewed"):
        regenerate_workbook(str(path), str(destination))
    assert not list(destination.iterdir())


def test_disabling_focused_pack_does_not_hide_unresolved_claim_status(reviewed, tmp_path):
    root, snapshot = reviewed
    native = ReviewSession(snapshot)
    native.clear_claim("C1")
    native.set_complaint_context("", "", False)
    incomplete = analyse_session(native, root / "without-focused-pack.json")
    destination = tmp_path / "report-only"
    destination.mkdir()
    outcome = run(
        incomplete.prepared.extraction.request,
        {},
        result_path=str(destination / "result.json"),
        manifest_path=str(destination / "run_manifest.json"),
        accepted_result=incomplete,
        selection=RenderSelection(()),
    )
    assert outcome.status == "review_required", outcome.error
    assert object_value(outcome.manifest)["warnings"]
    assert not (destination / "complaint_pack").exists()


def test_withdrawn_history_is_retained_without_false_outstanding_work_or_a_complaint_pack(
    reviewed, tmp_path
):
    from edf_bill_fetcher.models.claim_tables import group_table, total_table
    from edf_bill_fetcher.models.output_overview import build_output_overview
    from edf_bill_fetcher.processors.claim_reconciliation import has_active_claims

    root, snapshot = reviewed
    session = ReviewSession(snapshot)
    fields = {k: v for k, v in session.claims["C1"].items() if k not in {"id", "source_digest"}}
    fields["status"] = "withdrawn"
    session.set_claim("C1", **fields)
    # Withdrawing a group does not silently withdraw its individual positions.
    assert has_active_claims(session.claim_model())
    for key in ("a", "b"):
        set_position(session, rid(key), status="withdrawn", first_demand_date="01/02/2024")
    withdrawn = analyse_session(session, root / "withdrawn.json")
    accepted = view(withdrawn)
    assert not has_active_claims(accepted.claims)
    assert accepted.claims["input"]["groups"][0]["status"] == "withdrawn"
    assert "No active claims" in total_table(accepted.claims)[1]
    assert "Withdrawn" in group_table(accepted.claims, accepted.rebilling_ledger)[1][0]
    assert all(p.sheet != "Reconciled claims" for p in build_output_overview(accepted).priorities)
    destination = tmp_path / "withdrawn-run"
    destination.mkdir()
    outcome = run(
        withdrawn.prepared.extraction.request,
        {},
        result_path=str(destination / "result.json"),
        manifest_path=str(destination / "run_manifest.json"),
        accepted_result=withdrawn,
        selection=RenderSelection(()),
    )
    assert outcome.status == "complete", outcome.error
    assert not object_value(outcome.manifest)["warnings"]
    assert not (destination / "complaint_pack").exists()
    assert withdrawn.summary == snapshot.summary
