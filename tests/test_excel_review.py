"""Human Excel corrections must rebuild financial analysis and preserve evidence."""

from __future__ import annotations

import hashlib
import json
import shutil
from dataclasses import replace
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

from edf_bill_fetcher.application.analysis_service import analyse_prepared
from edf_bill_fetcher.application.legacy_adapter import capture_legacy
from edf_bill_fetcher.application.preparation_service import prepare_snapshot
from edf_bill_fetcher.application.workbook_review import regenerate_workbook
from edf_bill_fetcher.domain.json_values import freeze_json, object_value, rows_value
from edf_bill_fetcher.domain.requests import RunRequest
from edf_bill_fetcher.io.persistence.result_store import load_result
from edf_bill_fetcher.io.review_workbook import (
    HEADERS,
    PAYLOAD_SHEET,
    REVIEW_SHEET,
    read_workbook,
)
from edf_bill_fetcher.io.review_workbook import (
    load_review_workbook as _load_review_workbook,
)
from edf_bill_fetcher.io.writers.export import render_xlsx
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view


def load_review_workbook(path: str) -> tuple[list[dict], dict, str]:
    """Require this test's newly generated workbook to use the lossless format."""
    imported = _load_review_workbook(path)
    assert imported is not None
    return imported


@pytest.fixture
def review_pack(tmp_path):
    root = tmp_path / "original"
    sources = root / "evidence_files"
    sources.mkdir(parents=True)
    manifest = {}
    ids = []
    for name in ("shared.pdf", "excluded.pdf", "duplicate.pdf"):
        path = sources / name
        path.write_bytes(b"Original source bytes: " + name.encode())
        identity = hashlib.sha256(name.encode()).hexdigest()
        ids.append(identity)
        manifest[identity] = {
            "Source ID": identity,
            "Source": "Local PDF",
            "Attachment Name": name,
            "Original Filename": name,
            "Content Hash": hashlib.sha256(path.read_bytes()).hexdigest(),
            "Saved Path": "evidence_files/" + name,
        }
    records = []
    for ordinal, (invoice, charge, balance, source) in enumerate(
        (("INV-A", 100, 250, ids[0]), ("INV-B", 50, 300, ids[0]), ("INV-C", 70, 370, ids[1]))
    ):
        records.append(
            {
                **manifest[source],
                "Date": f"0{ordinal + 1}/02/2024",
                "Entry Type": "New Bill",
                "Invoice #": invoice,
                "Period From": "01/01/2024",
                "Period To": "31/01/2024",
                "Period Charge (£)": charge,
                "Amount (£)": balance,
                "Account": "000123",
                "Fuel": "Electricity",
                "Page From": ordinal + 1 if source == ids[0] else 1,
                "Page To": ordinal + 1 if source == ids[0] else 1,
                "Source PDF Text": "Unaltered original extraction text",
            }
        )
    records.append({**records[0], **manifest[ids[2]], "Source": "PST PDF Attachment"})
    config = {"as_of": "2026-01-31", "analysis_min": 0, "use_dedup": True}
    request = RunRequest(
        "original-review",
        datetime(2026, 1, 31, tzinfo=timezone.utc),
        date(2026, 1, 31),
        (),
        freeze_json(config),
    )
    snapshot = analyse_prepared(
        prepare_snapshot(
            capture_legacy(
                records, {"source_manifest": manifest, "source_bundle_path": str(root)}, request
            )
        )
    )
    path = root / "Evidence.xlsx"
    render_xlsx(
        build_report_view(snapshot, RenderSelection(("xlsx",))),
        str(path),
        evidence_paths={
            identity: str(root / metadata["Saved Path"]) for identity, metadata in manifest.items()
        },
    )
    return path, snapshot, ids


def amend(path, invoice, **values):
    """Edit cells as a human would, without altering hidden provenance."""
    wb = load_workbook(path)
    ws = wb[REVIEW_SHEET]
    for row in range(5, ws.max_row + 1):
        if ws.cell(row, HEADERS.index("Invoice #") + 1).value == invoice:
            for field, value in values.items():
                ws.cell(row, HEADERS.index(field) + 1).value = value
            break
    else:
        raise AssertionError(f"Missing invoice {invoice}")
    wb.save(path)
    wb.close()


def test_workbook_is_useful_and_unchanged_roundtrip_is_lossless(review_pack):
    path, baseline, _ = review_pack
    wb = load_workbook(path)
    assert wb.sheetnames[:3] == ["Start here", REVIEW_SHEET, "Review history"]
    assert wb[PAYLOAD_SHEET].sheet_state == "veryHidden"
    ws = wb[REVIEW_SHEET]
    assert ws.freeze_panes == "F5"
    assert ws.tables
    assert ws.data_validations.count >= 3
    assert ws["F5"].protection.locked is False
    assert ws.cell(5, len(HEADERS)).protection.locked is True
    assert ws.cell(5, HEADERS.index("Account") + 1).number_format == "@"
    assert ws.cell(5, HEADERS.index("Account") + 1).value == "000123"
    assert ws.cell(5, HEADERS.index("Source file") + 1).hyperlink.target.startswith(
        "evidence_files/"
    )
    originals, state, as_of = load_review_workbook(str(path))
    assert originals == [object_value(row.raw) for row in baseline.prepared.extraction.observations]
    assert len(originals) == 4
    assert len(state["human_review"]["rows"]) == 3
    assert all(not decision["changes"] for decision in state["human_review"]["decisions"])
    assert state["review_history"] == []
    assert as_of == "2026-01-31"


def test_corrections_and_exclusions_recompute_all_reports_and_bundle(review_pack, tmp_path):
    path, baseline, ids = review_pack
    originals = {p.name: p.read_bytes() for p in path.parent.joinpath("evidence_files").iterdir()}
    amend(
        path,
        "INV-A",
        **{"Period Charge (£)": 125, "Reason for change": "Verified OCR against original"},
    )
    amend(path, "INV-C", **{"Include": "No", "Reason for change": "Unrelated disputed account"})
    input_bytes = path.read_bytes()
    directory, outcome = regenerate_workbook(str(path), str(tmp_path))
    assert outcome.status in {"complete", "review_required"}, outcome.error
    assert outcome.snapshot is not None
    assert outcome.snapshot.summary.charges == Decimal("175.00")
    assert outcome.snapshot.summary.closing_balance == Decimal("300.00")
    assert outcome.snapshot.prepared.extraction.request.as_of == date(2026, 1, 31)
    state = object_value(outcome.snapshot.prepared.extraction.legacy_state)
    assert len(state["review_excluded_records"]) == 1
    assert ids[1] not in state["source_manifest"]
    assert ids[1] in state["review_excluded_source_manifest"]
    assert {entry["field"] for entry in state["review_history"]} == {
        "Period Charge (£)",
        "Transaction Amount (£)",
        "Include",
    }
    assert {p.read_bytes() for p in Path(directory, "evidence_files").iterdir()} == {
        originals["shared.pdf"],
        originals["duplicate.pdf"],
    }
    assert path.read_bytes() == input_bytes
    assert originals == {
        p.name: p.read_bytes() for p in path.parent.joinpath("evidence_files").iterdir()
    }
    assert [row.raw for row in outcome.snapshot.prepared.extraction.observations] == [
        row.raw for row in baseline.prepared.extraction.observations
    ]
    stored = load_result(Path(directory, "result.json"))
    assert stored.summary == outcome.snapshot.summary
    manifest = json.loads(Path(directory, "run.json").read_text())
    assert set(manifest["artifacts"]) == {"xlsx", "pdf", "docx", "html"}
    text = Path(directory, "Reviewed_Evidence_Report.html").read_text()
    assert "Human-reviewed analysis" in text
    assert "175.00" in text
    document = Document(str(Path(directory, "Reviewed_Evidence_Report.docx")))
    assert "Human-reviewed analysis" in "\n".join(p.text for p in document.paragraphs)
    workbook = Path(directory, "Reviewed_Evidence.xlsx")
    _, next_state, _ = load_review_workbook(str(workbook))
    assert len(next_state["review_history"]) == 3
    assert len(next_state["review_excluded_records"]) == 1
    assert build_report_view(stored, RenderSelection()).review_notice


def test_repeated_review_can_reinstate_excluded_sources(review_pack, tmp_path):
    path, _, ids = review_pack
    amend(path, "INV-C", **{"Include": "No", "Reason for change": "Exclude pending review"})
    first, outcome = regenerate_workbook(str(path), str(tmp_path), formats=("xlsx",))
    assert outcome.error is None
    assert outcome.snapshot is not None
    workbook = Path(first, "Reviewed_Evidence.xlsx")
    amend(
        workbook,
        "INV-C",
        **{
            "Include": "Yes",
            "Period Charge (£)": 80,
            "Reason for change": "Confirmed relevant and corrected",
        },
    )
    second, outcome = regenerate_workbook(str(workbook), str(tmp_path), formats=("xlsx",))
    assert outcome.error is None
    assert outcome.snapshot is not None
    assert outcome.snapshot.summary.charges == Decimal("230.00")
    state = object_value(outcome.snapshot.prepared.extraction.legacy_state)
    assert not state["review_excluded_records"]
    assert state["source_manifest"][ids[1]]["Saved Path"]
    assert len(list(Path(second, "evidence_files").iterdir())) == 3
    assert sum(change["field"] == "Include" for change in state["review_history"]) == 2


def test_excluding_all_records_preserves_auditable_empty_analysis(review_pack, tmp_path):
    path, baseline, _ = review_pack
    for invoice in ("INV-A", "INV-B", "INV-C"):
        amend(path, invoice, **{"Include": "No", "Reason for change": "Remove from complaint"})
    directory, outcome = regenerate_workbook(str(path), str(tmp_path))
    assert outcome.error is None
    assert outcome.snapshot is not None
    assert outcome.snapshot.summary.charges == 0
    assert rows_value(outcome.snapshot.prepared.records) == []
    assert len(outcome.snapshot.prepared.extraction.observations) == len(
        baseline.prepared.extraction.observations
    )
    assert not list(Path(directory, "evidence_files").iterdir())
    assert (
        len(
            load_review_workbook(str(Path(directory, "Reviewed_Evidence.xlsx")))[1][
                "review_excluded_records"
            ]
        )
        == 3
    )


@pytest.mark.parametrize(
    ("field", "value", "expected"),
    [
        ("Period Charge (£)", 123.456, "decimal places"),
        ("Period Charge (£)", "NaN", "finite number"),
        ("Period Charge (£)", True, "finite number"),
        ("Date", "31/02/2024", "real Excel date"),
        ("Date", "02/03/24", "real Excel date"),
        ("Period From", "01/02/2024", "later than Period To"),
        ("Invoice #", 123456, "format identifiers as Text"),
        ("Balance (£)", "=1+2", "Formulas"),
        ("Balance (£)", "#VALUE!", "Excel errors"),
        ("Include", "Maybe", "Include must be"),
        ("Amount Kind", "Whatever", "Amount Kind"),
    ],
)
def test_bad_review_inputs_fail_with_row_and_field_context(review_pack, field, value, expected):
    path, _, _ = review_pack
    amend(path, "INV-A", **{field: value, "Reason for change": "Review"})
    with pytest.raises(ValueError, match=expected) as error:
        load_review_workbook(str(path))
    assert "row 5" in str(error.value)


def test_changes_need_reasons_and_never_overwrite_inputs(review_pack, tmp_path):
    path, _, _ = review_pack
    amend(path, "INV-A", **{"Balance (£)": 275})
    before = path.read_bytes()
    with pytest.raises(ValueError, match="reason is required"):
        regenerate_workbook(str(path), str(tmp_path))
    assert not list(tmp_path.glob("reviewed_*"))
    assert path.read_bytes() == before


@pytest.mark.parametrize(
    "kind",
    ["delete", "duplicate", "source", "header", "reference", "payload", "both_sheets", "one_sheet"],
)
def test_damaged_or_misleading_workbooks_are_rejected(review_pack, kind):
    path, _, _ = review_pack
    wb = load_workbook(path)
    ws = wb[REVIEW_SHEET]
    if kind == "delete":
        ws.delete_rows(5)
    elif kind == "duplicate":
        for col in range(1, len(HEADERS) + 1):
            ws.cell(6, col).value = ws.cell(5, col).value
    elif kind == "source":
        ws.cell(5, HEADERS.index("Source file") + 1, "replacement.pdf")
    elif kind == "header":
        ws.cell(4, 1, "Remove?")
    elif kind == "reference":
        wb["EDF Evidence Report"]["H2"] = 123
    elif kind == "payload":
        wb[PAYLOAD_SHEET]["A2"] = "0" * 64
    elif kind == "both_sheets":
        del wb[REVIEW_SHEET]
        del wb[PAYLOAD_SHEET]
    else:
        del wb[PAYLOAD_SHEET]
    wb.save(path)
    with pytest.raises(ValueError):
        read_workbook(str(path))


def test_sorting_and_date_format_changes_preserve_identity(review_pack):
    path, _, _ = review_pack
    wb = load_workbook(path)
    ws = wb[REVIEW_SHEET]
    rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=5)]
    for index, values in enumerate(reversed(rows), 5):
        for column, value in enumerate(values, 1):
            ws.cell(index, column).value = value
    # Same date expressed in ISO text must not create a correction.
    ws.cell(7, HEADERS.index("Date") + 1, "2024-02-01")
    wb.save(path)
    _, state, _ = load_review_workbook(str(path))
    assert state["review_history"] == []


def test_clearing_money_does_not_resurrect_legacy_amounts(review_pack, tmp_path):
    path, _, _ = review_pack
    amend(
        path,
        "INV-A",
        **{
            "Balance (£)": None,
            "Period Charge (£)": None,
            "Reason for change": "Unreliable extraction; unknown amounts",
        },
    )
    _, outcome = regenerate_workbook(str(path), str(tmp_path), formats=("xlsx",))
    assert outcome.error is None
    assert outcome.snapshot is not None
    row = next(
        row for row in rows_value(outcome.snapshot.prepared.records) if row["Invoice #"] == "INV-A"
    )
    assert row["Balance (£)"] is None
    assert row["Transaction Amount (£)"] is None
    assert row["Period Charge (£)"] is None
    assert row["Amount (£)"] is None
    assert outcome.snapshot.summary.charges == Decimal("120.00")


def test_moved_pack_resolves_relative_sources_and_rejects_changed_bytes(review_pack, tmp_path):
    path, _, ids = review_pack
    moved = tmp_path / "moved"
    shutil.move(str(path.parent), moved)
    workbook = moved / path.name
    # Alter one source after extraction. The recorded hash must stop its copy.
    Path(moved, "evidence_files", "excluded.pdf").write_bytes(b"Changed original")
    directory, outcome = regenerate_workbook(str(workbook), str(tmp_path), formats=("xlsx",))
    assert outcome.status == "review_required"
    assert outcome.snapshot is not None
    state = object_value(outcome.snapshot.prepared.extraction.legacy_state)
    assert not state["source_manifest"][ids[1]]["Saved Path"]
    assert len(list(Path(directory, "evidence_files").iterdir())) == 2
    assert any("missing source" in note for note in state["error_log"])


def test_legacy_workbook_still_has_explicit_provenance_warning(tmp_path):
    path = tmp_path / "old.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame([{"Date": "01/01/2024", "Amount (£)": 42}]).to_excel(
            writer, sheet_name="EDF Evidence Report", index=False
        )
    records, state, _ = read_workbook(str(path))
    assert records[0]["Amount (£)"] == 42
    assert state["provenance"] == "lossy_workbook_import"
    assert "original observations" in state["error_log"][0]


def test_literal_formula_like_text_round_trips_without_execution(review_pack):
    path, _, _ = review_pack
    wb = load_workbook(path)
    cell = wb[REVIEW_SHEET].cell(5, HEADERS.index("Details") + 1)
    cell.value = "=literal evidence, not a formula"
    cell.data_type = "s"
    wb[REVIEW_SHEET].cell(5, 2, "Quoted source wording")
    wb.save(path)
    _, state, _ = load_review_workbook(str(path))
    assert state["human_review"]["decisions"][0]["changes"]["Details"].startswith("=literal")


def test_sap_rows_survive_excel_and_reconciliation_is_recomputed(review_pack, tmp_path):
    path, baseline, _ = review_pack
    sap_rows = [
        {"Posting Date": "01/02/2024", "Account": "000123", "Amount": 100, "Document No.": "SAP-1"}
    ]
    state = object_value(baseline.prepared.extraction.legacy_state)
    state["sap_financial_rows"] = sap_rows
    raw = [object_value(row.raw) for row in baseline.prepared.extraction.observations]
    snapshot = analyse_prepared(
        prepare_snapshot(capture_legacy(raw, state, baseline.prepared.extraction.request))
    )
    original_view = build_report_view(snapshot, RenderSelection(("xlsx",)))
    assert not any(
        row["side"] == "sap"
        for row in original_view.model("reconciliation")["details"]
        if row["entity"] == "Financial"
    )
    render_xlsx(original_view, str(path))
    amend(
        path,
        "INV-A",
        **{"Period Charge (£)": 125, "Reason for change": "Corrected original invoice"},
    )
    _, outcome = regenerate_workbook(str(path), str(tmp_path), formats=("xlsx",))
    assert outcome.error is None
    assert outcome.snapshot is not None
    assert (
        object_value(outcome.snapshot.prepared.extraction.legacy_state)["sap_financial_rows"]
        == sap_rows
    )
    reviewed_view = build_report_view(outcome.snapshot, RenderSelection(("xlsx",)))
    assert any(
        row["side"] == "sap"
        for row in reviewed_view.model("reconciliation")["details"]
        if row["entity"] == "Financial"
    )


def test_review_command_creates_fresh_workbook_even_for_html_only(review_pack, tmp_path, capsys):
    from edf_bill_fetcher.io.cli import run_cli_review

    path, _, _ = review_pack
    amend(path, "INV-A", **{"Period Charge (£)": 125, "Reason for change": "Source verified"})
    output = tmp_path / "cli-output"
    run_cli_review(["-i", str(path), "-o", str(output), "--formats", "html"])
    directory = next(output.glob("reviewed_*"))
    assert Path(directory, "Reviewed_Evidence.xlsx").is_file()
    assert Path(directory, "Reviewed_Evidence_Report.html").is_file()
    assert load_result(directory / "result.json").summary.charges == Decimal("245.00")
    assert "Reviewed pack:" in capsys.readouterr().out


def test_review_payload_size_is_enforced_before_decompression(review_pack):
    path, _, _ = review_pack
    wb = load_workbook(path)
    wb[PAYLOAD_SHEET]["A3"] = 1
    wb.save(path)
    with pytest.raises(ValueError, match="payload length"):
        load_review_workbook(str(path))


def test_long_history_and_forecast_survive_excel_numeric_serialization(review_pack):
    path, baseline, _ = review_pack
    records = [
        {
            "Date": f"28/{index % 12 + 1:02}/{2015 + index // 12}",
            "Period From": f"01/{index % 12 + 1:02}/{2015 + index // 12}",
            "Period To": f"28/{index % 12 + 1:02}/{2015 + index // 12}",
            "Invoice #": f"B{index:03}",
            "Entry Type": "New Bill",
            "Amount (£)": 100 + index * 20.13,
            "Period Charge (£)": 83 + index * 2.13,
            "Units (kWh)": 131 + index * 1.235736,
            "Energy Unit Rate (p/kWh)": 20.135796,
            "Source": "Local PDF",
            "Account": "0001",
        }
        for index in range(48)
    ]
    snapshot = analyse_prepared(
        prepare_snapshot(capture_legacy(records, {}, baseline.prepared.extraction.request))
    )
    render_xlsx(build_report_view(snapshot, RenderSelection(("xlsx",))), str(path))
    originals, state, _ = load_review_workbook(str(path))
    assert originals == records
    assert state["review_history"] == []


def test_cleared_invoice_reference_stays_blank_in_next_review(review_pack, tmp_path):
    path, _, _ = review_pack
    amend(
        path, "INV-A", **{"Invoice #": None, "Reason for change": "No readable invoice reference"}
    )
    directory, outcome = regenerate_workbook(str(path), str(tmp_path), formats=("xlsx",))
    assert outcome.error is None
    _, state, _ = load_review_workbook(str(Path(directory, "Reviewed_Evidence.xlsx")))
    assert all(not decision["changes"] for decision in state["human_review"]["decisions"])


def test_moved_cli_style_pack_finds_bundle_below_workbook_folder(review_pack, tmp_path):
    path, baseline, _ = review_pack
    container = tmp_path / "portable"
    container.mkdir()
    shutil.move(str(path.parent), container / "original")
    state = object_value(baseline.prepared.extraction.legacy_state)
    state["source_bundle_path"] = str(container / "original")
    snapshot = replace(
        baseline,
        prepared=replace(
            baseline.prepared,
            extraction=replace(baseline.prepared.extraction, legacy_state=freeze_json(state)),
        ),
    )
    render_xlsx(
        build_report_view(snapshot, RenderSelection(("xlsx",))), str(container / "Evidence.xlsx")
    )
    moved = tmp_path / "moved"
    container.rename(moved)
    directory, outcome = regenerate_workbook(
        str(moved / "Evidence.xlsx"), str(tmp_path), formats=("xlsx",)
    )
    assert outcome.error is None
    assert len(list(Path(directory, "evidence_files").iterdir())) == 3
