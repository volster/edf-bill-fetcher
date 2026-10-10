"""Real-file regressions for evidence identity, filename mappings and portable links."""

import hashlib
import json
import os
import shutil
from datetime import date, datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile

import docx
import openpyxl
import pandas as pd
import pytest
from reportlab.pdfgen.canvas import Canvas

from edf_bill_fetcher.application.run_service import bundle_engine_sources, run
from edf_bill_fetcher.collectors.engine import EvidenceEngine
from edf_bill_fetcher.domain.json_values import freeze_json, object_value
from edf_bill_fetcher.domain.requests import RunRequest
from edf_bill_fetcher.io.persistence.run_manifest import is_run_directory
from edf_bill_fetcher.io.writers import evidence_bundle
from edf_bill_fetcher.io.writers.evidence_bundle import resolve_bundle_paths, save_evidence_files


def bill(path, invoices=("KI-10000000-0001",)):
    canvas = Canvas(str(path), invariant=True)
    for invoice in invoices:
        text = canvas.beginText(40, 800)
        for line in (
            "Your VAT invoice",
            "Invoice number: " + invoice,
            "Account number: A-10000000",
            "Date issued: 15 January 2026",
            "Your charges: 01 Dec 2025 - 31 Dec 2025",
            "Current balance £100.00 debit",
            "Total charges for this period £60.00 debit",
            "Electricity used 250 kWh",
        ):
            text.textLine(line)
        canvas.drawText(text)
        canvas.showPage()
    canvas.save()
    return path


def engine():
    return EvidenceEngine(
        {
            "as_of": "2026-01-31",
            "filter_below": False,
            "save_evidence_files": True,
            "use_anchors": True,
            "use_domain_filter": True,
            "domain_filter": "edfenergy.com",
        },
        lambda text: None,
    )


def pack(collector, directory):
    directory.mkdir()
    request = RunRequest(
        "bundle-regression",
        datetime(2026, 1, 31, tzinfo=timezone.utc),
        date(2026, 1, 31),
        (),
        freeze_json(collector.config),
    )
    return run(
        request,
        {"xlsx": str(directory / "evidence.xlsx")},
        result_path=str(directory / "result.json"),
        manifest_path=str(directory / "run.json"),
        collect=lambda: (collector.records, collector.__getstate__()),
        copy_sources=lambda prepared: bundle_engine_sources(collector, prepared, str(directory)),
        cleanup=collector.cleanup,
    )


def workbook_targets(directory, sheet="EDF Evidence Report"):
    workbook = openpyxl.load_workbook(directory / "evidence.xlsx")
    try:
        ws = workbook[sheet]
        column = [cell.value for cell in ws[1]].index("Attachment Name") + 1
        return [
            (
                ws.cell(row, column).value,
                ws.cell(row, column).hyperlink.target if ws.cell(row, column).hyperlink else None,
            )
            for row in range(2, ws.max_row + 1)
        ]
    finally:
        workbook.close()


def test_same_named_pdfs_from_different_folders_keep_correct_bytes_and_links(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    first_dir, second_dir = source / "one", source / "two"
    first_dir.mkdir()
    second_dir.mkdir()
    first = bill(first_dir / "invoice.pdf", ("KI-10000000-0001",))
    second = bill(second_dir / "invoice.pdf", ("KI-10000000-0002",))
    collector = engine()
    collector.crawl_local_pdfs(str(source))
    assert len(collector.records) == 2
    assert len({row["Attachment Name"] for row in collector.records}) == 2
    directory = tmp_path / "pack"
    outcome = pack(collector, directory)
    assert outcome.status == "complete", outcome.error
    manifest = json.loads((directory / "evidence_manifest.json").read_text())
    assert len(manifest["files"]) == 2
    originals = {"KI-10000000-0001": first, "KI-10000000-0002": second}
    mappings = {item["attachment_name"]: item for item in manifest["files"]}
    for att, target in workbook_targets(directory):
        entry = mappings[att]
        assert target == entry["saved_path"]
        saved = directory / target
        assert saved.read_bytes() == originals[entry["invoices"][0]].read_bytes()
        assert entry["original_filename"] == "invoice.pdf"
        assert entry["sha256"] == hashlib.sha256(saved.read_bytes()).hexdigest()
    state = object_value(outcome.snapshot.prepared.extraction.legacy_state)
    for source_id, metadata in state["source_manifest"].items():
        assert state["saved_evidence_paths"][source_id] == str(directory / metadata["Saved Path"])


def test_merged_pdf_is_copied_once_and_all_invoice_pages_are_mapped(tmp_path):
    original = bill(tmp_path / "merged.pdf", ("KI-10000000-0001", "KI-10000000-0002"))
    collector = engine()
    collector.process_pdf_file(
        str(original), "PDF", "merged.pdf", "01/01/2026", attachment_name="merged.pdf"
    )
    assert len(collector.records) == 2
    directory = tmp_path / "pack"
    outcome = pack(collector, directory)
    assert outcome.status == "complete", outcome.error
    copied = list((directory / "evidence_files").iterdir())
    assert [path.name for path in copied] == ["merged.pdf"]
    assert copied[0].read_bytes() == original.read_bytes()
    targets = workbook_targets(directory)
    assert len(targets) == 2
    assert {target for _, target in targets} == {"evidence_files/merged.pdf"}
    manifest = json.loads((directory / "evidence_manifest.json").read_text())
    sliced = [entry for entry in manifest["files"] if entry["invoices"]]
    assert {tuple(entry["invoices"]) for entry in sliced} == {
        ("KI-10000000-0001",),
        ("KI-10000000-0002",),
    }
    assert {tuple(entry["page_ranges"][0]) for entry in sliced} == {(1, 1), (2, 2)}
    index = docx.Document(directory / "evidence_index.docx")
    assert (
        len([p for p in index.paragraphs if p.style is not None and p.style.name == "List Bullet"])
        == 1
    )


def test_email_attachment_survives_cleanup_and_filename_collision(tmp_path):
    mail = tmp_path / "mail"
    mail.mkdir()
    originals = []
    for number in (1, 2):
        original = bill(tmp_path / f"original-{number}.pdf", (f"KI-10000000-000{number}",))
        originals.append(original.read_bytes())
        message = EmailMessage()
        message["From"] = "billing@edfenergy.com"
        message["Subject"] = "Your EDF bill"
        message["Date"] = "Thu, 15 Jan 2026 10:00:00 +0000"
        message.set_content("Please find the bill attached.")
        message.add_attachment(
            original.read_bytes(), maintype="application", subtype="pdf", filename="invoice.pdf"
        )
        (mail / f"message-{number}.eml").write_bytes(message.as_bytes())
    collector = engine()
    collector.process_eml_folder(str(mail))
    assert len(collector.records) == 2
    temporary = Path(collector._attachment_dir.name)
    directory = tmp_path / "pack"
    outcome = pack(collector, directory)
    assert outcome.status == "complete", outcome.error
    assert not temporary.exists()
    files = list((directory / "evidence_files").glob("*.pdf"))
    assert sorted(path.read_bytes() for path in files) == sorted(originals)
    assert len(list((directory / "evidence_files").glob("*.eml"))) == 2
    for _, target in workbook_targets(directory):
        assert (directory / target).read_bytes() in originals


def test_source_id_wins_over_a_stale_display_name_lookup(tmp_path):
    right, wrong = tmp_path / "right.pdf", tmp_path / "wrong.pdf"
    right.write_bytes(b"right evidence")
    wrong.write_bytes(b"wrong evidence")
    frame = pd.DataFrame(
        [
            {
                "Attachment Name": "invoice.pdf",
                "Source ID": "source-1",
                "Content Hash": hashlib.sha256(right.read_bytes()).hexdigest(),
            }
        ]
    )
    saved = save_evidence_files(
        frame, {"source-1": str(right), "invoice.pdf": str(wrong)}, str(tmp_path / "out")
    )
    assert Path(saved["invoice.pdf"]).read_bytes() == b"right evidence"


def test_ambiguous_same_name_source_identities_are_not_silently_mislinked(tmp_path):
    first, second = tmp_path / "one.pdf", tmp_path / "two.pdf"
    first.write_bytes(b"one")
    second.write_bytes(b"two")
    frame = pd.DataFrame(
        [
            {"Attachment Name": "same.pdf", "Source ID": "one"},
            {"Attachment Name": "same.pdf", "Source ID": "two"},
        ]
    )
    stats: dict[str, int] = {}
    logs: list[str] = []
    saved = save_evidence_files(
        frame,
        {"one": str(first), "two": str(second)},
        str(tmp_path / "out"),
        stats=stats,
        log=logs.append,
    )
    assert saved == {}
    assert stats == {"saved": 0, "missing": 1, "ambiguous": 1}
    assert any("ambiguous source identity" in line for line in logs)
    assert list((tmp_path / "out").iterdir()) == []


def test_changed_source_is_not_published_and_pack_requires_review(tmp_path):
    original = bill(tmp_path / "invoice.pdf")
    collector = engine()
    collector.process_pdf_file(str(original), "PDF", "invoice.pdf", "01/01/2026")
    original.write_bytes(b"changed after extraction")
    directory = tmp_path / "pack"
    outcome = pack(collector, directory)
    assert outcome.status == "review_required", outcome.error
    assert list((directory / "evidence_files").iterdir()) == []
    assert all(target is None for _, target in workbook_targets(directory))
    state = object_value(outcome.snapshot.prepared.extraction.legacy_state)
    assert any("changed after extraction" in note for note in state["error_log"])
    manifest = json.loads((directory / "evidence_manifest.json").read_text())
    assert manifest["files"][0]["status"] == "unavailable"
    assert manifest["files"][0]["saved_path"] is None


def test_partial_copy_failure_leaves_no_incomplete_file_or_false_mapping(tmp_path, monkeypatch):
    source = tmp_path / "source.pdf"
    source.write_bytes(b"complete original")

    def fail(src, dest):
        Path(dest).write_bytes(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(evidence_bundle.shutil, "copy2", fail)
    frame = pd.DataFrame([{"Attachment Name": "source.pdf"}])
    saved = save_evidence_files(frame, {"source.pdf": str(source)}, str(tmp_path / "out"))
    assert saved == {}
    assert list((tmp_path / "out").iterdir()) == []
    assert source.read_bytes() == b"complete original"


def test_case_and_sanitisation_collisions_keep_both_files(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    first.write_bytes(b"first")
    second.write_bytes(b"second")
    frame = pd.DataFrame([{"Attachment Name": "A:B.pdf"}, {"Attachment Name": "a?b.pdf"}])
    saved = save_evidence_files(
        frame, {"A:B.pdf": str(first), "a?b.pdf": str(second)}, str(tmp_path / "out")
    )
    assert len({Path(path).name.casefold() for path in saved.values()}) == 2
    assert {Path(path).read_bytes() for path in saved.values()} == {b"first", b"second"}


def test_duplicate_evidence_rows_keep_real_source_links(tmp_path):
    first = bill(tmp_path / "first.pdf")
    second = bill(tmp_path / "second.pdf")
    # Add different PDF metadata without changing the invoice: distinct originals.
    second.write_bytes(second.read_bytes() + b"\n% extra metadata\n")
    collector = engine()
    collector.config.update(use_dedup=True, save_dups=True)
    for source in (first, second):
        collector.process_pdf_file(str(source), "PDF", source.name, "01/01/2026")
    directory = tmp_path / "pack"
    outcome = pack(collector, directory)
    assert outcome.status == "complete", outcome.error
    duplicates = workbook_targets(directory, "Duplicate Entries")
    assert len(duplicates) == 1
    assert all(target and (directory / target).is_file() for _, target in duplicates)
    assert len(list((directory / "evidence_files").glob("*.pdf"))) == 2


def test_bundle_index_mapping_and_links_survive_moving_the_whole_pack(tmp_path):
    original = bill(tmp_path / "original.pdf")
    collector = engine()
    collector.process_pdf_file(str(original), "PDF", "original.pdf", "01/01/2026")
    directory = tmp_path / "pack"
    assert pack(collector, directory).status == "complete"
    relocated = tmp_path / "relocated"
    shutil.move(str(directory), relocated)
    for _, target in workbook_targets(relocated):
        assert not os.path.isabs(target)
        assert (relocated / target).read_bytes() == original.read_bytes()
    document = docx.Document(relocated / "evidence_index.docx")
    bullets = [
        p for p in document.paragraphs if p.style is not None and p.style.name == "List Bullet"
    ]
    assert len(bullets) == 1
    assert "original.pdf" in bullets[0].text
    with ZipFile(relocated / "evidence_index.docx") as archive:
        from xml.etree import ElementTree

        relationships = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
        links = [
            item.attrib["Target"]
            for item in relationships
            if item.attrib["Type"].endswith("/hyperlink")
        ]
    assert links == ["evidence_files/KI-10000000-0001.pdf"]
    assert (relocated / unquote(links[0])).read_bytes() == original.read_bytes()


def test_recursive_scans_skip_previous_output_but_keep_nested_originals(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    year = source / "2026"
    year.mkdir()
    original = bill(year / "original.pdf")
    collector = engine()
    collector.crawl_local_pdfs(str(source))
    directory = source / "old-pack"
    assert pack(collector, directory).status == "complete"
    bill(directory / "generated_report.pdf", ("KI-99999999-9999",))
    rerun = engine()
    rerun.crawl_local_pdfs(str(source))
    assert rerun.pdf_count == 1
    assert len(rerun.records) == 1
    assert rerun.records[0]["Original Path"] == str(original)
    assert is_run_directory(str(directory))


@pytest.mark.parametrize(
    "content",
    [
        "not JSON",
        "[]",
        '{"manifest_version":1}',
        '{"manifest_version":1,"run_id":"unrelated","requested":{"other":1},"artifacts":{}}',
    ],
)
def test_unrelated_run_json_does_not_hide_sources(tmp_path, content):
    (tmp_path / "run.json").write_text(content)
    original = bill(tmp_path / "source.pdf")
    assert not is_run_directory(str(tmp_path))
    collector = engine()
    collector.crawl_local_pdfs(str(tmp_path))
    assert collector.pdf_count == 1
    assert collector.records[0]["Original Path"] == str(original)


def test_message_scan_also_skips_copied_messages_in_previous_packs(tmp_path):
    message = EmailMessage()
    message["From"] = "billing@edfenergy.com"
    message["Subject"] = "Your EDF bill"
    message["Date"] = "Thu, 15 Jan 2026 10:00:00 +0000"
    message.set_content("Your new account balance £100.00")
    (tmp_path / "source.eml").write_bytes(message.as_bytes())
    collector = engine()
    collector.process_eml_folder(str(tmp_path))
    assert collector.email_count == 1
    directory = tmp_path / "old-pack"
    assert pack(collector, directory).status == "complete"
    rerun = engine()
    rerun.process_eml_folder(str(tmp_path))
    assert rerun.email_count == 1


def test_relative_source_mapping_cannot_escape_the_selected_bundle(tmp_path):
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"not this bundle's source")
    root = tmp_path / "bundle"
    root.mkdir()
    paths, notes = resolve_bundle_paths(
        {
            "source_manifest": {
                "id": {"Saved Path": "../outside.pdf", "Attachment Name": "invoice.pdf"}
            }
        },
        str(root),
    )
    assert paths == {}
    assert "unavailable" in notes[0]


def test_filename_whitespace_is_preserved_in_lookup_and_manifest(tmp_path):
    source = tmp_path / " bill.pdf"
    source.write_bytes(b"evidence")
    frame = pd.DataFrame([{"Attachment Name": source.name}])
    saved = save_evidence_files(frame, {source.name: str(source)}, str(tmp_path / "out"))
    assert Path(saved[source.name]).read_bytes() == b"evidence"


def test_conflicting_hashes_do_not_publish_a_falsely_verified_source(tmp_path):
    source = tmp_path / "source.pdf"
    source.write_bytes(b"evidence")
    frame = pd.DataFrame(
        [
            {"Attachment Name": "a.pdf", "Content Hash": "one"},
            {"Attachment Name": "b.pdf", "Content Hash": "two"},
        ]
    )
    saved = save_evidence_files(
        frame, {"a.pdf": str(source), "b.pdf": str(source)}, str(tmp_path / "out")
    )
    assert saved == {}
    assert list((tmp_path / "out").iterdir()) == []
