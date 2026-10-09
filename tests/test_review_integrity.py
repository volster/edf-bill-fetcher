"""Regression contracts for the application review's evidence integrity fixes."""

from __future__ import annotations

import json
from email.message import EmailMessage
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

from edf_bill_fetcher.helpers.artifacts import atomic_output
from edf_bill_fetcher.helpers.date_utils import _safe_to_datetime, parse_to_sort_date
from edf_bill_fetcher.io.adapters.eml import parse_eml_message
from edf_bill_fetcher.io.cli import _load_report_inputs, _safe_json_dump
from edf_bill_fetcher.io.reporters.pdf_report import generate_pdf_from_gui
from edf_bill_fetcher.io.writers.export import export_to_excel
from edf_bill_fetcher.io.writers.reconciliation import write_reconciliation_sheet
from edf_bill_fetcher.models.report_models import compute_forecast, compute_ofgem_comparison
from edf_bill_fetcher.processors.compensation import estimate_compensation
from edf_bill_fetcher.processors.detection import detect_meter_rollover
from edf_bill_fetcher.processors.financial import financial_totals, mean_daily_rate
from edf_bill_fetcher.processors.preparation import prepare_evidence


def bill(**overrides):
    return {
        "Date": "01/02/2025",
        "Period From": "01/01/2025",
        "Period To": "01/02/2025",
        "Entry Type": "New Bill",
        "Invoice #": "KI-1",
        "Source": "PDF",
        "Amount (£)": 1000,
        "Period Charge (£)": 100,
        "Account": "A-12345678",
        **overrides,
    }


def test_balance_observations_do_not_count_as_charges_or_payments():
    rows = [
        bill(),
        bill(
            **{
                "Source": "HTM Account History",
                "Entry Type": "Payment",
                "Period Charge (£)": 20,
                "Amount (£)": 800,
            }
        ),
        {"Source": "PDF", "Entry Type": "Credit", "Amount (£)": -10},
    ]
    assert financial_totals(pd.DataFrame(rows)) == (100, 30)
    assert mean_daily_rate(pd.DataFrame(rows)) == pytest.approx(100 / 31)


def test_distinct_invoice_and_account_identities_survive_deduplication():
    rows = [bill(), bill(**{"Invoice #": "KI-2"}), bill(**{"Account": "A-99999999"})]
    assert len(prepare_evidence(rows, {"save_dups": False}).records) == 3


def test_duplicate_admission_and_provenance_survive_hidden_duplicate_sheet():
    rows = [
        bill(**{"Source ID": "one"}),
        bill(**{"Source ID": "two", "Cancel/Rebill Admitted": True}),
    ]
    result = prepare_evidence(rows, {"save_dups": False})
    assert len(result.records) == 1
    assert result.records[0]["Cancel/Rebill Admitted"] is True
    assert result.records[0]["Source IDs"] == ["one", "two"]
    assert result.duplicates[0]["_matches_kept_idx"] == 0


def test_iso_and_uk_dates_match_in_scalar_and_vector_paths():
    values = ["2025-02-01", "01/02/2025", "2025-02-01T12:00:00Z", "N/A"]
    vector = _safe_to_datetime(pd.Series(values))
    for idx, value in enumerate(values):
        scalar = parse_to_sort_date(value)
        assert pd.isna(vector[idx]) if pd.isna(scalar) else vector[idx] == scalar
    assert vector[0] == vector[1]


def test_period_consumption_never_supplies_a_register_rollover():
    records = [
        bill(**{"Reading": "Actual", "Units (kWh)": 99999}),
        bill(**{"Date": "02/02/2025", "Reading": "Actual", "Units (kWh)": 1}),
    ]
    assert detect_meter_rollover(pd.DataFrame(records)).empty
    records[0]["Register Reading (kWh)"] = 99999
    records[1]["Register Reading (kWh)"] = 1
    assert len(detect_meter_rollover(pd.DataFrame(records))) == 1
    records[1]["Account"] = "A-99999999"
    assert detect_meter_rollover(pd.DataFrame(records)).empty


def test_reconciliation_suppresses_matched_financial_rows():
    wb = openpyxl.Workbook()
    summary, detail = wb.active, wb.create_sheet()
    evidence = pd.DataFrame([bill()])
    sap = [{"Posting Date": "2025-02-01", "Amount": "100", "Account": "A-12345678"}]
    write_reconciliation_sheet(
        summary, detail, [], pd.DataFrame(), [], pd.DataFrame(), sap, evidence
    )
    assert summary.cell(6, 4).value == 1
    assert summary.cell(6, 5).value == 0
    assert summary.cell(6, 6).value == 0
    assert summary.cell(6, 7).value == "MATCHED"
    assert detail.max_row == 3


def test_cap_checks_do_not_invent_energy_rates_or_compliance():
    comparison = compute_ofgem_comparison(pd.DataFrame([bill(**{"Units (kWh)": 100})]))
    assert comparison.rows == []
    assert comparison.overall_verdict == "NOT ASSESSED"
    assert compute_ofgem_comparison(pd.DataFrame()).overall_verdict == "NOT ASSESSED"


def test_forecasts_require_consecutive_monthly_balances():
    records = [bill(**{"Date": date}) for date in ("01/01/2025", "01/03/2025", "01/04/2025")]
    forecast = compute_forecast(pd.DataFrame(records))
    assert forecast.linear_forecast == []


def test_credit_interest_uses_one_episode_and_actual_refund_movements():
    records = [
        {"Date": "01/01/2025", "Amount (£)": -100, "Amount Kind": "Balance"},
        {"Date": "16/01/2025", "Amount (£)": -100, "Amount Kind": "Balance"},
        {"Date": "16/01/2025", "Amount (£)": 50, "Entry Type": "Refund"},
    ]
    config = {"as_of": "2025-02-01", "credit_hold_days": 0, "credit_interest_rate": 0.1}
    rows = estimate_compensation(pd.DataFrame(records), config)
    assert len(rows) == 1
    assert rows[0]["indicative_amount"] == round((100 * 15 + 50 * 16) * 0.1 / 365, 2)
    assert estimate_compensation(pd.DataFrame(records), {"as_of": "2025-02-01"}) == []


def test_atomic_output_keeps_previous_file_when_generation_fails(tmp_path):
    destination = tmp_path / "report.txt"
    destination.write_text("complete")

    @atomic_output()
    def generate(output_path):
        from pathlib import Path

        Path(output_path).write_text("partial")
        raise RuntimeError("generation failed")

    with pytest.raises(RuntimeError, match="generation failed"):
        generate(str(destination))
    assert destination.read_text() == "complete"
    assert len(list(tmp_path.iterdir())) == 1


def test_report_section_failure_is_not_reported_as_success(tmp_path, monkeypatch):
    import edf_bill_fetcher.io.reporters.pdf_report as reporter

    def fail(*args, **kwargs):
        raise RuntimeError("broken section")

    monkeypatch.setattr(reporter, "create_data_quality_section", fail)
    path = tmp_path / "report.pdf"
    success, message = generate_pdf_from_gui(
        [bill()], str(path), {"report_sections": ["data_quality"]}
    )
    assert success is False
    assert "incomplete" in message
    assert not path.exists()


def test_eml_attachment_bytes_are_kept_out_of_the_body(tmp_path):
    message = EmailMessage()
    message["From"] = "billing@edfenergy.com"
    message["Subject"] = "EDF bill"
    message.set_content("Your bill is attached")
    message.add_attachment(b"%PDF-test", maintype="application", subtype="pdf", filename="bill.pdf")
    path = tmp_path / "bill.eml"
    path.write_bytes(message.as_bytes())
    parsed = parse_eml_message(path)
    assert parsed["attachments"] == [{"name": "bill.pdf", "data": b"%PDF-test"}]
    assert parsed["body_text"].strip() == "Your bill is attached"


def test_sap_only_exports_keep_external_values_literal(tmp_path):
    path = tmp_path / "sap.xlsx"
    export_to_excel([], str(path), [], {}, sap_rows={"contract": [{"Product Code": "=1+1"}]})
    wb = openpyxl.load_workbook(path)
    matches = [cell for row in wb["SAP Contract History"] for cell in row if cell.value == "=1+1"]
    assert len(matches) == 1
    assert matches[0].data_type == "s"


def test_json_restores_run_configuration_counts_and_missing_values(tmp_path):
    path = tmp_path / "records.json"
    _safe_json_dump(
        {
            "records": [bill(**{"Units (kWh)": float("nan")})],
            "config": {"as_of": "2025-02-01"},
            "pdf_count": 3,
            "email_count": 4,
            "filtered_records": [bill()],
        },
        str(path),
    )
    data = json.loads(path.read_text())
    assert data["version"] == 2
    assert data["records"][0]["Units (kWh)"] is None
    records, config, engine, filtered = _load_report_inputs(
        SimpleNamespace(records=path, config=None, engine_data=None)
    )
    assert len(records) == 1
    assert config["as_of"] == "2025-02-01"
    assert (engine.pdf_count, engine.email_count, len(filtered)) == (3, 4, 1)


def test_repeated_payments_in_one_document_are_distinct_without_transaction_identity():
    payment = {
        "Date": "01/02/2025",
        "Entry Type": "Payment",
        "Amount (£)": 20,
        "Source ID": "statement",
        "Logic Used": "Reconciliation Statement",
        "Invoice #": "STATEMENT-1",
    }
    assert len(prepare_evidence([payment, payment]).records) == 2
    first = {**payment, "Source Row": 1}
    second = {**payment, "Source Row": 2}
    assert len(prepare_evidence([first, second, first]).records) == 2


def test_statement_charges_and_refunds_are_real_movements():
    rows = [
        {"Entry Type": "Charge", "Logic Used": "Reconciliation Statement Charge", "Amount (£)": 50},
        {"Entry Type": "Credit", "Amount (£)": -100},
        {"Entry Type": "Refund", "Amount (£)": 40},
    ]
    assert financial_totals(pd.DataFrame(rows)) == (50, 60)


def test_undated_balance_has_no_chronological_percentage_formula():
    from edf_bill_fetcher.io.writers.evidence import write_evidence_sheet

    ws = openpyxl.Workbook().active
    write_evidence_sheet(
        ws,
        pd.DataFrame(
            [
                {"Date": "01/01/2025", "Amount Kind": "Balance", "Amount (£)": 100},
                {"Date": "N/A", "Amount Kind": "Balance", "Amount (£)": 300},
            ]
        ),
    )
    assert ws["J3"].value is None
    assert ws["S3"].data_type != "f"


def test_json_retains_duplicate_decisions_when_canonical_records_are_supplied(tmp_path):
    originals = [bill(), bill(**{"Source": "HTM Account History"})]
    prepared = prepare_evidence(originals)
    path = tmp_path / "state.json"
    _safe_json_dump(
        {"records": prepared.records, "original_records": originals, "config": prepared.config},
        str(path),
    )
    loaded = json.loads(path.read_text())
    assert len(loaded["records"]) == 1
    assert len(loaded["duplicate_records"]) == 1
    assert len(loaded["original_records"]) == 2


def test_superseded_sheet_does_not_invent_source_links(tmp_path):
    rows = [
        bill(
            **{
                "Invoice #": "OLD",
                "Date": "01/03/2024",
                "Period From": "01/01/2022",
                "Period To": "01/02/2024",
            }
        ),
        bill(
            **{
                "Invoice #": "NEW",
                "Date": "01/04/2024",
                "Period From": "01/01/2022",
                "Period To": "01/02/2024",
                "Period Charge (£)": 120,
            }
        ),
    ]
    out = tmp_path / "evidence.xlsx"
    export_to_excel(rows, str(out), [], {"analysis_min": 0})
    wb = openpyxl.load_workbook(out)
    ws = wb["Superseded Reconciliation"]
    assert all(
        cell.hyperlink is None or cell.hyperlink.target is None for row in ws for cell in row
    )
