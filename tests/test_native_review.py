"""Native review and bill revision contracts, including real exports and draft recovery."""

import copy
import json
from dataclasses import replace
from decimal import Decimal

import pytest
from docx import Document
from openpyxl import load_workbook

from edf_bill_fetcher.application.analysis_service import analyse_prepared
from edf_bill_fetcher.application.evidence_review import build_evidence_review
from edf_bill_fetcher.application.legacy_adapter import capture_legacy
from edf_bill_fetcher.application.preparation_service import prepare_snapshot
from edf_bill_fetcher.application.review_session import MAX_DRAFT_BYTES, ReviewSession
from edf_bill_fetcher.application.run_service import RunOutcome
from edf_bill_fetcher.domain.json_values import freeze_json, object_value, rows_value
from edf_bill_fetcher.io.persistence.result_store import load_result, save_result
from edf_bill_fetcher.io.review_workbook import load_review_workbook as _load_review_workbook
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view
from edf_bill_fetcher.processors.revisions import build_revisions, validate_decisions
from edf_bill_fetcher.ui.qt_jobs import DesktopJob, RunWorker
from tests.test_excel_review import amend
from tests.test_excel_review import review_pack as review_pack


def load_review_workbook(path: str) -> tuple[list[dict], dict, str]:
    imported = _load_review_workbook(path)
    assert imported is not None
    return imported


def row_id(session, invoice):
    return next(row["id"] for row in session.rows if row["record"]["Invoice #"] == invoice)


def analyse_session(session, source):
    records, state, _ = session.inputs(str(source))
    request = replace(
        session.snapshot.prepared.extraction.request, config=freeze_json(state["review_config"])
    )
    return analyse_prepared(prepare_snapshot(capture_legacy(records, state, request)))


def test_native_excel_parity_and_undo_recovery(review_pack, tmp_path):
    path, snapshot, _ = review_pack
    original = path.read_bytes()
    session = ReviewSession(snapshot)
    first, third = row_id(session, "INV-A"), row_id(session, "INV-C")
    session.update(
        first,
        {"Period Charge (£)": "125", "Account": "000123"},
        True,
        "Verified original",
        "verified",
    )
    session.update(third, {}, False, "Unrelated account", "needs_information")
    assert session.record(first)["Transaction Amount (£)"] == 125
    session.undo()
    assert session.decision(third)["include"] is True
    session.redo()
    draft = tmp_path / "progress.json"
    session.save(draft)
    restored = ReviewSession.load(snapshot, draft)
    assert restored.payload() == session.payload()
    native = analyse_session(restored, path)
    assert native.summary.charges == Decimal("175.00")
    assert path.read_bytes() == original
    amend(path, "INV-A", **{"Period Charge (£)": 125, "Reason for change": "Verified original"})
    amend(path, "INV-C", **{"Include": "No", "Reason for change": "Unrelated account"})
    records, state, _ = load_review_workbook(str(path))
    excel = analyse_prepared(
        prepare_snapshot(capture_legacy(records, state, snapshot.prepared.extraction.request))
    )
    assert native.summary == excel.summary
    assert rows_value(native.prepared.records) == rows_value(excel.prepared.records)
    assert [r.raw for r in native.prepared.extraction.observations] == [
        r.raw for r in snapshot.prepared.extraction.observations
    ]
    assert (
        object_value(native.prepared.extraction.legacy_state)["human_review_statuses"][first][
            "status"
        ]
        == "verified"
    )


@pytest.mark.parametrize(
    "mutation",
    [
        "foreign",
        "digest",
        "version",
        "duplicate",
        "unknown",
        "source",
        "reason",
        "status",
        "revisions",
    ],
)
def test_draft_rejects_invalid_inputs_without_changing_session(review_pack, mutation):
    _, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    key = row_id(session, "INV-A")
    session.update(key, {"Period Charge (£)": 120}, True, "Source checked", "verified")
    before = session.payload()
    draft = copy.deepcopy(before)
    if mutation == "foreign":
        draft["base_fingerprint"] = "0" * 64
    elif mutation == "digest":
        draft["base_digest"] = "0" * 64
    elif mutation == "version":
        draft["version"] = True
    elif mutation == "duplicate":
        draft["decisions"] *= 2
    elif mutation == "unknown":
        draft["decisions"][0]["row_id"] = "missing"
    elif mutation == "source":
        draft["decisions"][0]["changes"]["Source PDF Text"] = "forged"
    elif mutation == "reason":
        draft["decisions"][0]["reason"] = ""
    elif mutation == "status":
        draft["statuses"][key]["status"] = []
    else:
        draft["revisions"] = [{}]
    with pytest.raises(ValueError):
        ReviewSession.restore(snapshot, draft)
    assert session.payload() == before


def test_draft_normalizes_numeric_strings_and_preserves_explicit_blanks(review_pack):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    key = row_id(session, "INV-A")
    session.update(
        key,
        {"Period Charge (£)": "120", "Period From": ""},
        True,
        "Correct OCR and remove guessed period",
        "unreviewed",
    )
    payload = session.payload()
    payload["decisions"][0]["changes"]["Period Charge (£)"] = "120"
    restored = ReviewSession.restore(snapshot, payload)
    assert restored.record(key)["Transaction Amount (£)"] == 120
    result = analyse_session(restored, path)
    assert (
        next(r for r in rows_value(result.prepared.records) if r["Invoice #"] == "INV-A")[
            "Period From"
        ]
        is None
    )
    payload["decisions"][0]["changes"]["Period Charge (£)"] = "=SUM(1,2)"
    with pytest.raises(ValueError):
        ReviewSession.restore(snapshot, payload)


def test_draft_size_corruption_and_stale_result(review_pack, tmp_path):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    draft = tmp_path / "draft.json"
    draft.write_text('{"version":1,"version":1}')
    with pytest.raises(ValueError, match="Duplicate"):
        ReviewSession.load(snapshot, draft)
    draft.write_bytes(b" " * (MAX_DRAFT_BYTES + 1))
    with pytest.raises(ValueError, match="8 MiB"):
        ReviewSession.load(snapshot, draft)
    payload = session.payload()
    session.update(
        row_id(session, "INV-A"), {"Period Charge (£)": 120}, True, "Correct amount", "verified"
    )
    with pytest.raises(ValueError, match="different or changed"):
        ReviewSession.restore(analyse_session(session, path), payload)


def items():
    base = {
        "Account": "000123",
        "Fuel": "Electricity",
        "Meter": "M1",
        "Entry Type": "New Bill",
        "Period From": "01/01/2024",
        "Period To": "31/01/2024",
    }
    return [
        {
            "id": "a",
            "included": True,
            "record": {
                **base,
                "Invoice #": "INV-1",
                "Date": "01/02/2024",
                "Period Charge (£)": 100,
            },
        },
        {
            "id": "b",
            "included": True,
            "record": {
                **base,
                "Invoice #": "INV-2",
                "Date": "02/02/2024",
                "Period Charge (£)": 120,
            },
        },
        {
            "id": "c",
            "included": True,
            "record": {
                **base,
                "Invoice #": "CR-1",
                "Date": "02/02/2024",
                "Entry Type": "Credit",
                "Details": "Cancels INV-1",
                "Transaction Amount (£)": -100,
            },
        },
    ]


def decision(older="a", newer="b", **overrides):
    return {
        "older": older,
        "newer": newer,
        "kind": "replacement",
        "status": "confirmed",
        "reason": "Checked original PDFs",
        **overrides,
    }


def test_revision_suggestions_require_identity_periods_or_explicit_credit_reference():
    data = items()
    model = build_revisions(data)
    assert {(r["older"], r["newer"], r["kind"], r["status"]) for r in model["relationships"]} == {
        ("a", "b", "replacement", "suggested"),
        ("a", "c", "cancellation_credit", "suggested"),
    }
    data[1]["record"]["Period To"] = "15/02/2024"
    assert any("partial" in r["basis"] for r in build_revisions(data)["relationships"])
    data[1]["record"]["Meter"] = "OTHER"
    assert all(r["newer"] != "b" for r in build_revisions(data)["relationships"])
    data[2]["record"]["Account"] = ""
    assert not build_revisions(data)["relationships"]


@pytest.mark.parametrize(
    "mutation",
    [
        "account",
        "unknown",
        "meter",
        "fuel",
        "register",
        "dates",
        "cycle",
        "kind",
        "status",
        "reason",
        "endpoints",
        "duplicate",
    ],
)
def test_revision_validation_never_crosses_identity_or_corrupt_graph(mutation):
    data = items()
    decisions = [decision()]
    if mutation in {"account", "unknown", "meter", "fuel", "register"}:
        field = {
            "account": "Account",
            "unknown": "Account",
            "meter": "Meter",
            "fuel": "Fuel",
            "register": "Register",
        }[mutation]
        data[0]["record"][field] = "A"
        data[1]["record"][field] = "" if mutation == "unknown" else "B"
    elif mutation == "dates":
        data[1]["record"]["Date"] = "bad date"
    elif mutation == "cycle":
        data[1]["record"]["Date"] = data[0]["record"]["Date"]
        decisions.append(decision("b", "a"))
    elif mutation in {"kind", "status"}:
        decisions[0][mutation] = []
    elif mutation == "reason":
        decisions[0]["reason"] = " "
    elif mutation == "endpoints":
        decisions[0]["newer"] = "missing"
    else:
        decisions *= 2
    with pytest.raises(ValueError):
        validate_decisions(data, decisions)


def test_manual_nonoverlap_rejected_links_and_caps_are_explicit():
    data = items()
    data[1]["record"]["Period From"] = "01/03/2024"
    data[1]["record"]["Period To"] = "31/03/2024"
    model = build_revisions(data, [decision()], limit=0)
    assert model["relationships"][0]["status"] == "confirmed"
    assert model["relationships"][0]["charge_difference"] == 20
    assert model["truncated"] is True
    model = build_revisions(items(), [decision(status="rejected", reason="Separate meter review")])
    assert next(r for r in model["relationships"] if r["newer"] == "b")["status"] == "rejected"


def test_confirmation_changes_no_financial_totals_and_blocks_conflicting_edits(review_pack):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    a, b, c = [row_id(session, "INV-" + k) for k in "ABC"]
    session.set_revision(a, b, "replacement", "confirmed", "Original and replacement checked")
    session.set_revision(b, c, "replacement", "confirmed", "Second replacement checked")
    after = analyse_session(session, path)
    assert after.summary == snapshot.summary
    assert after.input_fingerprint != snapshot.input_fingerprint
    model = build_report_view(after, RenderSelection()).revisions
    assert model is not None
    assert model["relationships"][0]["status"] == "confirmed"
    before = session.payload()
    with pytest.raises(ValueError, match="same known account"):
        session.update(b, {"Account": "OTHER"}, True, "Different account", "verified")
    assert session.payload() == before
    session.set_revision(a, b, "replacement", "suggested", "")
    session.undo()
    assert session.payload() == before


def test_native_worker_exports_all_formats_retains_originals_and_reinstates_sources(
    review_pack, tmp_path
):
    path, snapshot, ids = review_pack
    result_path = path.parent / "result.json"
    save_result(snapshot, result_path, result_id="native-base")
    original_files = {
        p.name: p.read_bytes() for p in path.parent.joinpath("evidence_files").iterdir()
    }
    session = ReviewSession(snapshot)
    a, b, c = [row_id(session, "INV-" + k) for k in "ABC"]
    session.update(a, {"Period Charge (£)": 125}, True, "Source PDF verified", "verified")
    session.update(c, {}, False, "Not relevant", "needs_information")
    session.set_revision(a, b, "replacement", "confirmed", "Same consumption, revised invoice")
    directory = tmp_path / "reviewed"
    directory.mkdir()
    job = DesktopJob(
        snapshot.prepared.extraction.request,
        str(directory),
        "Reviewed",
        ("xlsx", "pdf", "docx", "html"),
        ("detailed_findings",),
        "review",
        str(result_path),
        freeze_json(session.payload()),
    )
    outcomes: list[RunOutcome] = []
    worker = RunWorker(job)
    worker.outcome_ready.connect(outcomes.append)
    worker.run()
    assert outcomes
    assert outcomes[0].status in {"complete", "review_required"}, outcomes[0].error
    accepted = load_result(directory / "result.json")
    assert accepted.summary.charges == Decimal("175.00")
    assert accepted.prepared.extraction.request.as_of == snapshot.prepared.extraction.request.as_of
    assert {
        p.name: p.read_bytes() for p in path.parent.joinpath("evidence_files").iterdir()
    } == original_files
    wb = load_workbook(directory / "Reviewed.xlsx")
    assert wb["Bill revisions"]["A5"].value == "Confirmed"
    assert wb["Bill revisions"]["O5"].value == session.record(a)["Attachment Name"]
    assert wb["Bill revisions"].freeze_panes == "D5"
    wb.close()
    assert "Bill revision review" in (directory / "Reviewed_Report.html").read_text()
    assert "Same consumption, revised invoice" in (directory / "Reviewed_Report.html").read_text()
    doc = Document(directory / "Reviewed_Report.docx")
    assert "Bill revision review" in "\n".join(p.text for p in doc.paragraphs)
    assert any(
        "INV-A" in cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    import pdfplumber

    with pdfplumber.open(directory / "Reviewed_Report.pdf") as pdf:
        assert "Bill revision review" in "\n".join(p.extract_text() or "" for p in pdf.pages)
    review = build_evidence_review(accepted, str(directory))
    excluded = next(entry for entry in review.entries if entry.row_id == c)
    assert excluded.category == "Excluded"
    assert excluded.sources[0].path
    assert (
        ids[1]
        in object_value(accepted.prepared.extraction.legacy_state)[
            "review_excluded_source_manifest"
        ]
    )
    restored = ReviewSession(accepted)
    restored.update(c, {}, True, "Now relevant after checking original", "verified")
    directory2 = tmp_path / "reinstated"
    directory2.mkdir()
    next_job = replace(
        job,
        directory=str(directory2),
        formats=("xlsx",),
        input_path=str(directory / "result.json"),
        review_draft=freeze_json(restored.payload()),
    )
    next_worker = RunWorker(next_job)
    next_outcomes: list[RunOutcome] = []
    next_worker.outcome_ready.connect(next_outcomes.append)
    next_worker.run()
    assert next_outcomes[0].status in {"complete", "review_required"}, next_outcomes[0].error
    assert original_files["excluded.pdf"] in {
        p.read_bytes() for p in (directory2 / "evidence_files").iterdir()
    }
    assert load_result(directory2 / "result.json").summary.charges == Decimal("245.00")
    assert json.loads((directory / "run.json").read_text())["status"] in {
        "complete",
        "review_required",
    }


def test_excel_changes_invalidate_previous_verification(review_pack, tmp_path):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    key = row_id(session, "INV-A")
    session.update(key, {}, True, "Checked original source", "verified")
    reviewed = analyse_session(session, path)
    from edf_bill_fetcher.io.writers.export import render_xlsx

    workbook = tmp_path / "verified.xlsx"
    render_xlsx(build_report_view(reviewed, RenderSelection()), str(workbook))
    amend(
        workbook, "INV-A", **{"Period Charge (£)": 140, "Reason for change": "Further correction"}
    )
    _, state, _ = load_review_workbook(str(workbook))
    assert state["human_review_statuses"][key]["status"] == "unreviewed"
    assert any(
        change["field"] == "Review status"
        and change["before"] == "verified"
        and change["after"] == "unreviewed"
        for change in state["review_history"]
    )


def test_confirming_cancelled_bill_credit_and_replacement_never_double_adjusts():
    from tests.test_qt_desktop import result_snapshot

    baseline = result_snapshot()
    data = [item["record"] for item in items()]
    data[0]["Amount Kind"] = data[1]["Amount Kind"] = "Transaction"
    data[0]["Transaction Amount (£)"], data[1]["Transaction Amount (£)"] = 100, 120
    data[2]["Amount Kind"] = "Transaction"
    snapshot = analyse_prepared(
        prepare_snapshot(capture_legacy(data, {}, baseline.prepared.extraction.request))
    )
    session = ReviewSession(snapshot)
    a, b, c = [row_id(session, invoice) for invoice in ("INV-1", "INV-2", "CR-1")]
    session.set_revision(a, b, "replacement", "confirmed", "Replacement checked")
    session.set_revision(a, c, "cancellation_credit", "confirmed", "Credit offsets original charge")
    reviewed = analyse_session(session, "result.json")
    assert reviewed.summary == snapshot.summary
    assert reviewed.summary.charges == Decimal("220.00")
    assert reviewed.summary.net_credits == Decimal("100.00")
    assert len(rows_value(reviewed.prepared.records)) == 3


@pytest.mark.parametrize(
    "mutation", ["version", "membership", "status", "amount", "inclusion", "reason"]
)
def test_saved_revision_model_is_strict_and_never_recalculated(review_pack, mutation):
    _, snapshot, _ = review_pack
    from edf_bill_fetcher.io.persistence.result_codec import _validate_models

    models = object_value(snapshot.models)
    model = models["sap"]["payload"]["bill_revisions"]
    link = model["relationships"][0]
    if mutation == "version":
        model["version"] = True
    elif mutation == "membership":
        model["relationships"].append(copy.deepcopy(link))
    elif mutation == "status":
        link["status"] = "automatically confirmed"
    elif mutation == "amount":
        link["charge_difference"] = True
    elif mutation == "inclusion":
        link["earlier_included"] = "yes"
    else:
        link["status"] = "confirmed"
        link["reason"] = ""
    with pytest.raises(ValueError):
        _validate_models(models)


def test_long_revision_reasons_can_span_pdf_pages(review_pack, tmp_path):
    import pdfplumber

    from edf_bill_fetcher.application.run_service import default_renderers

    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    session.set_revision(
        row_id(session, "INV-A"),
        row_id(session, "INV-B"),
        "replacement",
        "confirmed",
        "Checked the original source invoice. " * 100 + "END_OF_REASON",
    )
    reviewed = analyse_session(session, path)
    target = tmp_path / "long-reason.pdf"
    default_renderers()["pdf"](
        build_report_view(reviewed, RenderSelection(("pdf",), ("detailed_findings",))), str(target)
    )
    with pdfplumber.open(target) as pdf:
        assert len(pdf.pages) > 1
        assert "END_OF_REASON" in "\n".join(page.extract_text() or "" for page in pdf.pages)


def test_draft_size_limit_applies_to_the_actual_saved_bytes(review_pack, tmp_path, monkeypatch):
    from edf_bill_fetcher.application import review_session
    from edf_bill_fetcher.domain.json_values import canonical_text

    _, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    session.update(row_id(session, "INV-A"), {}, True, "原本を確認" * 100, "verified")
    target = tmp_path / "draft.json"
    session.save(target)
    original = target.read_bytes()
    canonical_size = len(canonical_text(session.payload()).encode())
    assert len(original) > canonical_size
    monkeypatch.setattr(review_session, "MAX_DRAFT_BYTES", canonical_size)
    with pytest.raises(ValueError, match="8 MiB"):
        session.save(target)
    assert target.read_bytes() == original


def test_correction_that_merges_duplicates_keeps_review_reopenable(review_pack):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    a, b = row_id(session, "INV-A"), row_id(session, "INV-B")
    session.update(a, {}, True, "Verified first extraction", "verified")
    reference = session.record(a)
    session.update(
        b,
        {field: reference[field] for field in ("Invoice #", "Period Charge (£)", "Balance (£)")},
        True,
        "Correct duplicate extraction against source",
        "verified",
    )
    merged = analyse_session(session, path)
    active = ReviewSession(merged)
    assert len(active.rows) == 2
    state = object_value(merged.prepared.extraction.legacy_state)
    assert len(state["human_review_statuses"]) == 1
    assert len(state["archived_review_statuses"]) == 1
    assert next(iter(state["archived_review_statuses"].values()))["status"] == "verified"
    assert any(
        change["after"] == "archived_after_deduplication" for change in state["review_history"]
    )
    assert [r.raw for r in merged.prepared.extraction.observations] == [
        r.raw for r in snapshot.prepared.extraction.observations
    ]
