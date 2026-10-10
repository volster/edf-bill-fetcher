"""Desktop workflow contracts using actual Qt widgets, signals and worker threads."""

import json
import os
import shutil
from datetime import date, datetime, timezone
from pathlib import Path
from threading import Event
from time import monotonic, sleep

import pandas as pd
import pytest

# Set the backend before the first QApplication; no desktop session is needed.
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QCoreApplication, QEvent, Qt, QTimer  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402
from PySide6.QtWidgets import (  # noqa: E402
    QApplication,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QLineEdit,
    QPlainTextEdit,
    QScrollArea,
)

from edf_bill_fetcher.application.analysis_service import analyse_prepared  # noqa: E402
from edf_bill_fetcher.application.legacy_adapter import capture_legacy  # noqa: E402
from edf_bill_fetcher.application.preparation_service import prepare_snapshot  # noqa: E402
from edf_bill_fetcher.application.run_service import RunOutcome  # noqa: E402
from edf_bill_fetcher.domain.json_values import freeze_json, object_value  # noqa: E402
from edf_bill_fetcher.domain.requests import RunRequest  # noqa: E402
from edf_bill_fetcher.io.persistence.result_store import load_result, save_result  # noqa: E402
from edf_bill_fetcher.ui import qt_app, qt_jobs  # noqa: E402
from edf_bill_fetcher.ui.qt_app import EvidenceWindow, ReportDialog  # noqa: E402
from edf_bill_fetcher.ui.qt_theme import apply_theme  # noqa: E402


@pytest.fixture(scope="session")
def application():
    app = QApplication.instance() or QApplication([])
    assert isinstance(app, QApplication)
    app.setQuitOnLastWindowClosed(False)
    apply_theme(app)
    yield app
    for widget in app.topLevelWidgets():
        widget.close()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    app.processEvents()


@pytest.fixture
def font_metrics(application, request):
    """Exercise native platform fonts at normal and larger text sizes."""
    stylesheet = application.styleSheet()
    pixels = getattr(request, "param", 13)
    application.setStyleSheet(stylesheet + f"\nQWidget {{ font-size: {pixels}px; }}")
    yield
    application.setStyleSheet(stylesheet)
    application.processEvents()


@pytest.fixture
def window(application, font_metrics, tmp_path):
    widget = EvidenceWindow(tmp_path / "preferences.json")
    widget.show()
    application.processEvents()
    yield widget
    if widget.worker is not None:
        widget.worker.cancel_event.set()
        assert widget.worker.wait(10000), "Worker did not finish cleanup"
        application.processEvents()
    widget.close()
    widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    application.processEvents()


def wait_until(application, predicate):
    deadline = monotonic() + 15
    while not predicate() and monotonic() < deadline:
        application.processEvents()
        sleep(0.01)  # Release the GIL so Python worker code can make progress.
    application.processEvents()
    assert predicate(), "Qt operation did not finish within 15 seconds"


def set_text(window, key, value):
    field = window.fields[key]
    assert isinstance(field, QLineEdit)
    field.setText(str(value))


def source_folder(window, tmp_path):
    folder = tmp_path / "sources"
    folder.mkdir(exist_ok=True)
    window.sources["pdf_dir"].entry.setText(str(folder))
    return folder


def result_snapshot():
    request = RunRequest(
        "original",
        datetime(2026, 1, 31, tzinfo=timezone.utc),
        date(2026, 1, 31),
        (),
        freeze_json({"as_of": "2026-01-31", "use_dedup": True}),
    )
    records = [
        {
            "Account": "A-100",
            "Fuel": "Electricity",
            "Date": "01/01/2026",
            "Entry Type": "New Bill",
            "Amount Kind": "Balance",
            "Amount (£)": 100,
            "Period Charge (£)": 60,
            "Invoice #": "B1",
            "Period From": "01/12/2025",
            "Period To": "31/12/2025",
        }
    ]
    return analyse_prepared(prepare_snapshot(capture_legacy(records, {}, request)))


def execute(job, application):
    outcomes: list[RunOutcome] = []
    worker = qt_jobs.RunWorker(job)
    worker.outcome_ready.connect(outcomes.append)
    worker.start()
    try:
        wait_until(application, lambda: not worker.isRunning() and bool(outcomes))
        assert worker.wait(1000)
        return outcomes[0]
    finally:
        worker.cancel_event.set()
        assert worker.wait(10000)


@pytest.mark.parametrize(("width", "height"), [(1160, 720), (820, 600), (920, 640)])
@pytest.mark.parametrize("font_metrics", [13, 20], indirect=True)
def test_action_stays_visible_and_source_steps_adapt(window, application, width, height):
    window.resize(width, height)
    application.processEvents()
    assert window.width() == width
    assert window.run_button.isVisible()
    corner = window.run_button.mapTo(window, window.run_button.rect().bottomRight())
    assert window.rect().contains(corner)
    assert not window.run_button.isEnabled()
    assert window.source_grid.getItemPosition(1)[:2] == ((0, 1) if width >= 1080 else (1, 0))
    for index in (1, 2, 0):
        window.nav_buttons[index].click()
        application.processEvents()
        assert window.pages.currentIndex() == index
        assert window.run_button.isVisible()
        page = window.pages.currentWidget()
        wait_until(application, lambda page=page: page.horizontalScrollBar().maximum() == 0)
        assert page.horizontalScrollBar().maximum() == 0, (
            f"Page {index} at window width {width}: viewport={page.viewport().width()}, "
            f"minimum content width={page.widget().minimumSizeHint().width()}"
        )
    window._navigate(2)
    window.run_button.click()
    assert window.pages.currentIndex() == 0


def test_browse_buttons_capture_the_correct_source_and_cancel_keeps_it(window, monkeypatch):
    directories = iter(["/chosen/pdfs", "", "/chosen/mail"])
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: next(directories))
    window.sources["pdf_dir"].browse_button.click()
    window.sources["pdf_dir"].browse_button.click()
    window.sources["eml_msg_dir"].browse_button.click()
    assert window.sources["pdf_dir"].value() == "/chosen/pdfs"
    assert window.sources["eml_msg_dir"].value() == "/chosen/mail"
    files = iter([("/chosen/archive.ost", ""), ("/chosen/history.htm", ""), ("", "")])
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: next(files))
    window.sources["pst_path"].browse_button.click()
    window.sources["htm_path"].browse_button.click()
    window.sources["htm_path"].browse_button.click()
    assert window.sources["pst_path"].value() == "/chosen/archive.ost"
    assert window.sources["htm_path"].value() == "/chosen/history.htm"
    assert window.run_button.isEnabled()
    assert window.sources["pdf_dir"].entry.accessibleName() == "PDF bills"


def test_output_destination_is_shared_between_workflows(window, monkeypatch):
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *args: "/chosen/output")
    window._pick_output()
    report_entries = window.pages.widget(1).findChildren(QLineEdit)
    output = next(field for field in report_entries if field.accessibleName() == "Save to")
    assert output.text() == "/chosen/output"
    output.setText("/changed/output")
    assert window._field_value("output_folder") == "/changed/output"


def test_report_picker_and_saved_choices_update_both_summaries(window, monkeypatch):
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: ("/saved/result.json", ""))
    window._pick_report_input()
    assert window.report_input.value() == "/saved/result.json"

    def save(dialog):
        dialog.formats.button(3).setChecked(True)
        dialog._select(False)
        dialog.section_checks["payment"].setChecked(True)
        dialog.accept()
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(ReportDialog, "exec", save)
    window._open_report_settings()
    assert window.report_options == {"format": "html", "sections": ["payment"]}
    assert window.report_summary.text() == "HTML · 1 section"
    assert window.saved_report_summary.text() == window.report_summary.text()
    monkeypatch.setattr(ReportDialog, "exec", lambda self: QDialog.DialogCode.Rejected)
    window._open_report_settings()
    assert window.report_options == {"format": "html", "sections": ["payment"]}


@pytest.mark.parametrize("font_metrics", [13, 20], indirect=True)
def test_report_dialog_requires_content_remembers_choices_and_scrolls(application, font_metrics):
    options = {"format": "docx", "sections": ["payment", "timeline"]}
    dialog = ReportDialog(options)
    dialog.resize(560, 460)
    dialog.show()
    application.processEvents()
    assert dialog.formats.checkedId() == 2
    assert dialog.section_count.text() == "2 of 17 sections"
    assert dialog.section_grid.columnCount() >= 1
    page = dialog.findChild(QScrollArea)
    assert page is not None
    for width in (560, 900, 560):
        dialog.resize(width, 460)
        wait_until(application, lambda: page.horizontalScrollBar().maximum() == 0)
        for index, radio in enumerate(dialog.formats.buttons()):
            page.ensureWidgetVisible(radio)
            application.processEvents()
            QTest.mouseClick(radio, Qt.MouseButton.LeftButton)
            assert dialog.formats.checkedId() == index
    dialog.formats.button(2).setChecked(True)
    dialog._select(False)
    actions = dialog.findChild(QDialogButtonBox)
    assert actions is not None
    save = actions.button(QDialogButtonBox.StandardButton.Save)
    save.click()
    assert dialog.selection is None
    assert "at least one" in dialog.validation.text()
    assert dialog.isVisible()
    dialog.section_checks["exec_summary"].setChecked(True)
    save.click()
    assert dialog.selection == {"format": "docx", "sections": ["exec_summary"]}
    assert options == {"format": "docx", "sections": ["payment", "timeline"]}
    dialog.deleteLater()
    invalid = ReportDialog({"format": "unknown"})
    assert invalid.formats.checkedId() == 0
    invalid.reject()


def test_preferences_restore_legacy_aliases_and_keep_sensitive_sources_out(application, tmp_path):
    path = tmp_path / "preferences.json"
    path.write_text(
        json.dumps(
            {
                "gui_state": {
                    "use_acc_filt": True,
                    "use_reading_class": False,
                    "acc_num": "A-100",
                    "output_name": "case.v2.xlsx",
                    "min_amount": 12.50,
                    "analysis_min": float("nan"),
                    "auto_generate_report": True,
                },
                "report_options": {"format": "html", "sections": ["payment", "payment"]},
            }
        )
    )
    widget = EvidenceWindow(path)
    try:
        assert widget.checks["use_acc_filter"].isChecked()
        assert not widget.checks["use_reading_classification"].isChecked()
        assert widget._field_value("output_name") == "case.v2"
        assert widget._field_value("min_amount") == 12.5
        assert widget._field_value("analysis_min") == 0
        assert widget.include_reports.isChecked()
        assert widget.report_options == {"format": "html", "sections": ["payment"]}
        source_folder(widget, tmp_path)
        widget.save_settings()
        stored = json.loads(path.read_text())
        assert stored["gui_state"]["use_acc_filter"] is True
        assert stored["gui_state"]["output_name"] == "case.v2.xlsx"
        assert "sources" not in path.read_text()
        assert not any(key in stored["gui_state"] for key in widget.sources)
        if os.name != "nt":
            assert path.stat().st_mode & 0o777 == 0o600
    finally:
        widget.close()


@pytest.mark.parametrize(
    "content",
    [
        "broken",
        "[]",
        '{"gui_state":[]}',
        '{"report_options":{"format":"bad"}}',
        '{"report_options":{"format":[]}}',
        '{"report_options":{"format":"pdf","sections":[[]]}}',
    ],
)
def test_bad_preferences_do_not_break_the_desktop(application, tmp_path, content):
    path = tmp_path / "bad.json"
    path.write_text(content)
    widget = EvidenceWindow(path)
    assert widget._field_value("output_name") == "EDF_Dispute_Evidence"
    assert widget.report_options["format"] == "both"
    widget.close()


def test_preferences_dependants_keep_values_when_disabled(window):
    window.checks["use_acc_filter"].setChecked(True)
    set_text(window, "acc_num", "A-100")
    window.checks["use_acc_filter"].setChecked(False)
    assert not window.fields["acc_num"].isEnabled()
    assert window._field_value("acc_num") == "A-100"
    window.checks["amalgamate_duplicates"].setChecked(True)
    window.checks["use_dedup"].setChecked(False)
    assert not window.checks["save_dups"].isEnabled()
    assert not window.checks["amalgamate_duplicates"].isEnabled()
    assert window.checks["amalgamate_duplicates"].isChecked()


def test_job_capture_freezes_settings_and_reserves_unique_outputs(window, tmp_path):
    source_folder(window, tmp_path)
    window.include_reports.setChecked(True)
    set_text(window, "output_folder", tmp_path / "exports")
    first = window.capture_job()
    set_text(window, "acc_num", "changed after capture")
    window.report_options["sections"] = ["payment"]
    second = window.capture_job()
    assert first.directory != second.directory
    assert first.request.run_id != second.request.run_id
    assert set(first.outputs) == {"xlsx", "pdf", "docx", "html"}
    assert all(Path(path).parent == Path(first.directory) for path in first.outputs.values())
    assert object_value(first.request.config)["acc_num"] == ""
    assert len(first.sections) == 17
    assert second.sections == ("payment",)
    assert first.request.created_at.tzinfo == timezone.utc


def test_default_output_folder_is_beside_source_not_inside_it(window, tmp_path):
    folder = source_folder(window, tmp_path)
    job = window.capture_job()
    assert Path(job.directory).parent == folder.parent
    assert not Path(job.directory).is_relative_to(folder)


@pytest.mark.parametrize("name", ["", "../escape", "CON", "bad?name", "case."])
def test_unsafe_pack_names_are_rejected_before_starting(window, tmp_path, name):
    source_folder(window, tmp_path)
    set_text(window, "output_name", name)
    with pytest.raises(ValueError, match="pack name"):
        window.capture_job()


def test_source_validation_is_inline_and_does_not_start_a_worker(window, tmp_path):
    with pytest.raises(ValueError, match="at least one"):
        window.capture_job()
    window.sources["pdf_dir"].entry.setText(str(tmp_path / "missing"))
    window.run_button.click()
    assert "Source not found" in window.feedback.text()
    assert window.feedback.isVisible()
    assert window.feedback.textFormat() == Qt.TextFormat.PlainText
    assert window.worker is None
    assert not window._busy
    source_folder(window, tmp_path)
    window.checks["use_acc_filter"].setChecked(True)
    with pytest.raises(ValueError, match="account number"):
        window.capture_job()
    window.checks["use_acc_filter"].setChecked(False)
    window.sources["eml_msg_dir"].entry.setText(str(tmp_path / "sources"))
    set_text(window, "domain_filter", "")
    with pytest.raises(ValueError, match="sender domain"):
        window.capture_job()


def test_email_and_ost_capture_choose_the_right_collectors(window, tmp_path):
    folder = source_folder(window, tmp_path)
    archive = tmp_path / "archive.OST"
    archive.touch()
    history = tmp_path / "account.htm"
    history.touch()
    window.sources["pst_path"].entry.setText(str(archive))
    window.sources["htm_path"].entry.setText(str(history))
    window.sources["eml_msg_dir"].entry.setText(str(folder))
    assert [s.kind for s in window.capture_job().request.sources] == [
        "pdf_dir",
        "htm",
        "ost",
        "eml_dir",
        "msg_dir",
    ]


def test_report_validation_and_optional_workbook(window, tmp_path):
    window._navigate(1)
    with pytest.raises(ValueError, match="existing saved analysis"):
        window.capture_job()
    path = tmp_path / "analysis.json"
    path.touch()
    window.report_input.entry.setText(str(path))
    window.report_excel.setChecked(True)
    job = window.capture_job()
    assert job.mode == "saved_result"
    assert job.input_path == str(path)
    assert job.formats == ("xlsx", "pdf", "docx", "html")


def test_real_pdf_extraction_delivers_all_outputs_and_original_evidence(
    window, application, tmp_path
):
    folder = source_folder(window, tmp_path)
    original = folder / "bill.pdf"
    shutil.copy2(Path(__file__).parent / "fixtures/sample_bill.pdf", original)
    window.include_reports.setChecked(True)
    window.report_options = {"format": "both", "sections": ["cover", "payment", "evidence_index"]}
    window.run_button.click()
    assert window._busy
    assert not window.pages.isEnabled()
    wait_until(application, lambda: window.last_outcome is not None and not window._busy)
    outcome = window.last_outcome
    assert outcome.status == "complete", outcome.error
    assert window.worker is None
    assert window.pages.isEnabled()
    assert not window.progress.isVisible()
    assert window.status.text() == "Your files are ready"
    assert window.open_folder_button.isVisible()
    manifest = object_value(outcome.manifest)
    assert set(manifest["artifacts"]) == {"xlsx", "pdf", "docx", "html"}
    directory = Path(window.current_job.directory)
    assert load_result(directory / "result.json") == outcome.snapshot
    copied = list((directory / "evidence_files").glob("*.pdf"))
    assert len(copied) == 1
    assert copied[0].read_bytes() == original.read_bytes()
    assert (directory / "evidence_index.docx").is_file()


def test_saved_result_renders_without_reanalysis_or_modifying_original(
    window, application, tmp_path, monkeypatch
):
    path = tmp_path / "result.json"
    snapshot = result_snapshot()
    save_result(snapshot, path, result_id="accepted")
    original = path.read_bytes()

    def forbidden(*args, **kwargs):
        raise AssertionError("Accepted analyses must not be recomputed")

    monkeypatch.setattr("edf_bill_fetcher.application.run_service.analyse_prepared", forbidden)
    monkeypatch.setattr("edf_bill_fetcher.application.run_service.prepare_snapshot", forbidden)
    window._navigate(1)
    window.report_input.entry.setText(str(path))
    window.report_options = {"format": "html", "sections": ["payment"]}
    job = window.capture_job()
    outcome = execute(job, application)
    assert outcome.status == "complete", outcome.error
    assert outcome.snapshot == snapshot
    assert path.read_bytes() == original
    assert load_result(Path(job.directory) / "result.json") == snapshot
    assert "Last Payment Amount" in Path(job.outputs["html"]).read_text()


@pytest.mark.parametrize("change", ["moved", "missing", "modified"])
def test_saved_workbook_links_resolve_moved_evidence_without_recalculating(
    window, application, tmp_path, monkeypatch, change
):
    import openpyxl

    source = source_folder(window, tmp_path)
    shutil.copy2(Path(__file__).parent / "fixtures/sample_bill.pdf", source / "bill.pdf")
    window.run_button.click()
    wait_until(application, lambda: not window._busy)
    original_snapshot = window.last_outcome.snapshot
    old_directory = Path(window.current_job.directory)
    relocated = tmp_path / "relocated"
    shutil.move(str(old_directory), relocated)
    originals = list((relocated / "evidence_files").glob("*.pdf"))
    assert len(originals) == 1
    if change == "missing":
        originals[0].unlink()
    elif change == "modified":
        originals[0].write_bytes(b"different evidence")
    original_result = (relocated / "result.json").read_bytes()

    def forbidden(*args, **kwargs):
        raise AssertionError("Re-rendering must preserve accepted calculations")

    monkeypatch.setattr("edf_bill_fetcher.application.run_service.analyse_prepared", forbidden)
    monkeypatch.setattr("edf_bill_fetcher.application.run_service.prepare_snapshot", forbidden)
    window._navigate(1)
    window.report_input.entry.setText(str(relocated / "result.json"))
    window.report_options = {"format": "html", "sections": ["payment"]}
    window.report_excel.setChecked(True)
    job = window.capture_job()
    outcome = execute(job, application)
    assert outcome.snapshot == original_snapshot
    assert (relocated / "result.json").read_bytes() == original_result
    workbook = openpyxl.load_workbook(job.outputs["xlsx"])
    try:
        ws = workbook["EDF Evidence Report"]
        column = [cell.value for cell in ws[1]].index("Attachment Name") + 1
        cell = ws.cell(2, column)
        if change == "moved":
            assert outcome.status == "complete", outcome.error
            assert (Path(job.directory) / cell.hyperlink.target).resolve() == originals[0]
        else:
            assert outcome.status == "review_required", outcome.error
            assert cell.hyperlink is None
            assert object_value(outcome.manifest)["warnings"]
    finally:
        workbook.close()


@pytest.mark.parametrize("content", ['{"version":999}', '{"amount":NaN}', "not JSON"])
def test_invalid_saved_result_publishes_failure_not_success(window, application, tmp_path, content):
    path = tmp_path / "bad.json"
    path.write_text(content)
    window._navigate(1)
    window.report_input.entry.setText(str(path))
    job = window.capture_job()
    outcome = execute(job, application)
    assert outcome.status == "failed"
    assert outcome.snapshot is None
    manifest = json.loads((Path(job.directory) / "run.json").read_text())
    assert manifest["stage"] == "loading_input"
    assert manifest["artifacts"] == {}
    assert manifest["result"] is None
    assert path.read_text() == content


def test_workbook_import_preserves_date_and_shows_provenance_limit(window, application, tmp_path):
    path = tmp_path / "legacy.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame(
            [{"Date": "01/01/2026", "Entry Type": "New Bill", "Amount (£)": 100}]
        ).to_excel(writer, sheet_name="EDF Evidence Report", index=False)
        pd.DataFrame([["as_of", "2026-01-31"]]).to_excel(
            writer, sheet_name="Provenance", index=False, header=False
        )
        pd.DataFrame([{"Amount (£)": 2}]).to_excel(
            writer, sheet_name="Filtered (Below Min)", index=False
        )
    window._navigate(1)
    window.report_input.entry.setText(str(path))
    window.report_options = {"format": "html", "sections": ["payment"]}
    outcome = execute(window.capture_job(), application)
    assert outcome.status == "review_required", outcome.error
    assert outcome.snapshot.prepared.extraction.request.as_of == date(2026, 1, 31)
    state = object_value(outcome.snapshot.prepared.extraction.legacy_state)
    assert state["provenance"] == "lossy_workbook_import"
    assert len(state["filtered_records"]) == 1
    assert "original observations" in state["error_log"][0]


def test_export_failure_keeps_saved_analysis_and_reports_partial_artifacts(
    window, application, tmp_path, monkeypatch
):
    path = tmp_path / "result.json"
    save_result(result_snapshot(), path, result_id="accepted")
    window._navigate(1)
    window.report_input.entry.setText(str(path))
    window.report_options = {"format": "both", "sections": ["payment"]}

    def fail(*args):
        raise OSError("Word writer failed")

    monkeypatch.setattr("edf_bill_fetcher.io.reporters.docx_report.render_docx", fail)
    window.run_button.click()
    wait_until(application, lambda: window.last_outcome is not None and not window._busy)
    assert window.last_outcome.status == "failed"
    assert set(object_value(window.last_outcome.manifest)["artifacts"]) == {"pdf"}
    assert "saved analysis is available" in window.feedback.text()
    assert window.status.text() == "Run needs attention"
    assert window.details_button.isVisible()


@pytest.mark.parametrize("close", [False, True])
def test_cancellation_and_window_close_wait_for_worker_cleanup(
    window, application, tmp_path, monkeypatch, close
):
    source_folder(window, tmp_path)
    entered, release, cleaned = Event(), Event(), Event()

    class Engine:
        def __init__(self, config, log, progress, cancel):
            self.records = []
            self.cancel_event = cancel
            self.progress = progress

        def crawl_local_pdfs(self, path):
            self.progress(1, 2, "Reading controlled source")
            entered.set()
            assert release.wait(10)

        def __getstate__(self):
            return {"error_log": []}

        def cleanup(self):
            cleaned.set()

    monkeypatch.setattr("edf_bill_fetcher.collectors.engine.EvidenceEngine", Engine)
    window.run_button.click()
    wait_until(application, entered.is_set)
    first_worker = window.worker
    try:
        if close:
            window.close()
            assert window.isVisible()
        else:
            window.run_button.click()
        assert first_worker.cancel_event.is_set()
        assert window.worker is first_worker
        assert not window.run_button.isEnabled()
        assert not cleaned.is_set()
    finally:
        release.set()
    wait_until(application, lambda: not window._busy)
    assert cleaned.is_set()
    assert window.last_outcome.status == "cancelled"
    assert object_value(window.last_outcome.manifest)["artifacts"] == {}
    assert window.isVisible() is (not close)


def test_empty_source_has_no_reports_and_a_truthful_status(window, application, tmp_path):
    source_folder(window, tmp_path)
    window.run_button.click()
    wait_until(application, lambda: not window._busy)
    assert window.last_outcome.status == "empty"
    assert window.status.text() == "No billing records found"
    assert "0 exported files" in window.readiness.text()
    assert window.details_button.isVisible()


def test_engine_constructor_failure_is_recorded(window, application, tmp_path, monkeypatch):
    source_folder(window, tmp_path)

    def fail(*args):
        raise RuntimeError("cannot initialise parser")

    monkeypatch.setattr("edf_bill_fetcher.collectors.engine.EvidenceEngine", fail)
    outcome = execute(window.capture_job(), application)
    assert outcome.status == "failed"
    assert "cannot initialise parser" in outcome.error
    assert object_value(outcome.manifest)["stage"] == "collecting"


def test_precancelled_worker_reports_no_exports(window, tmp_path):
    source_folder(window, tmp_path)
    worker = qt_jobs.RunWorker(window.capture_job())
    results: list[RunOutcome] = []
    worker.outcome_ready.connect(results.append)
    worker.cancel_event.set()
    worker.run()
    assert results[0].status == "cancelled"
    assert object_value(results[0].manifest)["artifacts"] == {}


@pytest.mark.parametrize(
    ("status", "has_result", "files"),
    [("review_required", True, 1), ("cancelled", True, 1), ("failed", False, 0)],
)
@pytest.mark.parametrize("font_metrics", [13, 20], indirect=True)
def test_terminal_status_does_not_hide_partial_outputs(
    window, application, tmp_path, status, has_result, files
):
    source_folder(window, tmp_path)
    window.current_job = window.capture_job()
    outcome = RunOutcome(
        status,
        result_snapshot() if has_result else None,
        freeze_json(
            {
                "result": "result.json" if has_result else None,
                "artifacts": {"html": "report"} if files else {},
            }
        ),
        "<b>source failed</b>" if status == "failed" else None,
    )
    window.last_outcome = outcome
    window._present_outcome(outcome)
    window.resize(820, 600)
    application.processEvents()
    assert window.width() == 820
    assert window.details_button.isVisible()
    assert window.open_folder_button.isVisible()
    assert window.rect().contains(
        window.run_button.mapTo(window, window.run_button.rect().bottomRight())
    )
    if status == "failed":
        assert "<b>source failed</b>" in window.feedback.text()
        assert "saved analysis is available" not in window.feedback.text()


def test_long_option_caption_click_and_keyboard_toggle_the_captured_setting(window, application):
    window.resize(820, 600)
    window._navigate(2)
    application.processEvents()
    check = window.checks["use_large"]
    initial = check.isChecked()
    assert check.accessibleName() == "Use amount fallback when bill fields are missing"
    QTest.mouseClick(check.caption, Qt.MouseButton.LeftButton)
    assert check.isChecked() != initial
    assert window._configuration()["use_large"] != initial
    check.setFocus()
    QTest.keyClick(check, Qt.Key.Key_Space)
    assert check.isChecked() == initial
    assert window._configuration()["use_large"] == initial


def test_details_are_readonly_and_folder_is_opened_only_on_request(window, tmp_path, monkeypatch):
    source_folder(window, tmp_path)
    window.current_job = window.capture_job()
    window.last_outcome = RunOutcome("failed", result_snapshot(), freeze_json({}), "problem")
    opened = []

    def open_url(url):
        opened.append(Path(url.toLocalFile()))
        return False

    monkeypatch.setattr(qt_app.QDesktopServices, "openUrl", open_url)
    window._present_outcome(window.last_outcome)
    assert opened == []
    window.open_folder_button.click()
    assert opened == [Path(window.current_job.directory)]
    assert window.current_job.directory in window.feedback.text()
    captured = []

    def inspect(dialog):
        text = dialog.findChild(QPlainTextEdit)
        captured.append((text.isReadOnly(), text.toPlainText()))
        return QDialog.DialogCode.Rejected

    monkeypatch.setattr(QDialog, "exec", inspect)
    window.details_button.click()
    assert captured[0][0]
    assert "problem" in captured[0][1]
    assert window.current_job.directory in captured[0][1]


def test_settings_write_failure_is_nonblocking(window, monkeypatch):
    def fail(*args, **kwargs):
        raise OSError("read only")

    monkeypatch.setattr(qt_app, "write_json", fail)
    window.save_settings()
    assert "Preferences could not be saved" in window.feedback.text()
    assert window.feedback.property("severity") == "warning"


def test_qt_launcher_runs_real_event_loop_without_tk(application, tmp_path, monkeypatch):
    real_window = EvidenceWindow
    windows = []
    visible = []

    def build():
        widget = real_window(tmp_path / "launcher.json")
        windows.append(widget)

        def finish():
            visible.append(widget.isVisible())
            application.quit()

        QTimer.singleShot(30, finish)
        return widget

    monkeypatch.setattr(qt_app, "EvidenceWindow", build)
    assert qt_app.launch([]) == 0
    assert visible == [True]
    windows[0].close()
