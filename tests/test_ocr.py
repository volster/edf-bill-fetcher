"""OCR regressions: real image-only evidence, mixed pages and bounded failures."""

import os
import sys
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pdfplumber
import pypdfium2 as pdfium
import pytest
from reportlab.pdfgen.canvas import Canvas

from edf_bill_fetcher.application.run_service import run
from edf_bill_fetcher.collectors.engine import EvidenceEngine
from edf_bill_fetcher.domain.json_values import freeze_json, object_value
from edf_bill_fetcher.domain.requests import RunRequest
from edf_bill_fetcher.io.adapters import ocr
from edf_bill_fetcher.io.persistence.result_store import load_result

SCAN = Path(__file__).parent / "fixtures" / "scanned_bill.pdf"
HEADER = "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tconf\ttext\n"


@pytest.fixture
def engine():
    value = EvidenceEngine({"filter_below": False, "use_acc_filter": False}, lambda _: None)
    yield value
    value.cleanup()


@pytest.fixture
def real_ocr():
    try:
        return ocr.find_tesseract()
    except RuntimeError:
        if os.environ.get("GITHUB_ACTIONS") == "true":
            raise
        pytest.skip("Tesseract is required for the real OCR integration cases")


def _native_pdf(invoice="KI-56789", balance=500):
    output = BytesIO()
    canvas = Canvas(output)
    lines = [
        f"Invoice number: {invoice}",
        "Account number: A-12345678",
        "Date issued: 15 February 2024",
        "Your charges: 01 Feb 2024 - 29 Feb 2024",
        "Total charges for this period £400.00 debit",
        f"Current balance £{balance}.00 debit",
    ]
    for index, line in enumerate(lines):
        canvas.drawString(40, 780 - index * 30, line)
    canvas.save()
    return output.getvalue()


def _collect(engine, path):
    engine.process_pdf_file(str(path), "PDF", path.name, "15/01/2024")
    return engine.records, engine.__getstate__()


def test_real_image_only_bill_retains_correct_money_and_original_bytes(engine, real_ocr):
    before = SCAN.read_bytes()
    with pdfplumber.open(SCAN) as document:
        assert not document.pages[0].extract_text()
    _collect(engine, SCAN)
    assert len(engine.records) == 1
    row = engine.records[0]
    assert row["Invoice #"] == "KI-12345"
    assert row["Balance (£)"] == 250
    assert row["Period Charge (£)"] == 125
    assert row["Text Extraction"] == "OCR"
    assert row["OCR Pages"] == [1]
    assert row["OCR Review Required"] is True
    assert 0 <= row["OCR Minimum Confidence (%)"] <= 100
    assert (row["Page From"], row["Page To"]) == (1, 1)
    assert "£125.00" in row["Source PDF Text"]
    manifest = engine.source_manifest[row["Source ID"]]
    assert manifest["OCR Page Results"]["1"]["Status"] == "recognized"
    assert SCAN.read_bytes() == before
    assert any("verify amounts" in entry for entry in engine.error_log)


@pytest.mark.parametrize("scan_first", [False, True])
def test_real_mixed_pdf_only_ocrs_scanned_page_and_keeps_slice_provenance(
    tmp_path, engine, real_ocr, scan_first, monkeypatch
):
    combined = tmp_path / "mixed.pdf"
    native = pdfium.PdfDocument(_native_pdf())
    scan = pdfium.PdfDocument(SCAN)
    result = pdfium.PdfDocument.new()
    try:
        for part in (scan, native) if scan_first else (native, scan):
            result.import_pages(part)
        result.save(combined)
    finally:
        result.close()
        native.close()
        scan.close()
    calls = []
    original = ocr.PdfOcrReader.read_page

    def recognize(self, data, index, cancelled):
        calls.append(index)
        return original(self, data, index, cancelled)

    monkeypatch.setattr(ocr.PdfOcrReader, "read_page", recognize)
    _collect(engine, combined)
    assert calls == [0 if scan_first else 1]
    assert len(engine.records) == 2
    rows = {row["Invoice #"]: row for row in engine.records}
    digital, scanned = rows["KI-56789"], rows["KI-12345"]
    assert digital["Text Extraction"] == "PDF text"
    assert "OCR Pages" not in digital
    assert "OCR Review Required" not in digital
    assert scanned["OCR Pages"] == [1 if scan_first else 2]
    assert scanned["Page From"] == (1 if scan_first else 2)
    assert digital["Page From"] == (2 if scan_first else 1)
    assert scanned["Balance (£)"] == 250
    assert digital["Balance (£)"] == 500


def test_ocr_evidence_saves_reloads_and_renders_four_formats(tmp_path, engine, real_ocr):
    request = RunRequest(
        "ocr-integration",
        datetime(2024, 4, 1, tzinfo=timezone.utc),
        date(2024, 4, 1),
        (),
        freeze_json({"as_of": "2024-04-01", "analysis_min": 0}),
    )
    result = run(
        request,
        {fmt: str(tmp_path / ("report." + fmt)) for fmt in ("xlsx", "pdf", "docx", "html")},
        result_path=str(tmp_path / "result.json"),
        manifest_path=str(tmp_path / "run.json"),
        collect=lambda: _collect(engine, SCAN),
    )
    assert result.status == "review_required", result.error
    restored = load_result(tmp_path / "result.json")
    assert restored == result.snapshot
    raw = object_value(restored.prepared.extraction.observations[0].raw)
    assert raw["Text Extraction"] == "OCR"
    assert raw["OCR Pages"] == [1]
    assert raw["Balance (£)"] == 250
    assert set(object_value(result.manifest)["artifacts"]) == {"xlsx", "pdf", "docx", "html"}
    from docx import Document
    from openpyxl import load_workbook

    with pdfplumber.open(tmp_path / "report.pdf") as report:
        assert "OCR review required" in " ".join(page.extract_text() or "" for page in report.pages)
    assert "OCR review required" in " ".join(
        p.text for p in Document(tmp_path / "report.docx").paragraphs
    )
    assert "OCR review required" in (tmp_path / "report.html").read_text(encoding="utf-8")
    workbook = load_workbook(tmp_path / "report.xlsx")
    try:
        sheet = workbook["EDF Evidence Report"]
        headers = [cell.value for cell in sheet[1]]
        assert sheet.cell(2, headers.index("Text Extraction") + 1).value == "OCR"
        assert sheet.cell(2, headers.index("OCR Review Required") + 1).value is True
    finally:
        workbook.close()


def test_text_only_bill_does_not_need_ocr(tmp_path, engine, monkeypatch):
    path = tmp_path / "native.pdf"
    path.write_bytes(_native_pdf())
    recognize = MagicMock(side_effect=AssertionError("Text PDF must not invoke OCR"))
    monkeypatch.setattr(ocr.PdfOcrReader, "read_page", recognize)
    _collect(engine, path)
    recognize.assert_not_called()
    assert engine.records[0]["Balance (£)"] == 500
    assert engine.error_log == []


def test_failed_or_empty_ocr_keeps_later_readable_pages(tmp_path, engine, monkeypatch):
    path = tmp_path / "partial.pdf"
    output = BytesIO()
    canvas = Canvas(output)
    canvas.showPage()
    canvas.save()
    blank = pdfium.PdfDocument(output.getvalue())
    native = pdfium.PdfDocument(_native_pdf())
    result = pdfium.PdfDocument.new()
    try:
        result.import_pages(blank)
        result.import_pages(native)
        result.save(path)
    finally:
        result.close()
        native.close()
        blank.close()
    monkeypatch.setattr(
        ocr.PdfOcrReader, "read_page", MagicMock(side_effect=RuntimeError("OCR unavailable"))
    )
    _collect(engine, path)
    assert engine.records[0]["Balance (£)"] == 500
    assert any("OCR unavailable" in entry for entry in engine.error_log)
    assert (
        next(iter(engine.source_manifest.values()))["OCR Page Results"]["1"]["Status"] == "failed"
    )


def test_cancelled_ocr_stops_without_parsing_partial_page(engine, monkeypatch):
    def cancel(*args):
        engine.cancel_event.set()
        raise ocr.OcrCancelled()

    monkeypatch.setattr(ocr.PdfOcrReader, "read_page", cancel)
    _collect(engine, SCAN)
    assert engine.is_cancelled()
    assert engine.records == []
    assert engine.error_log == []


def test_tsv_retains_low_confidence_words_and_line_boundaries():
    page = ocr._parse_tsv(
        HEADER
        + "5\t1\t1\t1\t1\t1\t96.0\tBalance\n"
        + "5\t1\t1\t1\t1\t2\t12.5\t£250.00\n"
        + "5\t1\t1\t1\t2\t1\t99\tdebit\n"
    )
    assert page.text == "Balance £250.00\ndebit"
    assert page.minimum_confidence == 12.5
    assert ocr._parse_tsv(HEADER) == ocr.OcrPage("", None)


@pytest.mark.parametrize("output", ["bad response", HEADER + "5\t1\t1\t1\t1\t1\tnan\tvalue\n"])
def test_invalid_ocr_output_is_rejected(output):
    with pytest.raises(ValueError):
        ocr._parse_tsv(output)


def test_frozen_ocr_requires_bundled_executable_and_language_data(tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)
    with pytest.raises(RuntimeError, match="executable is missing"):
        ocr.find_tesseract()
    directory = tmp_path / "ocr"
    directory.mkdir()
    executable = directory / ("tesseract.exe" if os.name == "nt" else "tesseract")
    executable.touch()
    with pytest.raises(RuntimeError, match="English language data is missing"):
        ocr.find_tesseract()
    data = directory / "tessdata"
    data.mkdir()
    (data / "eng.traineddata").touch()
    runtime = ocr.find_tesseract()
    assert runtime.executable == executable
    assert runtime.tessdata == data


@pytest.mark.parametrize("mode", ["cancel", "timeout", "failure", "success"])
def test_ocr_subprocess_is_bounded_and_temporary_image_is_cleaned(tmp_path, monkeypatch, mode):
    monkeypatch.setattr(
        ocr, "find_tesseract", lambda: ocr.TesseractRuntime(tmp_path / "tesseract", tmp_path)
    )
    images = []

    def render(data, index, path):
        path.write_bytes(b"image")
        images.append(path)

    monkeypatch.setattr(ocr, "_render_page", render)
    process = MagicMock()
    process.returncode = 1 if mode == "failure" else 0
    process.communicate.return_value = (HEADER.encode(), b"engine failed")
    process.poll.return_value = None if mode in {"cancel", "timeout"} else process.returncode
    monkeypatch.setattr(ocr.subprocess, "Popen", lambda *args, **kwargs: process)
    times = iter([0.0, 61.0] if mode == "timeout" else [0.0, 0.0])
    monkeypatch.setattr(ocr, "time", SimpleNamespace(monotonic=lambda: next(times)))
    reader = ocr.PdfOcrReader()

    def cancelled():
        return mode == "cancel" and bool(process.communicate.call_count)

    if mode == "cancel":
        process.communicate.side_effect = [
            ocr.subprocess.TimeoutExpired("tesseract", 0.2),
            (b"", b""),
        ]
    if mode in {"cancel", "timeout", "failure"}:
        exception = {"cancel": ocr.OcrCancelled, "timeout": TimeoutError, "failure": RuntimeError}[
            mode
        ]
        with pytest.raises(exception):
            reader.read_page(b"pdf", 0, cancelled)
    else:
        assert reader.read_page(b"pdf", 0, cancelled) == ocr.OcrPage("", None)
    assert images
    assert not images[0].exists()
    assert process.kill.called is (mode in {"cancel", "timeout"})


@pytest.mark.parametrize("explicit", [None, "2"])
def test_ocr_bounds_default_threads_without_changing_host_preferences(
    tmp_path, monkeypatch, explicit
):
    if explicit is None:
        monkeypatch.delenv("OMP_THREAD_LIMIT", raising=False)
    else:
        monkeypatch.setenv("OMP_THREAD_LIMIT", explicit)
    monkeypatch.setattr(
        ocr, "find_tesseract", lambda: ocr.TesseractRuntime(tmp_path / "tesseract", None)
    )
    monkeypatch.setattr(ocr, "_render_page", lambda *args: None)
    process = MagicMock()
    process.returncode = 0
    process.communicate.return_value = (HEADER.encode(), b"")
    process.poll.return_value = 0
    launch = MagicMock(return_value=process)
    monkeypatch.setattr(ocr.subprocess, "Popen", launch)
    assert ocr.PdfOcrReader().read_page(b"pdf", 0, lambda: False) == ocr.OcrPage("", None)
    assert launch.call_args.kwargs["env"]["OMP_THREAD_LIMIT"] == (explicit or "1")
    assert os.environ.get("OMP_THREAD_LIMIT") == explicit


def test_sparse_selectable_header_does_not_hide_scanned_bill(tmp_path, engine, real_ocr):
    from reportlab.lib.utils import ImageReader

    rendered = tmp_path / "source.png"
    ocr._render_page(SCAN.read_bytes(), 0, rendered)
    path = tmp_path / "header-and-scan.pdf"
    canvas = Canvas(str(path), pagesize=(650, 930))
    canvas.drawImage(ImageReader(str(rendered)), 0, 0, width=650, height=900)
    canvas.drawString(30, 910, "EDF scanned bill")
    canvas.save()
    with pdfplumber.open(path) as document:
        assert document.pages[0].extract_text() == "EDF scanned bill"
    _collect(engine, path)
    assert engine.records[0]["Balance (£)"] == 250
    assert engine.records[0]["Text Extraction"] == "OCR"


@pytest.mark.parametrize("save_error", [False, True])
def test_rasterization_bounds_memory_and_closes_native_handles(monkeypatch, tmp_path, save_error):
    bitmap = MagicMock()
    image = bitmap.to_pil.return_value.convert.return_value
    if save_error:
        image.save.side_effect = OSError("disk full")
    page = MagicMock()
    page.get_size.return_value = (100_000, 100_000)
    page.render.return_value = bitmap
    document = MagicMock()
    document.__getitem__.return_value = page
    monkeypatch.setattr(pdfium, "PdfDocument", lambda data: document)
    if save_error:
        with pytest.raises(OSError, match="disk full"):
            ocr._render_page(b"pdf", 0, tmp_path / "page.png")
    else:
        ocr._render_page(b"pdf", 0, tmp_path / "page.png")
    scale = page.render.call_args.kwargs["scale"]
    assert (100_000 * scale) ** 2 <= 10_000_001
    assert 100_000 * scale <= 4000
    bitmap.close.assert_called_once()
    page.close.assert_called_once()
    document.close.assert_called_once()
    image.close.assert_called_once()


@pytest.mark.parametrize(
    ("text", "images", "expected"),
    [
        ("", [], True),
        ("EDF scanned bill", [{"width": 600, "height": 800}], True),
        ("Get in touch", [{"width": 50, "height": 50}], False),
        ("Your new account balance £200", [{"width": 600, "height": 800}], False),
        ("A long native bill text. " * 8, [{"width": 600, "height": 800}], False),
    ],
)
def test_scan_detection_distinguishes_headers_and_small_logos(text, images, expected):
    page = SimpleNamespace(width=650, height=900, images=images)
    assert ocr.requires_ocr(text, page) is expected
