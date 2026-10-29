"""Quality checks use applicable evidence, distinguish duplicate disposition and preserve saves."""

import copy

import pandas as pd
import pytest
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import Workbook
from reportlab.platypus import Table

from edf_bill_fetcher.application.analysis_service import analyse_prepared
from edf_bill_fetcher.application.legacy_adapter import capture_legacy
from edf_bill_fetcher.application.preparation_service import prepare_snapshot
from edf_bill_fetcher.domain.json_values import freeze_json
from edf_bill_fetcher.domain.requests import RunRequest
from edf_bill_fetcher.io.persistence.result_codec import decode_result, encode_result
from edf_bill_fetcher.io.reporters import docx_report, html_report, pdf_report
from edf_bill_fetcher.io.writers.data_quality import write_data_quality_sheet
from edf_bill_fetcher.io.writers.workbook_presentation import polish_reference_sheets
from edf_bill_fetcher.models.quality_tables import quality_notice, quality_table
from edf_bill_fetcher.models.report_models import compute_data_quality_report
from edf_bill_fetcher.models.snapshot_adapter import from_snapshot_json, to_snapshot_json
from edf_bill_fetcher.processors.preparation import normalize_frame


def _bill(**fields):
    return {
        "Account": "A-45678901",
        "Date": "01/02/2025",
        "Period From": "01/01/2025",
        "Period To": "31/01/2025",
        "Entry Type": "New Bill",
        "Amount (£)": 120,
        "Period Charge (£)": 120,
        "Invoice #": "TEST-001",
        "Reading": "Actual",
        "Energy Unit Rate (p/kWh)": 30,
        "Unit Rate Basis": "Extracted",
        "Unit Rate (p/kWh)": 30,
        "Source": "Local PDF Folder",
        "Source ID": "bill",
        **fields,
    }


def _mixed():
    return pd.DataFrame(
        [
            _bill(),
            {
                "Entry Type": "Ongoing Balance",
                "Document Role": "Payment reminder",
                "Date": "02/02/2025",
                "Amount (£)": 120,
                "Source": "Email Body",
            },
            {
                "Entry Type": "Payment",
                "Date": "03/02/2025",
                "Amount (£)": -20,
                "Transaction Amount (£)": -20,
                "Source": "HTM Account History",
            },
            {
                "Entry Type": "Credit",
                "Date": "04/02/2025",
                "Amount (£)": -10,
                "Transaction Amount (£)": -10,
                "Source": "Local PDF Folder",
            },
            {"Entry Type": "Supporting Document", "Amount Kind": "Not monetary", "Source": "PDF"},
        ]
    )


def _rows(report):
    return {row[0]: row[1:] for row in quality_table(report)[1:]}


def test_supporting_documents_and_reminders_do_not_lower_bill_completeness():
    frame = _mixed()
    before = frame.copy(deep=True)
    report = compute_data_quality_report(frame, duplicate_count=3)
    rows = _rows(report)
    assert rows["Financial record dates"] == ["4", "4", "100.0%", "PASS"]
    assert rows["Financial amounts present"] == ["4", "4", "100.0%", "PASS"]
    assert rows["Bill periods complete"] == ["1", "1", "100.0%", "PASS"]
    assert rows["Bill reading basis"] == ["1", "1", "100.0%", "PASS"]
    assert rows["Bill energy rates"] == ["1", "1", "100.0%", "INFO"]
    assert rows["Duplicate observations removed"] == ["3", "8", "37.5%", "INFO"]
    assert rows["Duplicate observations retained"] == ["0", "5", "0.0%", "PASS"]
    assert report.duplicate_rate == 3 / 8
    pd.testing.assert_frame_equal(frame, before)


@pytest.mark.parametrize("provided", [None, 0])
def test_supported_duplicates_still_included_require_review(provided):
    report = compute_data_quality_report(
        pd.DataFrame([_bill(), _bill(**{"Source ID": "copy"})]), duplicate_count=provided
    )
    assert report.duplicate_count == 1
    assert report.duplicate_rate == 0.5
    assert _rows(report)["Duplicate observations removed"][0] == "0"
    assert _rows(report)["Duplicate observations retained"] == ["1", "2", "50.0%", "REVIEW"]


def test_equal_unidentified_payments_do_not_become_proven_duplicate_decisions():
    payment = {"Date": "01/01/2025", "Entry Type": "Payment", "Amount (£)": -50}
    report = compute_data_quality_report(pd.DataFrame([payment, payment]))
    assert report.duplicate_count == 0


@pytest.mark.parametrize("empty", [False, True])
def test_nonfinancial_and_empty_inputs_have_no_spurious_completeness_failures(empty):
    frame = pd.DataFrame() if empty else _mixed().iloc[[-1]]
    report = compute_data_quality_report(frame)
    rows = _rows(report)
    assert rows["Financial record dates"] == ["0", "0", "—", "N/A"]
    assert rows["Bill reading basis"] == ["0", "0", "—", "N/A"]
    assert not any(row[-1] in {"FAIL", "REVIEW"} for row in rows.values())
    assert "None" not in str(quality_table(report))


def test_missing_or_reversed_bill_periods_and_unknown_reading_remain_review_work():
    report = compute_data_quality_report(
        pd.DataFrame([_bill(**{"Period From": "02/02/2025", "Reading": "Unknown"})])
    )
    assert _rows(report)["Bill periods complete"][-1] == "REVIEW"
    assert _rows(report)["Bill reading basis"][-1] == "REVIEW"


def test_a_balance_figure_cannot_satisfy_missing_bill_charge_or_payment_value():
    report = compute_data_quality_report(
        pd.DataFrame(
            [
                _bill(**{"Period Charge (£)": None}),
                {
                    "Entry Type": "Payment",
                    "Amount Kind": "Balance",
                    "Amount (£)": 120,
                    "Source": "HTM Account History",
                },
            ]
        )
    )
    assert _rows(report)["Financial amounts present"][-1] == "PASS"
    assert _rows(report)["Bill charge amounts"] == ["0", "1", "0.0%", "REVIEW"]
    assert _rows(report)["Transaction values"] == ["0", "1", "0.0%", "REVIEW"]


@pytest.mark.parametrize("kind", ["New Bill", "Ongoing Balance", "Payment", "Credit"])
def test_an_explicitly_unknown_amount_is_not_reclassified_from_its_document_type(kind):
    original = {"Entry Type": kind, "Amount Kind": "Unknown", "Amount (£)": 120}
    frame = normalize_frame(pd.DataFrame([original]))
    row = frame.iloc[0]
    assert row["Amount Kind"] == "Unknown"
    assert row["Balance (£)"] is None
    assert row["Transaction Amount (£)"] is None
    assert _rows(compute_data_quality_report(frame))["Amounts with unverified meaning"][0] == "1"


@pytest.mark.parametrize("basis", ["Customer", "Deemed / estimated"])
def test_source_qualified_bill_reading_bases_are_recognized(basis):
    report = compute_data_quality_report(pd.DataFrame([_bill(Reading=basis)]))
    assert _rows(report)["Bill reading basis"][0] == "1"


def test_unverified_meaning_and_ambiguous_postings_are_visible_in_every_format():
    report = compute_data_quality_report(
        pd.DataFrame(
            [
                {"Amount (£)": 500, "Amount Kind": "Unknown", "Entry Type": "Ongoing Balance"},
                {
                    "Amount (£)": -50,
                    "Entry Type": "Payment",
                    "Transaction Identity Review": "Compare originals",
                    "Statement Reconciliation Status": "Review required",
                },
            ]
        )
    )
    rows = _rows(report)
    for check in (
        "Amounts with unverified meaning",
        "Ambiguous transaction identity",
        "Statements needing reconciliation",
    ):
        assert rows[check][0] == "1"
        assert rows[check][-1] == "REVIEW"


def test_all_output_formats_share_accepted_checks_without_recalculation(monkeypatch):
    report = compute_data_quality_report(_mixed(), duplicate_count=3)
    expected = quality_table(report)

    def forbidden(*args, **kwargs):
        raise AssertionError("Rendering must use accepted assessment facts")

    monkeypatch.setattr(
        "edf_bill_fetcher.models.report_models.compute_data_quality_report", forbidden
    )
    monkeypatch.setattr(
        "edf_bill_fetcher.processors.quality_assessment.build_quality_assessment", forbidden
    )
    frame = pd.DataFrame()
    html = BeautifulSoup(
        html_report.create_data_quality_section(frame, payload=report), "html.parser"
    )
    assert [
        [cell.get_text() for cell in row.select("th,td")]
        for row in html.select("table")[0].select("tr")
    ] == expected
    doc = Document()
    docx_report.create_data_quality_section(
        doc, docx_report._get_or_create_styles(doc), frame, payload=report
    )
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == expected
    table = next(
        element
        for element in pdf_report.create_data_quality_section(frame, payload=report)
        if isinstance(element, Table)
    )
    assert [[cell.getPlainText() for cell in row] for row in table._cellvalues] == expected
    sheet = Workbook().active
    write_data_quality_sheet(sheet, frame, payload=report)
    polish_reference_sheets(sheet.parent)
    assert sheet.print_title_rows == "$1:$5"
    assert [
        [str(sheet.cell(row, column).value) for column in range(1, 6)]
        for row in range(5, 5 + len(expected))
    ] == expected
    assert any("12 checks" in str(cell.value) for row in sheet for cell in row)


def test_legacy_saved_metrics_roundtrip_exactly_without_reanalysis(monkeypatch):
    payload = to_snapshot_json(compute_data_quality_report(_mixed(), duplicate_count=3))
    del payload["fields"]["assessment"]
    payload["fields"]["duplicate_rate"] = 0.6
    original = copy.deepcopy(payload)

    def forbidden(*args, **kwargs):
        raise AssertionError("A saved result must retain its accepted scope")

    monkeypatch.setattr(
        "edf_bill_fetcher.processors.quality_assessment.build_quality_assessment", forbidden
    )
    restored = from_snapshot_json(payload)
    assert restored.assessment is None
    assert to_snapshot_json(restored) == original
    assert _rows(restored)["Recorded duplicate decisions"] == ["3", "5", "60.0%", "INFO"]
    assert "predates evidence-type eligibility" in quality_notice(restored)


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("version", 2),
        ("bills", 10),
        ("complete_bill_periods", 2),
        ("financial_amounts", -1),
        ("financial_records", True),
        ("duplicates_removed", 0),
        ("unknown_field", 1),
    ],
)
def test_saved_assessment_scope_and_disposition_are_strictly_validated(field, value):
    payload = to_snapshot_json(compute_data_quality_report(_mixed(), duplicate_count=3))
    payload["fields"]["assessment"][field] = value
    with pytest.raises(ValueError):
        from_snapshot_json(payload)


def test_saved_duplicate_rate_cannot_disagree_with_its_recorded_denominator():
    payload = to_snapshot_json(compute_data_quality_report(_mixed(), duplicate_count=3))
    payload["fields"]["duplicate_rate"] = 0.6
    with pytest.raises(ValueError, match="duplicate rate"):
        from_snapshot_json(payload)


def test_only_the_exact_legacy_shape_can_omit_quality_scope():
    payload = to_snapshot_json(compute_data_quality_report(_mixed()))
    payload["fields"]["assessment"] = None
    with pytest.raises(ValueError, match="legacy shape"):
        from_snapshot_json(payload)
    del payload["fields"]["assessment"]
    del payload["fields"]["total_records"]
    with pytest.raises(ValueError, match="fields disagree"):
        from_snapshot_json(payload)


def test_acceptance_and_strict_result_roundtrip_include_quality_disposition():
    from datetime import date, datetime, timezone

    request = RunRequest(
        "quality-fixture",
        datetime(2025, 2, 1, tzinfo=timezone.utc),
        date(2025, 2, 1),
        (),
        freeze_json({"as_of": "2025-02-01", "use_dedup": True}),
    )
    rows = [_bill(), _bill(**{"Source ID": "copy"})]
    result = analyse_prepared(prepare_snapshot(capture_legacy(rows, {}, request)))
    encoded = encode_result(result, "test")
    assert encode_result(decode_result(encoded), "test") == encoded
    fields = encoded["analysis"]["models"]["data_quality"]["payload"]["fields"]
    assert fields["assessment"]["duplicates_removed"] == 1
    assert fields["assessment"]["duplicates_retained"] == 0
    assert fields["duplicate_rate"] == 0.5


@pytest.mark.parametrize("deduplicate", [True, False])
def test_run_status_distinguishes_removed_duplicates_from_duplicates_still_included(
    tmp_path, deduplicate
):
    from datetime import date, datetime, timezone

    from edf_bill_fetcher.application.run_service import run

    request = RunRequest(
        "quality-run",
        datetime(2025, 2, 1, tzinfo=timezone.utc),
        date(2025, 2, 1),
        (),
        freeze_json({"as_of": "2025-02-01", "use_dedup": deduplicate}),
    )

    def render(view, path):
        from pathlib import Path

        Path(path).write_text("Complete report", encoding="utf-8")

    outcome = run(
        request,
        {"html": str(tmp_path / "report.html")},
        result_path=str(tmp_path / "result.json"),
        manifest_path=str(tmp_path / "run.json"),
        collect=lambda: ([_bill(), _bill(**{"Source ID": "copy"})], {}),
        renderers={"html": render},
    )
    assert outcome.error is None
    assert outcome.status == ("complete" if deduplicate else "review_required")
