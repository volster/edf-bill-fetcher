"""Synthetic reproductions of failures verified against privately supplied bills."""

from __future__ import annotations

import base64
import json
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook

from edf_bill_fetcher.collectors.engine import EvidenceEngine
from edf_bill_fetcher.helpers.excel_utils import (
    build_evidence_index,
    evidence_report_hyperlink_cell,
    matching_evidence_row,
)
from edf_bill_fetcher.helpers.identity import document_identity
from edf_bill_fetcher.io.reporters.html_report import _source_links, create_evidence_index
from edf_bill_fetcher.io.writers.evidence_bundle import _infer_section
from edf_bill_fetcher.io.writers.meter import write_meter_readings_sheet
from edf_bill_fetcher.processors.bill_facts import bill_facts, payment_history
from edf_bill_fetcher.processors.detection import _ambiguous_invoice_references, detect_back_billing
from edf_bill_fetcher.processors.preparation import prepare_evidence
from edf_bill_fetcher.processors.sap_parsers import (
    extract_new_invoice_fields,
    extract_reconciliation_statement_rows,
)

OLD_BILL = """Account number
securely 24/7 via MyAccount
601 234 567 890 edfenergy.com/247
Invoice number: T12345678901
Bill date: 4 Sep 2020
Your business electricity bill
For 14 June 2019 - 04 September 2020 (448 days)
Your new account balance £8,000.00 in debit
Meter: D99T12345 Tariff: Example
14 Jun 19 - 31 Dec 19 1000 OUR READ 2000 ESTIMATED 1000 kWh 20.00p £200.00
01 Jan 20 - 04 Sep 20 2000 ESTIMATED 4000 YOUR READ 2000 kWh 25.00p £500.00
Total electricity charges for this period £700.00
Total charges for this period £500.00
credit
Your usage this bill 3,000 kWh This period last year 2,900 kWh
Your payment history Page 4 of 4
21 August 2020 £100.00
4 September 2020 £100.00
Frequently asked questions
"""

STATEMENT = """Bill reference: 12345678 (7 April 2026)
Account number: A-12345678
Balance on your last bill £1000.00 debit
Charges
Electricity 1 July 2024 - 31 July 2024 £200.00
Late Payment Charge £60.00
Credits
Reversed electricity charge 5 Sept. 2024 £200.00
(1 July 2024 - 31 July 2024)
Payments
1 April 2026 £100.00
Your new balance £960.00 debit
Your charges in detail
Other charges
25 Apr 2025 Late Payment Charge £30.00
26 Apr 2025 Late Payment Charge £30.00
"""


def test_old_credit_balance_usage_gross_and_payments_are_distinct():
    engine = EvidenceEngine({"filter_below": False}, lambda _: None)
    engine.process_text(OLD_BILL, "Local PDF Folder", "synthetic.pdf", "01/01/2026")
    bill, *payments = engine.records
    assert bill["Balance (£)"] == 8000
    assert bill["Transaction Amount (£)"] == -500
    assert bill["Period Charge (£)"] == -500
    assert bill["Gross Electricity Charges (£)"] == 700
    assert bill["Units (kWh)"] == 3000
    assert bill["Reading"] == "Mixed"
    assert bill["Account"] == "601234567890"
    assert bill["Meter"] == "D99T12345"
    assert bill["Reading Date"] == "04/09/2020"
    assert bill["Register Reading (kWh)"] == 4000
    assert bill["Closing Reading Type"] == "Customer"
    assert [(r["Date"], r["Transaction Amount (£)"]) for r in payments] == [
        ("21/08/2020", -100),
        ("04/09/2020", -100),
    ]
    assert all(r["Balance (£)"] is None for r in payments)
    assert all(r.get("Register Reading (kWh)") is None for r in payments)
    backbilling = detect_back_billing(pd.DataFrame([bill]))
    assert len(backbilling) == 1
    assert backbilling.iloc[0]["Value Source"].startswith("Gross replacement")


@pytest.mark.parametrize(("side", "expected"), [("debit", 500), ("credit", -500), ("", 500)])
def test_old_net_side_is_local_to_total(side, expected):
    engine = EvidenceEngine({"filter_below": False}, lambda _: None)
    text = OLD_BILL.replace("£500.00\ncredit", f"£500.00 {side}")
    engine.process_text(text, "Local PDF Folder", "synthetic.pdf", "01/01/2026")
    assert engine.records[0]["Transaction Amount (£)"] == expected
    assert engine.records[0]["Balance (£)"] == 8000


def test_credit_electricity_component_is_not_a_positive_replacement_charge():
    facts = bill_facts("Total electricity charges for this period £200.00 credit", [])
    assert facts["Gross Electricity Charges (£)"] == -200
    frame = pd.DataFrame(
        [
            {
                "Entry Type": "New Bill",
                "Date": "01/01/2026",
                "Period From": "01/01/2020",
                "Period To": "01/02/2020",
                "Period Charge (£)": -200,
                **facts,
            }
        ]
    )
    assert detect_back_billing(frame).empty


def test_new_wrapped_usage_tariff_and_mixed_sign_cover():
    text = """Invoice number: KI-12345678-0001
Account number: A-12345678
Date issued: 06 August 2024
Your charges: 1 July 2024 - 31 July 2024
Electricity -£190.48 -£9.52 -£200.00
Total charges for this period £200.00 debit
Current balance £1000.00 debit
Supply address: Example Farm, Tariff name Freedom
AB12 3CD
Payment type Non-Direct Debit
Electricity charges for meter D99T12345
1 Jul 2024 1000.0000 Estimated reading
1 Aug 2024 1500.0000 Customer reading
500.0000 kWh @
Electricity used 30.000p/kWh £150.00
"""
    fields = extract_new_invoice_fields(text)
    assert fields["tariff"] == "Freedom"
    assert float(fields["units_used"]) == 500
    assert fields["period_charge"] == 200
    facts = bill_facts(text, [])
    assert facts["Reading Date"] == "01/08/2024"
    assert facts["Register Reading (kWh)"] == 1500
    assert facts["Closing Reading Type"] == "Customer"


@pytest.mark.parametrize(
    ("text", "account"),
    [
        (
            "Sort code 40-00-00 Account number 12345678 Reference A-12345678\nAccount number: A-12345678",
            "A-12345678",
        ),
        ("Account number 12345678 Sort code 40-00-00", ""),
        ("Account number A-12345678 Account number A-87654321", ""),
        ("Account number 601 234 567 890", "601234567890"),
    ],
)
def test_customer_identity_does_not_use_bank_or_ambiguous_accounts(text, account):
    assert document_identity(text)["Account"] == account


def test_statement_summary_and_detail_fees_reconcile_without_losing_equal_events():
    rows = extract_reconciliation_statement_rows(STATEMENT, "synthetic.pdf")
    fees = [r for r in rows if r["Entry Type"] == "Late Payment"]
    assert [(r["Date"], r["Amount (£)"]) for r in fees] == [("25/04/2025", 30), ("26/04/2025", 30)]
    assert rows[-1]["Statement Reconciliation Difference (£)"] == 0
    assert rows[-1]["Statement Reconciliation Status"] == "Balanced"
    assert rows[0]["Document Date"] == "07/04/2026"
    assert rows[0]["First Billed Date"] == "N/A"
    assert rows[1]["Period From"] == "01/07/2024"
    assert detect_back_billing(pd.DataFrame(rows)).empty


def test_statement_mismatch_is_disclosed_not_forced_to_balance():
    rows = extract_reconciliation_statement_rows(
        STATEMENT.replace("£960.00", "£959.00"), "test.pdf"
    )
    assert rows[-1]["Statement Reconciliation Difference (£)"] == 1
    assert rows[-1]["Statement Reconciliation Status"] == "Review required"


def test_statement_container_reference_is_not_a_conflicting_invoice_version():
    rows = extract_reconciliation_statement_rows(STATEMENT, "test.pdf")
    rows.append({**rows[0], "Period From": "01/08/2024", "Period To": "31/08/2024"})
    assert _ambiguous_invoice_references(pd.DataFrame(rows)) == set()
    originals = pd.DataFrame(
        [
            {
                "Entry Type": "New Bill",
                "Invoice #": "KI-1",
                "Date": "01/09/2024",
                "Period Charge (£)": 100,
            },
            {
                "Entry Type": "New Bill",
                "Invoice #": "KI-1",
                "Date": "01/09/2024",
                "Period Charge (£)": 101,
            },
        ]
    )
    assert _ambiguous_invoice_references(originals) == {"KI-1"}


def test_fee_summary_mismatch_preserves_details_without_counting_both():
    rows = extract_reconciliation_statement_rows(
        STATEMENT.replace("Charge £60.00", "Charge £90.00"), "test.pdf"
    )
    assert sum(r["Amount (£)"] for r in rows if r["Entry Type"] == "Late Payment") == 60
    assert rows[-1]["Statement Reconciliation Status"] == "Review required"
    assert "only details counted" in rows[-1]["Reconciliation Review"]


def test_credit_statement_balances_have_signed_meaning():
    rows = extract_reconciliation_statement_rows(
        "Bill reference: 12345678 (7 April 2026)\nAccount number: A-12345678\n"
        "Balance on your last bill £100.00 credit\nPayments\n1 April 2026 £20.00\n"
        "Your new balance £120.00 credit",
        "test.pdf",
    )
    assert rows[-1]["Amount (£)"] == -120
    assert rows[-1]["Statement Reconciliation Difference (£)"] == 0


def test_original_invoice_corroborates_statement_line_and_keeps_both_occurrences():
    statement = extract_reconciliation_statement_rows(STATEMENT, "statement.pdf")[0]
    statement["Source ID"] = "statement-source"
    original = {
        "Source": "Local PDF Folder",
        "Source ID": "original-source",
        "Account": "A-12345678",
        "Invoice #": "KI-12345678-0001",
        "Date": "06/08/2024",
        "Entry Type": "New Bill",
        "Period From": "01/07/2024",
        "Period To": "31/07/2024",
        "Period Charge (£)": 200,
    }
    result = prepare_evidence([statement, original], {"as_of": "2026-10-17"})
    assert len(result.records) == len(result.duplicates) == 1
    assert result.records[0]["Invoice #"] == "KI-12345678-0001"
    assert result.records[0]["Source IDs"] == ["original-source", "statement-source"]
    assert result.records[0]["Source Occurrences"][0]["Document Date"] == "07/04/2026"
    assert result.duplicates[0]["First Billed Date"] == "06/08/2024"
    for changed in (
        {"Account": ""},
        {"Account": "A-87654321"},
        {"Period Charge (£)": 201},
        {"Meter": "M2"},
    ):
        candidate = {**original, "Meter": "M1"}
        stmt = {**statement, "Meter": "M1", **changed}
        if "Period Charge (£)" in changed:
            stmt["Transaction Amount (£)"] = changed["Period Charge (£)"]
        assert len(prepare_evidence([stmt, candidate]).records) == 2
    ambiguous = {**original, "Invoice #": "KI-12345678-0002"}
    assert len(prepare_evidence([statement, original, ambiguous]).records) == 3


def test_payment_history_does_not_invent_cover_payments_or_collapse_identical_lines():
    assert payment_history("You paid us £200.00 (2 payments of £100.00)") == []
    rows = payment_history("Your payment history\n1 May 2020 £100.00\n1 May 2020 £100.00")
    assert len(rows) == 2
    assert rows[0]["Source Row"] != rows[1]["Source Row"]
    assert len(prepare_evidence(rows).records) == 2


def test_review_edit_revokes_a_previously_supported_statement_link():
    statement = extract_reconciliation_statement_rows(STATEMENT, "statement.pdf")[0]
    original = {
        "Account": "A-12345678",
        "Entry Type": "New Bill",
        "Invoice #": "KI-12345678-0001",
        "Date": "06/08/2024",
        "Period From": "01/07/2024",
        "Period To": "31/07/2024",
        "Period Charge (£)": 200,
    }
    first = prepare_evidence([statement, original])
    edited = {**first.duplicates[0], "Transaction Amount (£)": 201, "Period Charge (£)": 201}
    second = prepare_evidence([first.records[0], edited])
    assert len(second.records) == 2
    unresolved = next(row for row in second.records if row.get("Document Role"))
    assert unresolved["First Billed Date"] == "N/A"
    assert "Reconciled Event ID" not in unresolved


def test_identical_statement_lines_are_ambiguous_even_with_one_original():
    statement = extract_reconciliation_statement_rows(STATEMENT, "statement.pdf")[0]
    original = {
        "Account": "A-12345678",
        "Entry Type": "New Bill",
        "Invoice #": "KI-12345678-0001",
        "Date": "06/08/2024",
        "Period From": "01/07/2024",
        "Period To": "31/07/2024",
        "Period Charge (£)": 200,
    }
    result = prepare_evidence([statement, {**statement, "Source Row": 99}, original])
    assert len(result.records) == 3
    assert all("Reconciliation Review" in row for row in result.records if row.get("Document Role"))


def test_missing_and_ambiguous_references_never_link_unrelated_rows():
    df = pd.DataFrame(
        [
            {"Invoice #": "N/A", "Source ID": "terms", "Source PDF Text": "Terms"},
            {"Invoice #": "123", "Source ID": "statement", "Source Row": 1},
            {"Invoice #": "123", "Source ID": "statement", "Source Row": 2},
        ]
    )
    assert matching_evidence_row(df, {"Invoice #": "N/A"}) is None
    assert matching_evidence_row(df, {"Invoice #": "123"}) is None
    match = matching_evidence_row(df, {"Source ID": "statement", "Source Row": 2})
    assert match is not None
    assert match[0] == 4
    assert "inv:123" not in build_evidence_index(df)
    ws = Workbook().active
    evidence_report_hyperlink_cell(ws, 1, 1, df, "N/A")
    assert ws["A1"].hyperlink is None


def test_account_filename_is_not_regulatory_section():
    assert _infer_section("A-12345678-bill.pdf") == ("D", False)
    assert _infer_section("A1 - regulator letter.pdf") == ("A", False)


def test_html_links_use_saved_manifest_identity_and_escape_text(tmp_path: Path) -> None:
    folder = tmp_path / "evidence_files"
    folder.mkdir()
    (folder / "invoice & copy.pdf").write_bytes(b"%PDF-1.4")
    files = [
        {
            "source_id": "source",
            "attachment_name": "bill.pdf",
            "status": "saved",
            "saved_path": "evidence_files/invoice & copy.pdf",
        }
    ]
    (tmp_path / "evidence_manifest.json").write_text(json.dumps({"files": files}))
    row = {
        "Source": "Local PDF Folder",
        "Source ID": "source",
        "Date": "01/01/2026",
        "Invoice #": "<script>x</script>",
        "Page From": 2,
    }
    output = create_evidence_index(pd.DataFrame([row]), SimpleNamespace(), bundle_dir=str(tmp_path))
    assert 'href="evidence_files/invoice%20%26%20copy.pdf#page=2"' in output
    assert "<script>" not in output
    assert '<details class="audit">' in output
    files += [{"source_id": "escape", "status": "saved", "saved_path": "../outside.pdf"}]
    (tmp_path / "evidence_manifest.json").write_text(json.dumps({"files": files}))
    assert "escape" not in _source_links(str(tmp_path))


def test_terms_only_document_cannot_turn_liability_limit_into_a_bill():
    engine = EvidenceEngine({"filter_below": False}, lambda _: None)
    engine.process_text(
        "Terms and conditions. Liability limit £10,000.00; estimated annual usage 300000 kWh",
        "Local PDF Folder",
        "terms.pdf",
        "01/01/2026",
    )
    assert len(engine.records) == 1
    assert engine.records[0]["Entry Type"] == "Supporting Document"
    assert engine.records[0]["Amount (£)"] is None
    assert engine.records[0]["Transaction Amount (£)"] is None


def test_pst_body_export_is_traceable_and_preserves_captured_bytes():
    raw = b"<p>Your new account balance &pound;750.00 in debit</p>"
    msg = SimpleNamespace(get_html_body=lambda: raw)
    engine = EvidenceEngine({"filter_below": False}, lambda _: None)
    engine._archive_origin = "/example/archive.pst"
    try:
        engine._process_pst_body(
            msg, "Synthetic bill", "01/01/2026", "billing@example.com", "Inbox, message index 7"
        )
        record = engine.records[0]
        source = Path(engine.source_paths[record["Source ID"]])
        exported = source.read_text()
        assert "Generated PST/OST message body export" in exported
        assert base64.b64encode(raw).decode("ascii") in exported
        assert "Inbox, message index 7" in exported
        assert (
            engine.source_manifest[record["Source ID"]]["Source Representation"]
            == "Generated PST/OST body text export"
        )
        first_id = record["Source ID"]
        engine._archive_origin = "/example/another-archive.pst"
        engine._process_pst_body(
            msg, "Synthetic bill", "01/01/2026", "billing@example.com", "Inbox, message index 7"
        )
        assert engine.records[-1]["Source ID"] != first_id
    finally:
        engine.cleanup()


def test_pst_attachments_keep_occurrences_with_identical_subjects_separate(monkeypatch):
    monkeypatch.setattr("edf_bill_fetcher.collectors.engine.HAS_PYPFF", True)
    attachment = SimpleNamespace(get_size=lambda: 9, read_buffer=lambda size: b"%PDF-test")
    message = SimpleNamespace(
        get_subject=lambda: "EDF bill",
        get_delivery_time=lambda: datetime(2026, 1, 1),
        get_html_body=lambda: b"<p>Your new account balance &pound;750.00</p>",
        get_number_of_attachments=lambda: 1,
        get_attachment=lambda index: attachment,
    )
    folder = SimpleNamespace(
        get_name=lambda: "Inbox",
        get_number_of_sub_messages=lambda: 2,
        get_sub_message=lambda index: message,
        get_number_of_sub_folders=lambda: 0,
    )
    engine = EvidenceEngine({"filter_below": False}, lambda _: None)
    engine._archive_origin = "/example/archive.pst"
    parents = []
    monkeypatch.setattr(engine, "_process_attachment", lambda *args: parents.append(args[-1]))
    try:
        engine.crawl_pst(folder)
        assert len(parents) == 2
        assert parents[0] != parents[1]
        assert "message index 0" in parents[0]
        assert "/example/archive.pst" in parents[1]
    finally:
        engine.cleanup()


def test_meter_history_shows_real_read_dates_and_sap_readings():
    ws = Workbook().active
    frame = pd.DataFrame(
        [
            {
                "Date": "06/08/2024",
                "Reading Date": "01/08/2024",
                "Register Reading (kWh)": 1500,
                "Invoice #": "KI-1",
                "Reading": "Estimated",
            },
            {"Date": "01/01/2025", "Invoice #": "N/A", "Register Reading (kWh)": "N/A"},
        ]
    )
    write_meter_readings_sheet(
        ws,
        frame,
        sap_readings=[
            {
                "Meter Read Date": "2024-09-01",
                "Reading (kWh)": "2000",
                "Read Type": "A",
                "Source File": "readings.pdf",
            }
        ],
    )
    assert ws["A8"].value == "01/08/2024"
    assert ws["B8"].value == 1500
    assert ws["B9"].value == 2000
    assert ws["C9"].value == "U"
    assert "raw type A" in ws["F9"].value
    assert "SAP reading" in ws["F9"].value
    assert ws["A10"].value is None


def test_old_credit_balance_does_not_become_a_positive_debt():
    engine = EvidenceEngine({"filter_below": False}, lambda _: None)
    engine.process_text(
        OLD_BILL.replace("£8,000.00 in debit", "£8,000.00 in credit"),
        "Local PDF Folder",
        "example.pdf",
        "01/01/2026",
    )
    assert engine.records[0]["Balance (£)"] == -8000
    assert engine.records[0]["Transaction Amount (£)"] == -500


def test_statement_links_never_merge_distinct_meter_ledgers_with_one_invoice_reference():
    statement = extract_reconciliation_statement_rows(STATEMENT, "statement.pdf")[0]
    original = {
        "Account": "A-12345678",
        "Entry Type": "New Bill",
        "Invoice #": "KI-1",
        "Date": "06/08/2024",
        "Period From": "01/07/2024",
        "Period To": "31/07/2024",
        "Period Charge (£)": 200,
    }
    originals = [{**original, "Meter": meter} for meter in ("M1", "M2")]
    ambiguous = prepare_evidence([statement, *originals])
    assert len(ambiguous.records) == 3
    statements = [{**statement, "Meter": meter} for meter in ("M1", "M2")]
    linked = prepare_evidence([*statements, *originals])
    assert len(linked.records) == 2
    assert len(linked.duplicates) == 2
    assert {row["Meter"] for row in linked.records} == {"M1", "M2"}
