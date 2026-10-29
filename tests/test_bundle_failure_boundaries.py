"""Broken evidence mappings remain explicit without losing valid originals or prior files."""

import hashlib
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest
from docx import Document

from edf_bill_fetcher.io.writers import evidence_bundle as bundle


@pytest.mark.parametrize("name", ["CON", "nul.pdf", "COM1.txt", "LPT9.pdf"])
def test_reserved_windows_names_remain_extractable_on_every_platform(name):
    saved = bundle.sanitise_filename(name)
    assert saved.startswith("_")
    assert saved.endswith(name)


@pytest.mark.parametrize(
    ("name", "section"),
    [("Financial-Transactions.pdf", "F"), ("Contract-History.pdf", "D"), ("Z1 - unknown.pdf", "D")],
)
def test_supporting_evidence_routing_has_a_stable_fallback(name, section):
    assert bundle._infer_section(name)[0] == section


@pytest.mark.parametrize(
    "manifest",
    [
        [],
        {"broken": None},
        {"broken": {"Saved Path": 7}},
        {"broken": {"Saved Path": "C:/outside.pdf"}},
    ],
)
def test_invalid_saved_mapping_returns_an_explanation(tmp_path, manifest):
    paths, notes = bundle.resolve_bundle_paths({"source_manifest": manifest}, str(tmp_path))
    assert paths == {}
    assert notes
    assert any("mapping is unavailable" in note or "path is invalid" in note for note in notes)


@pytest.mark.parametrize("legacy", [[], {"missing": None}, {"missing": ""}])
def test_invalid_legacy_mapping_never_guesses_a_file(tmp_path, legacy):
    paths, notes = bundle.resolve_bundle_paths({"saved_evidence_paths": legacy}, str(tmp_path))
    assert paths == {}
    assert notes


@pytest.mark.parametrize("failure", ["resolve", "hash"])
def test_one_unavailable_source_keeps_other_verified_sources_available(
    tmp_path, monkeypatch, failure
):
    metadata = {}
    for name in ("broken", "good"):
        path = tmp_path / (name + ".pdf")
        path.write_bytes(name.encode())
        metadata[name] = {
            "Saved Path": path.name,
            "Content Hash": hashlib.sha256(path.read_bytes()).hexdigest(),
        }
    if failure == "resolve":
        original = Path.resolve

        def unavailable(path, *args, **kwargs):
            if path.name == "broken.pdf":
                raise OSError("Simulated unavailable storage")
            return original(path, *args, **kwargs)

        monkeypatch.setattr(Path, "resolve", unavailable)
    else:
        original_record = bundle.artifact_record

        def unreadable(path):
            if Path(path).name == "broken.pdf":
                raise PermissionError("Simulated permission failure")
            return original_record(path)

        monkeypatch.setattr(bundle, "artifact_record", unreadable)
    paths, notes = bundle.resolve_bundle_paths({"source_manifest": metadata}, str(tmp_path))
    assert "broken" not in paths
    assert paths["good"] == str(tmp_path / "good.pdf")
    assert len(notes) == 1


def test_manifest_skips_unidentified_rows_and_repeated_page_ranges(tmp_path):
    original = tmp_path / "source.pdf"
    original.write_bytes(b"Synthetic original")
    row = {
        "Source ID": "source",
        "Attachment Name": "source.pdf",
        "Invoice #": "INV",
        "Page From": 1,
        "Page To": 2,
    }
    manifest = bundle.build_bundle_manifest(
        pd.DataFrame([{}, row, row]),
        {"source": str(original), "source.pdf": str(original)},
        str(tmp_path),
    )
    assert len(manifest["files"]) == 1
    assert manifest["files"][0]["page_ranges"] == [[1, 2]]


def test_index_preserves_account_and_only_links_available_originals(tmp_path):
    original = tmp_path / "evidence_files" / "source.pdf"
    original.parent.mkdir()
    original.write_bytes(b"Original bytes")
    frame = pd.DataFrame(
        [
            {"Attachment Name": None, "Invoice #": "MISSING"},
            {"Attachment Name": "source.pdf", "Invoice #": "INV"},
        ]
    )
    path = tmp_path / "index.docx"
    bundle.build_bundle_index(frame, {"source.pdf": str(original)}, str(path), account="000123")
    doc = Document(path)
    assert "000123" in doc.paragraphs[0].text
    with ZipFile(path) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode()
    assert "evidence_files/source.pdf" in rels
    assert "MISSING" not in "\n".join(p.text for p in doc.paragraphs)
