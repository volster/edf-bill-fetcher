"""Rebilling history, human positions and cross-format review must preserve financial meaning."""

import copy
from dataclasses import replace
from decimal import Decimal

import pytest
from docx import Document
from openpyxl import load_workbook

from edf_bill_fetcher.application.analysis_service import analyse_prepared
from edf_bill_fetcher.application.legacy_adapter import capture_legacy
from edf_bill_fetcher.application.preparation_service import prepare_snapshot
from edf_bill_fetcher.application.review_session import ReviewSession
from edf_bill_fetcher.application.workbook_review import regenerate_workbook
from edf_bill_fetcher.domain.analysis import AnalysisSnapshot
from edf_bill_fetcher.domain.json_values import freeze_json, object_value
from edf_bill_fetcher.io.dispute_workbook import HEADERS, SHEET
from edf_bill_fetcher.io.persistence.result_codec import (
    _validate_ledger,
    decode_result,
    encode_result,
)
from edf_bill_fetcher.io.review_workbook import load_review_workbook
from edf_bill_fetcher.io.writers.export import render_xlsx
from edf_bill_fetcher.models.rebilling_tables import comparison_table, schedule_table
from edf_bill_fetcher.models.report_view import RenderSelection, build_report_view
from edf_bill_fetcher.processors.bill_facts import bill_facts
from edf_bill_fetcher.processors.rebilling_ledger import (
    build_ledger,
    component_facts,
    components,
    money,
    schedule_digest,
    validate_schedule,
)
from edf_bill_fetcher.processors.revisions import build_revisions
from tests.test_excel_review import amend
from tests.test_excel_review import review_pack as review_pack
from tests.test_native_review import analyse_session, decision, items, row_id
from tests.test_qt_desktop import result_snapshot


def stored_ledger(snapshot: AnalysisSnapshot) -> dict:
    model = build_report_view(snapshot, RenderSelection()).rebilling_ledger
    assert model is not None
    return model


def position(key, **overrides):
    return {
        "row_id": key,
        "source_digest": "a" * 64,
        "status": "challenged",
        "challenged_amount": "150.00",
        "proposed_reduction": "150.00",
        "first_demand_date": None,
        "basis": "Challenge the full extra charge pending its earlier billing history",
        "edf_request": "Earliest demand and allocation of the extra units, including all intervening credits",
        **overrides,
    }


def set_position(session, key, **overrides):
    values = position(key, **overrides)
    values.pop("row_id")
    values.pop("source_digest")
    session.set_position(key, **values)


def ledger(data, decisions=None, positions=None, **options):
    return build_ledger(data, build_revisions(data, decisions, **options), positions)


def test_full_history_retains_higher_original_estimates_credit_once_and_partial_replacements():
    data = items()
    data[0]["record"].update(
        {"Period Charge (£)": 400, "Sub Periods": "01/01/2024|31/01/2024|1000|40|400"}
    )
    data[1]["record"].update(
        {"Period Charge (£)": 200, "Sub Periods": "01/01/2024|31/01/2024|500|40|200"}
    )
    final = copy.deepcopy(data[1])
    final.update(id="d")
    final["record"].update(
        {
            "Invoice #": "INV-3",
            "Date": "03/02/2024",
            "Period Charge (£)": 350,
            "Reading": "Actual",
            "Sub Periods": "01/01/2024|31/01/2024|700|50|350",
        }
    )
    data.append(final)
    links = [decision(), decision("b", "d"), decision("a", "c", kind="cancellation_credit")]
    model = ledger(data, links)
    assert len(model["documents"]) == 4
    assert [doc["row_id"] for doc in model["documents"]].count("c") == 1
    assert (
        next(doc for doc in model["documents"] if doc["row_id"] == "a")["record"][
            "Period Charge (£)"
        ]
        == 400
    )
    change = next(
        link for link in model["comparisons"] if link["older"] == "b" and link["newer"] == "d"
    )["components"][0]
    assert (
        change["charge_difference"],
        change["units_effect"],
        change["rate_effect"],
        change["residual"],
    ) == ("150.00", "80.00", "70.00", "0.00")
    assert "not write-offs" in model["notice"]
    _validate_ledger(model)
    data[1]["record"]["Period From"] = "15/01/2024"
    data[1]["record"]["Sub Periods"] = "15/01/2024|31/01/2024|500|40|200"
    partial = next(
        link
        for link in ledger(data, links)["comparisons"]
        if link["older"] == "a" and link["newer"] == "b"
    )
    assert not partial["same_period"]
    assert partial["components"][0]["status"] == "unallocated"
    assert partial["components"][0]["charge_difference"] is None


def test_ambiguous_register_components_never_get_a_fabricated_allocation():
    data = items()
    data[0]["record"]["Sub Periods"] = (
        "01/01/2024|31/01/2024|100|20|20; 01/01/2024|31/01/2024|50|10|5"
    )
    data[1]["record"]["Sub Periods"] = "01/01/2024|31/01/2024|200|20|40"
    model = ledger(data)
    assert model["comparisons"][0]["components"][0]["status"] == "unallocated"
    assert len(model["documents"][0]["components"]) == 2


def test_standing_vat_and_cancellations_are_observations_not_extra_transactions():
    old_text = "Standing charge (31 days @ 35.00p per day) £10.85\nVAT @ 5% of £110.85 £5.54"
    new_text = "Standing charge 31 days @ 50.000p/day £15.50\nVAT @ 5% on £215.50 £10.78\nWe cancelled your electricity charges (excluding VAT) £110.85\nCancelled VAT (5% on £110.85) was £5.54 credit\nDirect Debit Discount @ 7% of £100.00 £7.00"
    data = items()
    data[0]["record"]["Rebilling Components"] = component_facts(old_text)
    data[1]["record"]["Rebilling Components"] = component_facts(new_text)
    model = ledger(data)
    other = {part["kind"]: part for part in model["comparisons"][0]["other_components"]}
    assert other["standing"]["difference"] == "4.65"
    assert other["vat"]["difference"] == "5.24"
    assert other["discount"]["difference"] is None
    assert [
        p["amount"] for p in model["documents"][1]["components"] if p["kind"] == "cancellation"
    ] == ["-110.85", "-5.54"]
    assert "4.65" in str(comparison_table(model))
    data[1]["record"]["Period To"] = "15/02/2024"
    assert all(p["difference"] is None for p in ledger(data)["comparisons"][0]["other_components"])
    _validate_ledger(model)


def test_new_format_energy_sections_keep_their_own_periods_and_vat_basis():
    text = """Freedom (1 January 2024 - 31 January 2024)
Electricity charges for meter EXAMPLE
100.0000 kWh @
Electricity used 20.000p/kWh £20.00
Standing charge 31 days @ 50.000p/day £15.50
Subtotal of charges before VAT £35.50
VAT @ 5% on £35.50 £1.78
Total electricity charges for this period £37.28
Freedom (1 February 2024 - 29 February 2024)
200.0000 kWh @
Electricity used 30.000p/kWh £60.00
"""
    facts = bill_facts(text, [])
    assert facts["Gross Electricity Charges Basis"] == "Including VAT"
    observed = components(facts)
    energy = [part for part in observed if part["kind"] == "energy"]
    assert len(energy) == 2
    assert (energy[0]["period_from"], energy[1]["period_from"]) == ("2024-01-01", "2024-02-01")
    assert energy[0]["amount"] == "20.00"
    assert components({"Source PDF Text": text}) == observed


def test_no_inferred_first_demand_from_invoice_statement_or_reading_dates():
    data = items()
    data[1]["record"].update({"Reading Date": "20/01/2024", "First Billed Date": "01/02/2024"})
    doc = next(doc for doc in ledger(data)["documents"] if doc["row_id"] == "b")
    assert doc["invoice_date"] == "2024-02-02"
    assert doc["reading_date"] == "2024-01-20"
    assert doc["first_demand_date"] is None
    assert doc["cutoff_date"] is None
    p = position("b", first_demand_date="29/02/2024", source_digest=schedule_digest(data, [], "b"))
    doc = next(doc for doc in ledger(data, positions=[p])["documents"] if doc["row_id"] == "b")
    assert (doc["first_demand_date"], doc["cutoff_date"], doc["needs_review"]) == (
        "2024-02-29",
        "2023-02-28",
        False,
    )


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("status", "agreed write-off"),
        ("basis", " "),
        ("basis", 123),
        ("edf_request", "x" * 4001),
        ("challenged_amount", -1),
        ("challenged_amount", True),
        ("challenged_amount", "nan"),
        ("challenged_amount", "1.001"),
        ("challenged_amount", "1000000000001"),
        ("proposed_reduction", "151"),
        ("first_demand_date", "31/02/2024"),
        ("first_demand_date", "last July"),
        ("source_digest", "changed"),
        ("row_id", "c"),
    ],
)
def test_invalid_positions_fail_before_any_export(field, value):
    with pytest.raises(ValueError):
        validate_schedule(items(), [position("b", **{field: value})])


def test_net_credit_does_not_cap_a_positive_replacement_component_claim_and_no_totals_are_inferred():
    data = items()
    data[1]["record"].update({"Period Charge (£)": -20, "Gross Electricity Charges (£)": 500})
    positions = validate_schedule(
        data, [position("a"), position("b", challenged_amount=500, proposed_reduction=500)]
    )
    model = ledger(data, positions=positions)
    assert len(schedule_table(model)) == 3
    assert positions[1]["proposed_reduction"] == "500.00"
    assert not any("total" in key for key in model)
    assert "not summed" in model["notice"]
    assert money(Decimal("-.001")) == "0.00"
    assert money(Decimal("1.005")) == "1.01"
    assert money(None) is None


def test_capped_suggestions_never_remove_documents_or_explicit_history():
    model = ledger(items(), [decision()], limit=0)
    assert len(model["documents"]) == 3
    assert model["suggestions_truncated"]
    assert (
        next(doc for doc in model["documents"] if doc["row_id"] == "b")["incoming"][0]["status"]
        == "confirmed"
    )


def test_positions_undo_recover_and_become_stale_after_evidence_or_relationship_edits(
    review_pack, tmp_path
):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    a, b = row_id(session, "INV-A"), row_id(session, "INV-B")
    set_position(session, b, first_demand_date="01/02/2024")
    session.save(tmp_path / "draft.json")
    restored = ReviewSession.load(snapshot, tmp_path / "draft.json")
    assert restored.schedule == session.schedule
    session.undo()
    assert not session.schedule
    session.redo()
    session.set_revision(a, b, "replacement", "confirmed", "Full source history checked")
    assert next(doc for doc in session.ledger_model()["documents"] if doc["row_id"] == b)[
        "needs_review"
    ]
    set_position(session, b, first_demand_date="01/02/2024")
    assert not next(doc for doc in session.ledger_model()["documents"] if doc["row_id"] == b)[
        "needs_review"
    ]
    session.update(a, {"Period Charge (£)": 105}, True, "Correct earlier amount", "verified")
    assert next(doc for doc in session.ledger_model()["documents"] if doc["row_id"] == b)[
        "needs_review"
    ]
    accepted = analyse_session(session, path)
    assert accepted.summary.charges == snapshot.summary.charges + Decimal("5.00")
    assert accepted.summary.closing_balance == snapshot.summary.closing_balance
    assert "NEEDS REVIEW" in str(schedule_table(stored_ledger(accepted)))
    assert decode_result(encode_result(accepted, "ledger-roundtrip")) == accepted
    session.clear_position(b)
    assert not session.schedule
    session.undo()
    assert b in session.schedule


def edit_schedule(path, invoice="INV-B", **values):
    wb = load_workbook(path)
    ws = wb[SHEET]
    n = next(
        n
        for n in range(5, ws.max_row + 1)
        if ws.cell(n, HEADERS.index("Invoice") + 1).value == invoice
    )
    for name, value in values.items():
        ws.cell(n, HEADERS.index(name) + 1).value = value
    wb.save(path)
    wb.close()


def test_excel_positions_regenerate_all_formats_and_match_native_review(review_pack, tmp_path):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    key = row_id(session, "INV-B")
    set_position(session, key, first_demand_date="01/02/2024")
    native = analyse_session(session, path)
    edit_schedule(
        path,
        **{
            "Status": "challenged",
            "Challenged amount (£)": 150,
            "Proposed reduction (£)": 150,
            "First demand date": "2024-02-01",
            "Basis / scope": position(key)["basis"],
            "EDF must substantiate": position(key)["edf_request"],
        },
    )
    originals = path.read_bytes()
    imported = load_review_workbook(str(path))
    assert imported is not None
    records, state, _ = imported
    assert state["dispute_schedule"] == list(session.schedule.values())
    assert any(change["field"] == "Dispute position" for change in state["review_history"])
    excel = analyse_prepared(
        prepare_snapshot(
            capture_legacy(
                records,
                state,
                replace(
                    snapshot.prepared.extraction.request, config=freeze_json(state["review_config"])
                ),
            )
        )
    )
    assert excel.summary == native.summary == snapshot.summary
    assert (
        build_report_view(excel, RenderSelection()).rebilling_ledger
        == build_report_view(native, RenderSelection()).rebilling_ledger
    )
    directory, outcome = regenerate_workbook(
        str(path), str(tmp_path), sections=("detailed_findings",)
    )
    assert outcome.status in {"complete", "review_required"}, outcome.error
    assert outcome.snapshot is not None
    assert outcome.snapshot.summary == snapshot.summary
    assert path.read_bytes() == originals
    from pathlib import Path

    import pdfplumber

    root = Path(directory)
    html = (root / "Reviewed_Evidence_Report.html").read_text()
    doc = Document(str(root / "Reviewed_Evidence_Report.docx"))
    doc_text = " ".join(p.text for p in doc.paragraphs) + " ".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    with pdfplumber.open(root / "Reviewed_Evidence_Report.pdf") as pdf:
        pdf_text = " ".join(page.extract_text() or "" for page in pdf.pages)
    for content in (html, doc_text, pdf_text):
        assert "Dispute schedule" in content
        assert "150.00" in content
        assert "Earliest" in content
        assert "demand" in content
    assert load_review_workbook(str(root / "Reviewed_Evidence.xlsx")) is not None
    wb = load_workbook(root / "Reviewed_Evidence.xlsx")
    assert {"Rebilling ledger", "Rebilling components", SHEET} <= set(wb.sheetnames)
    assert wb[SHEET]["B5"].protection.locked is False
    assert wb[SHEET]["G5"].protection.locked is True
    assert wb["Rebilling ledger"]["R5"].hyperlink.target.endswith("#page=1")
    wb.close()


@pytest.mark.parametrize(
    "kind",
    [
        "formula",
        "negative",
        "reduction",
        "date",
        "reason",
        "identity",
        "delete",
        "duplicate",
        "sheet",
        "header",
    ],
)
def test_excel_rejects_misleading_dispute_inputs_without_creating_a_pack(
    review_pack, tmp_path, kind
):
    path, _, _ = review_pack
    if kind in {"formula", "negative", "reduction", "date", "reason", "identity"}:
        values = {
            "Status": "challenged",
            "Challenged amount (£)": 150,
            "Proposed reduction (£)": 150,
            "Basis / scope": "Checked source",
        }
        name, value = {
            "formula": ("Challenged amount (£)", "=100+50"),
            "negative": ("Challenged amount (£)", -1),
            "reduction": ("Proposed reduction (£)", 151),
            "date": ("First demand date", "next week"),
            "reason": ("Basis / scope", ""),
            "identity": ("Invoice", "Wrong invoice"),
        }[kind]
        values[name] = value
        edit_schedule(path, **values)
    else:
        wb = load_workbook(path)
        ws = wb[SHEET]
        if kind == "delete":
            ws.delete_rows(5)
        elif kind == "duplicate":
            for c in range(1, len(HEADERS) + 1):
                ws.cell(6, c).value = ws.cell(5, c).value
        elif kind == "sheet":
            del wb[SHEET]
        else:
            ws.cell(4, 1).value = "Wrong header"
        wb.save(path)
        wb.close()
    with pytest.raises(ValueError):
        regenerate_workbook(str(path), str(tmp_path))
    assert not list(tmp_path.glob("reviewed_*"))


def test_unchanged_excel_position_stays_stale_after_a_record_correction_and_exclusion(
    review_pack, tmp_path
):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    key = row_id(session, "INV-B")
    set_position(session, key)
    accepted = analyse_session(session, path)
    target = tmp_path / "positions.xlsx"
    render_xlsx(build_report_view(accepted, RenderSelection()), str(target))
    amend(
        target,
        "INV-B",
        **{
            "Period Charge (£)": 60,
            "Include": "No",
            "Reason for change": "Not part of this complaint",
        },
    )
    imported = load_review_workbook(str(target))
    assert imported is not None
    records, state, _ = imported
    reviewed = analyse_prepared(
        prepare_snapshot(capture_legacy(records, state, accepted.prepared.extraction.request))
    )
    doc = next(doc for doc in stored_ledger(reviewed)["documents"] if doc["row_id"] == key)
    assert doc["needs_review"]
    assert not doc["included"]
    assert "EXCLUDED" in str(schedule_table(stored_ledger(reviewed)))
    assert doc["position"]["proposed_reduction"] == "150.00"


@pytest.mark.parametrize(
    "mutation",
    [
        "version",
        "document",
        "duplicate",
        "date",
        "position",
        "component",
        "edge",
        "comparison",
        "amount",
    ],
)
def test_saved_ledger_rejects_corrupt_or_misleading_models_without_recalculation(mutation):
    data = items()
    data[0]["record"]["Sub Periods"] = "01/01/2024|31/01/2024|100|20|20"
    data[1]["record"]["Sub Periods"] = "01/01/2024|31/01/2024|200|20|40"
    model = ledger(data, positions=[position("b")])
    if mutation == "version":
        model["version"] = True
    elif mutation == "document":
        model["documents"][0]["included"] = "Yes"
    elif mutation == "duplicate":
        model["documents"].append(copy.deepcopy(model["documents"][0]))
    elif mutation == "date":
        model["documents"][0]["reading_date"] = "2024-02-31"
    elif mutation == "position":
        model["documents"][1]["position"]["proposed_reduction"] = -10
    elif mutation == "component":
        model["documents"][0]["components"][0]["kind"] = "refund entitlement"
    elif mutation == "edge":
        model["documents"][0]["outgoing"][0]["newer"] = "missing"
    elif mutation == "comparison":
        model["comparisons"][0]["newer"] = "missing"
    else:
        model["comparisons"][0]["components"][0]["charge_difference"] = True
    with pytest.raises(ValueError):
        _validate_ledger(model)


def test_old_accepted_results_do_not_silently_gain_recomputed_ledger_tables():
    snapshot = result_snapshot()
    encoded = encode_result(snapshot, "old-without-ledger")
    encoded["analysis"]["models"]["sap"]["payload"].pop("rebilling_ledger", None)
    restored = decode_result(encoded)
    assert build_report_view(restored, RenderSelection()).rebilling_ledger is None
    assert len(ReviewSession(restored).ledger_model()["documents"]) >= 1


def test_displayed_units_rate_and_residual_reconcile_to_the_observed_penny_difference():
    data = items()
    data[0]["record"]["Sub Periods"] = "01/01/2024|31/01/2024|3|33.333|1.00"
    data[1]["record"]["Sub Periods"] = "01/01/2024|31/01/2024|4|66.666|2.67"
    change = ledger(data)["comparisons"][0]["components"][0]
    assert change["residual"] == "0.01"
    assert sum(
        Decimal(change[field]) for field in ("units_effect", "rate_effect", "residual")
    ) == Decimal(change["charge_difference"])


def test_absorbed_duplicate_proposals_are_archived_without_transferring_or_summing_them(
    review_pack,
):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    a, b = row_id(session, "INV-A"), row_id(session, "INV-B")
    set_position(session, b)
    original = session.record(a)
    session.update(
        b,
        {field: original[field] for field in ("Invoice #", "Period Charge (£)", "Balance (£)")},
        True,
        "Duplicate extraction verified against original",
        "verified",
    )
    merged = analyse_session(session, path)
    reopened = ReviewSession(merged)
    assert not reopened.schedule
    state = object_value(merged.prepared.extraction.legacy_state)
    assert state["archived_dispute_schedule"][0]["row_id"] == b
    assert any(
        change["field"] == "Dispute position" and change["after"] == "archived_after_deduplication"
        for change in state["review_history"]
    )


def test_earlier_only_components_stay_unallocated_and_visible():
    data = items()
    data[0]["record"]["Sub Periods"] = (
        "01/01/2024|14/01/2024|100|20|20; 15/01/2024|31/01/2024|200|20|40"
    )
    data[1]["record"]["Sub Periods"] = "15/01/2024|31/01/2024|250|20|50"
    comparisons = ledger(data)["comparisons"][0]["components"]
    assert len(comparisons) == 2
    assert {part["status"] for part in comparisons} == {"matched", "unallocated"}
    assert (
        next(part for part in comparisons if part["status"] == "unallocated")["period_from"]
        == "2024-01-01"
    )


def test_excel_can_clear_a_position_with_an_audited_round_trip(review_pack, tmp_path):
    path, snapshot, _ = review_pack
    session = ReviewSession(snapshot)
    set_position(session, row_id(session, "INV-B"))
    accepted = analyse_session(session, path)
    target = tmp_path / "clear.xlsx"
    render_xlsx(build_report_view(accepted, RenderSelection()), str(target))
    edit_schedule(
        target,
        **{
            "Status": "unreviewed",
            "Challenged amount (£)": None,
            "Proposed reduction (£)": None,
            "First demand date": None,
            "Basis / scope": None,
            "EDF must substantiate": None,
        },
    )
    imported = load_review_workbook(str(target))
    assert imported is not None
    assert not imported[1]["dispute_schedule"]
    assert any(
        change["field"] == "Dispute position" and change["after"] is None
        for change in imported[1]["review_history"]
    )
