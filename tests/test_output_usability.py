"""Document navigation and financial meaning must survive presentation changes."""

from __future__ import annotations

import re
from dataclasses import replace
from types import SimpleNamespace

import pandas as pd
import pdfplumber
import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.section import WD_ORIENT
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from edf_bill_fetcher.helpers.formatting import fmt_money
from edf_bill_fetcher.io.persistence.result_codec import encode_result
from edf_bill_fetcher.io.reporters.docx_report import render_docx
from edf_bill_fetcher.io.reporters.html_report import render_html
from edf_bill_fetcher.io.reporters.pdf_report import render_pdf
from edf_bill_fetcher.io.review_workbook import HEADERS, REVIEW_SHEET, load_review_workbook
from edf_bill_fetcher.io.writers.export import render_xlsx
from edf_bill_fetcher.models.output_overview import build_output_overview, report_account_label
from edf_bill_fetcher.models.output_tables import EVIDENCE_HEADERS, evidence_cells, timeline_table
from edf_bill_fetcher.models.report_sections import REPORT_SECTIONS
from edf_bill_fetcher.models.report_view import RenderSelection, capture_report_view

INVOICE = "KI-123456789-REFERENCE-END"
FILENAME = "original-bill-with-a-long-file-name-REFERENCE-END.pdf"


def bill(**overrides):
    return {
        "Date": "31/01/2024",
        "Period From": "01/01/2024",
        "Period To": "31/01/2024",
        "Entry Type": "New Bill",
        "Invoice #": INVOICE,
        "Period Charge (£)": 140,
        "Balance (£)": 250,
        "Account": "000123",
        "Fuel": "Electricity",
        "Source": "Local PDF",
        "Attachment Name": FILENAME,
        "Page From": 1,
        "Page To": 2,
        "Energy Unit Rate (p/kWh)": 28,
        "Units (kWh)": 400,
        "Payment Method": "Direct Debit",
        "Tariff Type": "Variable",
        "Reading": "Actual",
        "Details": "Verify source <text> & amounts",
        **overrides,
    }


def view_for(records, **config):
    return capture_report_view(
        records,
        {"as_of": "2024-02-01", "use_dedup": False, **config},
        formats=("xlsx", "pdf", "docx", "html"),
    )


@pytest.fixture(scope="module")
def output_pack(tmp_path_factory):
    folder = tmp_path_factory.mktemp("usable-outputs")
    view = capture_report_view(
        [bill(**{"OCR Review Required": True})],
        {"as_of": "2024-02-01", "use_dedup": False},
        SimpleNamespace(error_log=["[12:00:00] Attachment — Check source <scan> & amounts"]),
        formats=("xlsx", "pdf", "docx", "html"),
    )
    before = encode_result(view.snapshot, "usability")
    for extension, render in (
        ("pdf", render_pdf),
        ("docx", render_docx),
        ("html", render_html),
        ("xlsx", render_xlsx),
    ):
        render(view, str(folder / ("report." + extension)))
    assert encode_result(view.snapshot, "usability") == before
    return folder, view


def test_pdf_retains_sections_and_source_references_without_page_bloat(output_pack):
    folder, _ = output_pack
    with pdfplumber.open(folder / "report.pdf") as document:
        assert len(document.pages) <= 10
        texts = [page.extract_text() or "" for page in document.pages]
        assert all(len(text.strip()) > 80 for text in texts)
        # Extract cells separately: line-wrapped PDF columns otherwise interleave.
        cells = {
            re.sub(r"\s+", "", cell or "")
            for table in document.pages[-1].extract_tables()
            for row in table
            for cell in row
        }
        assert INVOICE in cells
        assert FILENAME in cells
        compact = re.sub(r"\s+", "", "\n".join(texts))
        assert "<b>PDFBills:</b>" not in compact
        outlines = list(document.doc.get_outlines())
        assert len(outlines) == len(REPORT_SECTIONS)
        page_ids = {page.page_obj.pageid for page in document.pages}
        assert all(destination[0].objid in page_ids for _, _, destination, _, _ in outlines)
        toc_links = document.pages[1].annots
        assert len(toc_links) == len(REPORT_SECTIONS)
        assert all(link["data"]["Dest"][0].objid in page_ids for link in toc_links)
        assert document.pages[-1].width > document.pages[-1].height
        assert "Balance(£)" in cells
        assert "Transaction(£)" in cells


def test_pdf_without_cover_still_has_headers_and_correct_page_numbers(output_pack, tmp_path):
    _, original = output_pack
    view = replace(
        original,
        selection=RenderSelection(("pdf",), ("exec_summary", "appendix_full_evidence")),
    )
    path = tmp_path / "no-cover.pdf"
    render_pdf(view, str(path))
    with pdfplumber.open(path) as document:
        text = document.pages[0].extract_text()
        assert "EDF Energy" in text
        assert "Page 1" in text
        assert document.pages[0].width < document.pages[0].height
        assert document.pages[-1].width > document.pages[-1].height


@pytest.mark.parametrize("framing", [(), ("cover",), ("toc",)])
def test_appendix_only_selection_has_no_empty_opening_page(tmp_path, framing):
    original = view_for([bill()])
    view = replace(
        original,
        selection=RenderSelection(("pdf", "docx"), (*framing, "appendix_full_evidence")),
    )
    render_pdf(view, str(tmp_path / "appendix.pdf"))
    with pdfplumber.open(tmp_path / "appendix.pdf") as document:
        assert len(document.pages) == 1 + len(framing)
        assert all(len((page.extract_text() or "").strip()) > 80 for page in document.pages)
        assert document.pages[-1].width > document.pages[-1].height
    render_docx(view, str(tmp_path / "appendix.docx"))
    doc = Document(tmp_path / "appendix.docx")
    assert len(doc.sections) == 1 + len(framing)
    assert not doc._element.xpath('//w:br[@w:type="page"]')
    assert doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE


def test_word_has_working_navigation_a4_pages_and_repeating_table_headings(output_pack):
    folder, _ = output_pack
    doc = Document(folder / "report.docx")
    bookmarks = set(doc._element.xpath("//w:bookmarkStart/@w:name"))
    targets = set(doc._element.xpath("//w:hyperlink/@w:anchor"))
    expected = {"sec_" + section.key for section in REPORT_SECTIONS}
    assert targets == expected
    assert targets <= bookmarks
    assert doc.styles["SectionHeader"].element.xpath("./w:pPr/w:outlineLvl/@w:val") == ["0"]
    assert doc.sections[0].page_width is not None
    assert doc.sections[0].page_height is not None
    assert doc.sections[-1].page_width is not None
    assert round(doc.sections[0].page_width.cm, 1) == 21
    assert round(doc.sections[0].page_height.cm, 1) == 29.7
    assert doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE
    assert round(doc.sections[-1].page_width.cm, 1) == 29.7
    appendix = doc.tables[-1]
    assert appendix.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    cells = [cell.text for row in appendix.rows for cell in row.cells]
    assert INVOICE in cells
    assert FILENAME in cells
    assert "1–2" in cells
    assert "£250.00" in cells
    assert "£140.00" in cells
    assert len(doc._element.xpath('//w:br[@w:type="page"]')) == 2  # Cover and contents.


def test_html_navigation_offline_access_and_user_text(output_pack):
    folder, _ = output_pack
    source = (folder / "report.html").read_text()
    soup = BeautifulSoup(source, "html.parser")
    ids = {element["id"] for element in soup.select("[id]")}
    assert all(link["href"][1:] in ids for link in soup.select('a[href^="#"]'))
    assert len(soup.select("nav[aria-label] a")) == len(REPORT_SECTIONS)
    assert soup.select_one('a[href="#report-content"]')
    assert soup.select_one('meta[name="viewport"]')
    assert soup.select("table thead th[scope=col]")
    assert not soup.select("script, link, iframe, img[src]")
    assert "Verify source <text> & amounts" in soup.get_text()
    assert "Read 1 extraction / review note" in soup.get_text()
    assert "&amp;lt;" not in source
    assert FILENAME in soup.get_text()
    assert INVOICE in soup.get_text()


def test_excel_dashboard_links_cover_the_visible_workbook_and_edits_remain_lossless(output_pack):
    folder, view = output_pack
    path = folder / "report.xlsx"
    wb = load_workbook(path)
    ws = wb["Start here"]
    destinations = {
        cell.hyperlink.target for row in ws for cell in row if cell.hyperlink is not None
    }
    assert destinations == {
        f"#'{sheet.title.replace(chr(39), chr(39) * 2)}'!A1"
        for sheet in wb
        if sheet.sheet_state == "visible" and sheet.title != "Start here"
    }
    assert ws["A6"].value == 140
    assert ws["A10"].value == 250
    assert ws["E10"].value == 250
    assert ws["I10"].value == 1
    assert ws["A6"].data_type == "n"
    assert "regenerate" in ws["A3"].value
    evidence = wb["EDF Evidence Report"]
    columns = {cell.value: get_column_letter(cell.column) for cell in evidence[1]}
    for name in ("Balance (£)", "Transaction Amount (£)", "Account", "Fuel"):
        assert not evidence.column_dimensions[columns[name]].hidden
    assert evidence.column_dimensions[columns["Amount (£)"]].hidden
    assert evidence.freeze_panes == "G2"
    review = wb[REVIEW_SHEET]
    assert review.protection.formatColumns is False
    assert review.protection.formatRows is False
    assert review.column_dimensions[get_column_letter(HEADERS.index("Meter") + 1)].hidden
    imported = load_review_workbook(str(path))
    assert imported is not None
    _, state, as_of = imported
    assert as_of == "2024-02-01"
    assert not state["review_history"]
    assert all(not item["changes"] for item in state["human_review"]["decisions"])
    assert view.snapshot.summary.charges == 140


def test_mixed_ledger_overview_never_presents_a_combined_account_balance(tmp_path):
    view = view_for([bill(), bill(**{"Account": "000456", "Invoice #": "OTHER"})])
    overview = build_output_overview(view)
    assert report_account_label(view) == "Multiple accounts (2)"
    assert overview.ledger_count == 2
    assert "mixed ledgers" in overview.balance_note
    assert dict(overview.metrics)["Last observed balance"] == "Unavailable"
    path = tmp_path / "mixed.xlsx"
    render_xlsx(view, str(path))
    wb = load_workbook(path)
    assert wb["Start here"]["A10"].value == "Unavailable"
    assert wb["Start here"]["E10"].value == "Unavailable"
    path = tmp_path / "mixed.docx"
    render_docx(view, str(path))
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "Average observed balance" not in text


def test_review_priorities_do_not_treat_payments_as_incomplete_bills():
    payment = view_for(
        [{"Date": "01/01/2024", "Entry Type": "Payment", "Transaction Amount (£)": 50}]
    )
    titles = [priority.title for priority in build_output_overview(payment).priorities]
    assert "Complete the bill references and amounts" not in titles
    assert "No payment, credit or refund rows supplied" not in titles
    missing = view_for([bill(**{"Invoice #": "", "Period From": "", "Period Charge (£)": None})])
    actions = " ".join(p.action for p in build_output_overview(missing).priorities)
    assert "missing an invoice reference: 1" in actions
    assert "missing/invalid period: 1" in actions
    assert "missing charge: 1" in actions


def test_unknown_account_and_amounts_are_not_manufactured():
    view = view_for([bill(), bill(**{"Account": "", "Invoice #": "UNASSIGNED"})])
    assert report_account_label(view) == "000123 + unassigned records"
    assert report_account_label(view_for([bill(**{"Account": ""})])) == "Not supplied"
    row = bill(**{"Balance (£)": None, "Transaction Amount (£)": 0})
    cells = dict(zip(EVIDENCE_HEADERS, evidence_cells(row), strict=True))
    assert cells["Balance (£)"] == ""
    assert cells["Transaction (£)"] == "£0.00"
    assert cells["Attachment"] == FILENAME
    assert cells["Invoice #"] == INVOICE
    events = timeline_table(pd.DataFrame([row]), [("CHECK", "31/01/2024", 999, "Inspect", "HIGH")])
    assert events[1][2:4] == ["", "£0.00"]
    assert events[2][2:4] == ["", ""]
    assert "Flag value: £999.00" in events[2][4]


def test_html_forecast_has_real_months_and_formatted_stored_predictions(tmp_path):
    records = [
        bill(
            **{
                "Date": f"28/{month:02}/2024",
                "Invoice #": str(month),
                "Balance (£)": 100 + month * 10,
            }
        )
        for month in range(1, 7)
    ]
    view = view_for(records)
    path = tmp_path / "forecast.html"
    render_html(view, str(path))
    soup = BeautifulSoup(path.read_text(), "html.parser")
    section = soup.select_one("#sec-forecast")
    assert section is not None
    text = section.get_text(" ")
    assert "Jul 2024" in text
    assert "Dec 2024" in text
    assert fmt_money(view.model("forecast").linear_forecast[0]) in text
    assert view.model("forecast").hw_forecast is not None
    assert fmt_money(view.model("forecast").hw_forecast[0]) in text
    assert "Linear Forecast" not in text  # No raw dataclass field labels/list dumps.
    assert len(section.select("tbody tr")) == 6
