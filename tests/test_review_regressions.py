"""End-to-end regressions for evidence preservation and record identity."""

from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from edf_bill_fetcher.collectors.engine import EvidenceEngine
from edf_bill_fetcher.io.adapters.html import parse_htm_account_history
from edf_bill_fetcher.io.writers.evidence_bundle import save_evidence_files
from edf_bill_fetcher.io.writers.export import export_to_excel
from edf_bill_fetcher.models.records import BillingRecord
from edf_bill_fetcher.processors.run_diff import diff_records


def test_missing_period_does_not_merge_bills_a_year_apart(tmp_path):
    rows = [
        BillingRecord(
            source="Local PDF Folder",
            entry_type="New Bill",
            logic_used="test",
            date=date,
            invoice_num=invoice,
            amount=100,
            period_charge=100,
        ).to_dict()
        for date, invoice in [("01/01/2024", "INV-A"), ("01/01/2025", "INV-B")]
    ]
    output = tmp_path / "evidence.xlsx"
    export_to_excel(rows, str(output), [], {"use_dedup": True, "save_dups": True})
    wb = load_workbook(output)
    ws = wb["EDF Evidence Report"]
    col = [c.value for c in ws[1]].index("Invoice #") + 1
    assert {ws.cell(r, col).value for r in range(2, ws.max_row + 1)} == {"INV-A", "INV-B"}
    wb.close()


@pytest.mark.parametrize(
    "action",
    [
        "We charged your account £100.00",
        "You paid us £100.00",
        "Reversed account charge £100.00",
    ],
)
@pytest.mark.parametrize(("direction", "expected"), [("debit", 200), ("credit", -200)])
def test_html_balance_direction(action, direction, expected):
    rows = parse_htm_account_history(f"1 January 2024 {action} Balance £200.00 in {direction}")
    assert len(rows) == 1
    assert rows[0]["Amount (£)"] == expected
    assert rows[0]["Period Charge (£)"] == 100


def test_standalone_html_credit_is_negative():
    rows = parse_htm_account_history("1 January 2024 Balance £200.00 in credit")
    assert rows[0]["Amount (£)"] == -200


def _pdf(pages):
    result = MagicMock()
    result.__enter__.return_value.pages = [
        MagicMock(extract_text=lambda text=t: text) for t in pages
    ]
    return result


def test_merged_invoice_bundle_copies_parent_for_each_slice(tmp_path):
    source = tmp_path / "merged.pdf"
    source.write_bytes(b"%PDF parent bytes")
    eng = EvidenceEngine({}, lambda *_: None)
    pages = [
        f"Invoice number: KI-31105244-000{i}\nDate issued: 06 August 2024\nPage 1 of 1"
        for i in (1, 2)
    ]
    with (
        patch("edf_bill_fetcher.collectors.engine.pdfplumber.open", return_value=_pdf(pages)),
        patch.object(eng, "_process_new_invoice") as dispatch,
    ):
        eng.process_pdf_file(str(source), "Local PDF Folder", source.name, "01/01/2024")
    names = [c.kwargs["attachment_name"] for c in dispatch.call_args_list]
    assert len(names) == 2
    rows = pd.DataFrame({"Attachment Name": names, "Invoice #": ["KI-1", "KI-2"]})
    saved = save_evidence_files(rows, eng.source_paths, str(tmp_path / "bundle"))
    assert set(saved) == set(names)
    for path in saved.values():
        assert Path(path).suffix == ".pdf"
        assert Path(path).read_bytes() == source.read_bytes()


def test_pst_attachment_survives_until_bundle_export(tmp_path):
    eng = EvidenceEngine({}, lambda *_: None)
    att = MagicMock()
    att.get_size.return_value = 20
    att.read_buffer.return_value = b"%PDF attachment bytes"
    msg = MagicMock()
    msg.get_subject.return_value = "EDF invoice"
    msg.get_delivery_time.return_value = None
    msg.get_html_body.return_value = None
    msg.get_plain_text_body.return_value = None
    msg.get_rtf_body.return_value = None
    msg.get_number_of_attachments.return_value = 1
    msg.get_attachment.return_value = att
    folder = MagicMock()
    folder.get_number_of_sub_messages.return_value = 1
    folder.get_sub_message.return_value = msg
    folder.get_number_of_sub_folders.return_value = 0
    try:
        with (
            patch("edf_bill_fetcher.collectors.engine.HAS_PYPFF", True),
            patch(
                "edf_bill_fetcher.collectors.engine._pst_attachment_filename",
                return_value="bill.pdf",
            ),
            patch(
                "edf_bill_fetcher.collectors.engine.pdfplumber.open", return_value=_pdf(["invoice"])
            ),
            patch.object(eng, "process_text"),
        ):
            eng.crawl_pst(folder)
        source = Path(eng.source_paths["bill.pdf"])
        assert source.exists()
        saved = save_evidence_files(
            pd.DataFrame({"Attachment Name": ["bill.pdf"]}),
            eng.source_paths,
            str(tmp_path / "bundle"),
        )
        assert Path(saved["bill.pdf"]).read_bytes() == att.read_buffer.return_value
    finally:
        eng.cleanup()
    assert not source.exists()
    eng.cleanup()  # cleanup is idempotent


def _row(invoice="N/A", **extra):
    return {
        "Date": "2024-01-01",
        "Amount (£)": 100,
        "Source": "Local PDF",
        "Invoice #": invoice,
        **extra,
    }


def test_diff_identical_colliding_invoices_and_reordered_rows():
    rows = [_row("INV-A"), _row("INV-B")]
    assert diff_records(rows, rows[::-1]) == {"added": [], "removed": [], "changed": []}


def test_diff_preserves_duplicate_multiplicity():
    row = _row()
    assert diff_records([row, row], [row]) == {"added": [], "removed": [row], "changed": []}


def test_diff_matches_unchanged_occurrences_before_changed_rows():
    old = [_row(Details="one"), _row(Details="two")]
    new = [_row(Details="three"), _row(Details="one")]
    result = diff_records(old, new)
    assert result["added"] == result["removed"] == []
    assert len(result["changed"]) == 1
    assert result["changed"][0]["old_values"] == ["two"]
    assert result["changed"][0]["new_values"] == ["three"]


def test_diff_never_pairs_different_invoice_numbers():
    old, new = _row("INV-A"), _row("INV-B")
    assert diff_records([old], [new]) == {"added": [new], "removed": [old], "changed": []}


def test_diff_amount_change_keeps_invoice_identity():
    old = [_row("INV-A"), _row("INV-B")]
    new = [_row("INV-B"), {**_row("INV-A"), "Amount (£)": 150}]
    result = diff_records(old, new)
    assert result["added"] == result["removed"] == []
    assert len(result["changed"]) == 1
    assert result["changed"][0]["changed_fields"] == ["Amount (£)"]
    assert result["changed"][0]["new_row"]["Invoice #"] == "INV-A"


@pytest.mark.parametrize("fail", [False, True])
def test_cli_releases_attachments_on_empty_run_or_failure(tmp_path, monkeypatch, fail):
    from edf_bill_fetcher.io.cli import run_cli_extract

    engine = EvidenceEngine({}, lambda *_: None)
    with (
        patch("edf_bill_fetcher.collectors.engine.EvidenceEngine", return_value=engine),
        patch.object(
            engine, "crawl_local_pdfs", side_effect=ValueError("parse failed") if fail else None
        ),
        patch.object(engine, "cleanup", wraps=engine.cleanup) as cleanup,
        pytest.raises(SystemExit) as exc,
    ):
        run_cli_extract(["--pdf-dir", str(tmp_path), "-o", str(tmp_path / "out.xlsx")])
    assert exc.value.code == (1 if fail else 0)
    cleanup.assert_called_once()
