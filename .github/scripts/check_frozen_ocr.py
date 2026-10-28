"""Exercise bundled OCR through the delivered application's real extraction command."""

import argparse
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--executable", type=Path)
options = parser.parse_args()
candidates = (
    [options.executable]
    if options.executable
    else [
        Path("dist/EDF_Evidence_Collector.app/Contents/MacOS/EDF_Evidence_Collector"),
        Path(
            "dist/EDF_Evidence_Collector.exe" if os.name == "nt" else "dist/EDF_Evidence_Collector"
        ),
    ]
)
executable = next(path.resolve() for path in candidates if path.is_file())
with tempfile.TemporaryDirectory(prefix="edf-frozen-ocr-") as directory:
    root = Path(directory)
    inputs = root / "input"
    inputs.mkdir()
    shutil.copyfile("tests/fixtures/scanned_bill.pdf", inputs / "scanned.pdf")
    records_path = root / "records.json"
    diagnostic_path = root / "diagnostics.json"
    # Hiding PATH prevents accidentally exercising a machine-installed engine.
    environment = os.environ.copy()
    environment["PATH"] = os.environ.get("SystemRoot", "/usr") + (
        "/System32" if os.name == "nt" else "/bin"
    )
    subprocess.run(
        [str(executable), "--doctor", "--output", str(diagnostic_path)],
        env=environment,
        check=True,
        timeout=60,
    )
    diagnostics = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    assert diagnostics["ready"] is True, diagnostics
    assert diagnostics["build"]["mode"] == "packaged", diagnostics
    assert diagnostics["build"]["commit"] != "unknown", diagnostics
    subprocess.run(
        [
            str(executable),
            "--extract",
            "--pdf-dir",
            str(inputs),
            "--output",
            str(root / "report.xlsx"),
            "--records-json",
            str(records_path),
            "--no-filter-below",
            "--strict",
        ],
        env=environment,
        check=True,
        timeout=180,
    )
    result = json.loads(records_path.read_text(encoding="utf-8"))
    rows = result["records"]
    assert len(rows) == 1, rows
    row = rows[0]
    assert row["Invoice #"] == "KI-12345", row
    assert row["Balance (£)"] == 250, row
    assert row["Period Charge (£)"] == 125, row
    assert row["Text Extraction"] == "OCR", row
    assert row["OCR Pages"] == [1], row
    assert row["OCR Review Required"] is True, row
    assert result["status"] == "review_required", result["status"]

    # Exercise the delivered binary's Excel round trip too: the amended charge
    # must rebuild analysis without changing the source OCR observations.
    workbook_path = root / "report.xlsx"
    workbook = load_workbook(workbook_path)
    review = workbook["Review & Corrections"]
    headers = {cell.value: cell.column for cell in review[4]}
    review.cell(5, headers["Period Charge (£)"], 140)
    review.cell(5, headers["Reason for change"], "Frozen-build review smoke test")
    schedule = workbook["Dispute schedule"]
    schedule_headers = {cell.value: cell.column for cell in schedule[4]}
    for field, value in {
        "Status": "challenged",
        "Challenged amount (£)": 140,
        "Proposed reduction (£)": 100,
        "Basis / scope": "Synthetic reviewed challenge to historic additional charges",
        "EDF must substantiate": "Earliest demand and allocation of additional units",
        "Claim groups": "C1",
    }.items():
        schedule.cell(5, schedule_headers[field], value)
    for column, value in enumerate(
        (
            "C1",
            "Synthetic complaint smoke test",
            "reviewed",
            140,
            100,
            "One consolidated challenge checked against the original scanned bill",
        ),
        1,
    ):
        workbook["Claim groups"].cell(5, column, value)
    workbook.save(workbook_path)
    workbook.close()
    subprocess.run(
        [
            str(executable),
            "--review-workbook",
            "--workbook",
            str(workbook_path),
            "--output-dir",
            str(root / "reviewed"),
            "--formats",
            "xlsx,html",
        ],
        env=environment,
        check=True,
        timeout=180,
    )
    reviewed_path = next((root / "reviewed").glob("reviewed_*/result.json"))
    reviewed = json.loads(reviewed_path.read_text(encoding="utf-8"))
    assert reviewed["analysis"]["summary"]["charges"] == "140.00", reviewed["analysis"]["summary"]
    assert reviewed["evidence"]["observations"][0]["raw"]["Period Charge (£)"] == 125
    changes = reviewed["evidence"]["legacy_state"]["review_history"]
    assert {change["field"] for change in changes} == {
        "Period Charge (£)",
        "Transaction Amount (£)",
        "Dispute position",
        "Claim reconciliation",
    }
    assert (reviewed_path.parent / "Reviewed_Evidence.xlsx").is_file()
    assert (reviewed_path.parent / "Reviewed_Evidence_Report.html").is_file()
    claims = reviewed["analysis"]["models"]["sap"]["payload"]["claim_reconciliation"]
    assert claims["ready"] is True, claims
    assert claims["proposed_reduction"] == "100.00", claims
    focused = reviewed_path.parent / "complaint_pack"
    for name in (
        "Complaint_Draft.html",
        "Complaint_Draft.pdf",
        "Complaint_Draft.docx",
        "Claim_Schedule.xlsx",
        "evidence_manifest.json",
    ):
        assert (focused / name).is_file(), name
    assert (reviewed_path.parent / "Complaint_pack.zip").is_file()
    manifest = json.loads((focused / "evidence_manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["files"]) == 1, manifest
    assert manifest["files"][0]["status"] == "saved", manifest
    assert "100.00" in (focused / "Complaint_Draft.html").read_text(encoding="utf-8")
